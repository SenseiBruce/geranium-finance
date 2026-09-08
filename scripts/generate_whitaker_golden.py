#!/usr/bin/env python3
"""Generate golden loan underwriting decision workbook for Whitaker."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

from sanitize_office import cache_xlsx_formula_values, sanitize_xlsx
from whitaker_constants import (
    ADJ_EBITDA,
    APPLICATION_ID,
    AS_OF,
    BANKABLE,
    BUILDING_RP_NOLV,
    CNC_ADVANCE,
    CNC_DMG_NOLV,
    CNC_HAAS_NOLV,
    CNC_MAZAK_NOLV,
    COLLATERAL_CSV,
    COMMITMENT,
    COMMITTEE_DATE,
    CONDITIONS_TENOR_MO,
    COND_DSCR_FLOOR,
    DELIVERABLE,
    DSCR_FULL,
    DSCR_SIZED,
    ELIGIBLE_NOLV,
    ENTITY,
    EXISTING_ANNUAL_DS,
    EXISTING_DEBT_BAL,
    FIN_XLSX,
    FIXTURE_ADVANCE,
    FIXTURES_NOLV,
    FORKLIFT_LEASE_NOLV,
    LEVERAGE_SIZED,
    LTV_FULL,
    LTV_SIZED,
    MAX_BY_LTV,
    MAX_LEVERAGE,
    MAX_LTV,
    MEMO_TXT,
    MGMT_EBITDA,
    MIN_DSCR,
    OBSOLETE_NOLV,
    ONE_TIME_GAIN,
    OWNER_ADD_ALLOWED,
    OWNER_ADD_CLAIMED,
    PATH,
    PRICING_FACTOR_PER_1000_MO,
    PROPOSED_DS_FULL,
    PROPOSED_DS_SIZED,
    RELATED_PARTY_RENT_ADD_CLAIMED,
    REPORTED_EBITDA,
    REQUESTED,
    REQUESTED_TENOR_MO,
    SLUG,
    SOFT_COST,
    TOTAL_DEBT_SIZED,
    TOTAL_DS_FULL,
    TOTAL_DS_SIZED,
)

OUT = Path(__file__).resolve().parent.parent / "tasks" / SLUG / "golden" / DELIVERABLE

header_fill = PatternFill("solid", fgColor="37474F")
header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
title_font = Font(bold=True, size=13, name="Calibri")
section_font = Font(bold=True, size=11, name="Calibri")
body = Font(name="Calibri", size=10)
thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
pass_fill = PatternFill("solid", fgColor="C8E6C9")
warn_fill = PatternFill("solid", fgColor="FFECB3")
fail_fill = PatternFill("solid", fgColor="FFCDD2")
money_fmt = "#,##0.00"
ratio_fmt = "0.0000"
pct_fmt = "0.00%"


def money(x: float) -> float:
    return round(float(x), 2)


def style_header(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin


def border_row(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        ws.cell(row, c).border = thin
        if ws.cell(row, c).font.name is None:
            ws.cell(row, c).font = body


def autosize(ws, max_w: int = 42) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min(max(len(str(c.value or "")) for c in col) + 2, max_w)
        ws.column_dimensions[letter].width = max(width, 10)


def main() -> None:
    wb = Workbook()

    # ------------------------------------------------------------------
    # Assumptions
    # ------------------------------------------------------------------
    ws = wb.active
    ws.title = "Assumptions"
    ws["A1"] = f"{ENTITY}: credit decision assumptions"
    ws["A1"].font = title_font
    ws["A2"] = (
        f"Application {APPLICATION_ID} | Financials as-of {AS_OF} | "
        f"Committee {COMMITTEE_DATE} | Governing: {MEMO_TXT}"
    )
    ws["A2"].font = body

    headers = ["Parameter", "Value", "Source / rule"]
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header(ws, 4, 3)

    assumptions = [
        ("Requested commitment", REQUESTED, f"{FIN_XLSX} Application"),
        ("Requested tenor (months)", REQUESTED_TENOR_MO, f"{FIN_XLSX} Application"),
        ("Pricing factor per $1,000 / month", PRICING_FACTOR_PER_1000_MO, f"{MEMO_TXT} §4"),
        ("Existing bank annual debt service", EXISTING_ANNUAL_DS, f"{FIN_XLSX} Debt_Schedule"),
        ("Existing bank UPB", EXISTING_DEBT_BAL, f"{FIN_XLSX} Debt_Schedule"),
        ("Reported EBITDA FY2025", REPORTED_EBITDA, f"{FIN_XLSX} Income"),
        ("One-time parcel gain in reported EBITDA", ONE_TIME_GAIN, f"{FIN_XLSX} Income"),
        ("Owner add-back claimed (RM)", OWNER_ADD_CLAIMED, f"{FIN_XLSX} RM_Notes"),
        ("Owner add-back allowed (policy)", OWNER_ADD_ALLOWED, f"{MEMO_TXT} §3(b)"),
        ("Related-party rent add-back claimed", RELATED_PARTY_RENT_ADD_CLAIMED, f"{MEMO_TXT} §3(c) — disallow"),
        ("Minimum DSCR (Approve)", MIN_DSCR, f"{MEMO_TXT} §5"),
        ("Conditional DSCR floor", COND_DSCR_FLOOR, f"{MEMO_TXT} §5"),
        ("Maximum LTV", MAX_LTV, f"{MEMO_TXT} §7"),
        ("Maximum leverage", MAX_LEVERAGE, f"{MEMO_TXT} §5"),
        ("CNC advance rate", CNC_ADVANCE, f"{MEMO_TXT} §6"),
        ("Fixture advance rate", FIXTURE_ADVANCE, f"{MEMO_TXT} §6"),
        ("Hierarchy", "Memo supersedes RM notes", f"{MEMO_TXT} §1"),
    ]
    for r, (p, v, s) in enumerate(assumptions, 5):
        ws.cell(r, 1, p).font = body
        cell = ws.cell(r, 2, v)
        cell.font = body
        ws.cell(r, 3, s).font = body
        border_row(ws, r, 3)
        if isinstance(v, float) and v > 20:
            cell.number_format = money_fmt
        elif isinstance(v, float) and v <= 5:
            cell.number_format = ratio_fmt
    autosize(ws)

    # ------------------------------------------------------------------
    # EBITDA_Build
    # ------------------------------------------------------------------
    ws = wb.create_sheet("EBITDA_Build")
    ws["A1"] = "Adjusted EBITDA build (policy basis)"
    ws["A1"].font = title_font
    ws["A2"] = f"Sources: {FIN_XLSX} Income; rules in {MEMO_TXT} §3"
    headers = ["Step", "Amount", "Include", "Rationale"]
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header(ws, 4, 4)

    ebitda_rows = [
        ("Reported EBITDA FY2025", REPORTED_EBITDA, "Yes", "Starting point per memo §3"),
        ("Less: one-time parcel sale gain", -ONE_TIME_GAIN, "Yes", "Non-operating gain embedded in reported EBITDA"),
        ("Plus: owner add-back allowed", OWNER_ADD_ALLOWED, "Yes", f"Capped at {OWNER_ADD_ALLOWED:,.2f}; not RM claim {OWNER_ADD_CLAIMED:,.2f}"),
        ("Related-party rent add-back (claimed)", RELATED_PARTY_RENT_ADD_CLAIMED, "No", "Disallowed — continuing cash rent; no CA exception"),
        ("Management EBITDA (RM — not used)", MGMT_EBITDA, "No", "Keeps gain; adds full owner + rent — superseded"),
    ]
    for r, (step, amt, inc, rationale) in enumerate(ebitda_rows, 5):
        ws.cell(r, 1, step).font = body
        cell = ws.cell(r, 2, money(amt))
        cell.number_format = money_fmt
        cell.font = body
        ws.cell(r, 3, inc).font = body
        ws.cell(r, 4, rationale).font = body
        border_row(ws, r, 4)
        if inc == "No":
            ws.cell(r, 3).fill = fail_fill
        else:
            ws.cell(r, 3).fill = pass_fill

    ws["A11"] = "Adjusted EBITDA"
    ws["A11"].font = section_font
    # B11 = reported - gain + allowed owner  (typed components with formula)
    ws["B5"] = money(REPORTED_EBITDA)
    ws["B5"].number_format = money_fmt
    ws["B6"] = money(-ONE_TIME_GAIN)
    ws["B6"].number_format = money_fmt
    ws["B7"] = money(OWNER_ADD_ALLOWED)
    ws["B7"].number_format = money_fmt
    ws["B11"] = "=B5+B6+B7"
    ws["B11"].number_format = money_fmt
    ws["B11"].font = section_font
    ws["B11"].fill = pass_fill
    ws["C11"] = "Policy measure for DSCR / leverage"
    border_row(ws, 11, 4)

    ws["A13"] = "Counterfactual: if RM management EBITDA were used (prohibited)"
    ws["A13"].font = section_font
    ws["B13"] = money(MGMT_EBITDA)
    ws["B13"].number_format = money_fmt
    ws["B13"].fill = warn_fill
    autosize(ws)

    # ------------------------------------------------------------------
    # Collateral_LTV
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Collateral_LTV")
    ws["A1"] = "Collateral eligibility, advances, and LTV"
    ws["A1"].font = title_font
    ws["A2"] = f"Sources: {COLLATERAL_CSV}; eligibility/advances per {MEMO_TXT} §6–7"
    headers = [
        "Asset ID",
        "Description",
        "Category",
        "NOLV",
        "Eligible",
        "Advance %",
        "Bankable",
        "Disposition",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header(ws, 4, 8)

    haas_b = money(CNC_HAAS_NOLV * CNC_ADVANCE)
    dmg_b = money(CNC_DMG_NOLV * CNC_ADVANCE)
    mazak_b = money(CNC_MAZAK_NOLV * CNC_ADVANCE)
    fix_b = money(FIXTURES_NOLV * FIXTURE_ADVANCE)

    coll_rows = [
        ("CNC-4418-01", "Haas VF-4SS", "CNC_new", CNC_HAAS_NOLV, "Yes", CNC_ADVANCE, haas_b, "Eligible — 80% NOLV"),
        ("CNC-4418-02", "DMG MORI NLX2500", "CNC_new", CNC_DMG_NOLV, "Yes", CNC_ADVANCE, dmg_b, "Eligible — 80% NOLV"),
        ("CNC-4418-03", "Mazak Integrex i-200S", "CNC_new", CNC_MAZAK_NOLV, "Yes", CNC_ADVANCE, mazak_b, "Eligible — 80% NOLV"),
        ("FIX-4418-10", "Fixtures / tombstones", "Fixtures", FIXTURES_NOLV, "Yes", FIXTURE_ADVANCE, fix_b, "Eligible — 50% NOLV (not RM 70%)"),
        ("BLD-ROCK-01", "Plant building (related)", "RealEstate_related", BUILDING_RP_NOLV, "No", 0.0, 0.0, "Ineligible — related-party title"),
        ("MILL-1998-04", "Bridgeport Series I", "Obsolete_manual", OBSOLETE_NOLV, "No", 0.0, 0.0, "Ineligible — age cutoff"),
        ("SOFT-4418", "Install / freight / deposits", "SoftCost", SOFT_COST, "No", 0.0, 0.0, "Ineligible — soft costs"),
        ("FL-LEASE-02", "Toyota forklift", "Leased_equipment", FORKLIFT_LEASE_NOLV, "No", 0.0, 0.0, "Ineligible — lessor title"),
        ("LEASEHOLD-IMP", "Leasehold improvements", "Leasehold", 22440.55, "No", 0.0, 0.0, "Ineligible — related realty"),
    ]
    for r, row in enumerate(coll_rows, 5):
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val)
            cell.font = body
        ws.cell(r, 4).number_format = money_fmt
        ws.cell(r, 6).number_format = pct_fmt
        ws.cell(r, 7).number_format = money_fmt
        border_row(ws, r, 8)
        if row[4] == "No":
            ws.cell(r, 5).fill = fail_fill
        else:
            ws.cell(r, 5).fill = pass_fill

    # Eligible NOLV = sum of Yes rows NOLV (D5:D8)
    ws["A15"] = "Eligible NOLV"
    ws["B15"] = "=D5+D6+D7+D8"
    ws["B15"].number_format = money_fmt
    ws["C15"] = f"Must equal {ELIGIBLE_NOLV:,.2f}"
    border_row(ws, 15, 3)

    ws["A16"] = "Bankable collateral"
    ws["B16"] = "=G5+G6+G7+G8"
    ws["B16"].number_format = money_fmt
    ws["C16"] = f"Must equal {BANKABLE:,.2f}"
    border_row(ws, 16, 3)

    ws["A17"] = "75% of eligible NOLV"
    ws["B17"] = f"=ROUND(B15*{MAX_LTV},2)"
    ws["B17"].number_format = money_fmt
    border_row(ws, 17, 3)

    ws["A18"] = "Binding maximum commitment"
    ws["B18"] = "=MIN(B16,B17)"
    ws["B18"].number_format = money_fmt
    ws["B18"].fill = pass_fill
    ws["C18"] = "min(bankable, 75% eligible NOLV)"
    border_row(ws, 18, 3)

    ws["A20"] = "LTV at full request"
    ws["B20"] = f"=ROUND({REQUESTED}/B15,4)"
    ws["B20"].number_format = ratio_fmt
    ws["C20"] = "Fails max 0.75"
    ws["C20"].fill = fail_fill
    border_row(ws, 20, 3)

    ws["A21"] = "LTV at binding max"
    ws["B21"] = "=ROUND(B18/B15,4)"
    ws["B21"].number_format = ratio_fmt
    ws["C21"] = "Passes at 0.75"
    ws["C21"].fill = pass_fill
    border_row(ws, 21, 3)
    autosize(ws)

    # ------------------------------------------------------------------
    # DSCR_Leverage
    # ------------------------------------------------------------------
    ws = wb.create_sheet("DSCR_Leverage")
    ws["A1"] = "DSCR and leverage tests"
    ws["A1"].font = title_font
    ws["A2"] = f"Debt service factor {PRICING_FACTOR_PER_1000_MO} per $1,000 / month ({MEMO_TXT} §4–5)"

    headers = ["Metric", "Full request", "Sized commitment", "Threshold", "Result"]
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header(ws, 4, 5)

    # Row 5 commitment amounts
    ws["A5"] = "Commitment tested"
    ws["B5"] = money(REQUESTED)
    ws["C5"] = "=Collateral_LTV!B18"
    ws["B5"].number_format = money_fmt
    ws["C5"].number_format = money_fmt
    ws["D5"] = "n/a"
    ws["E5"] = ""
    border_row(ws, 5, 5)

    # Proposed DS = ROUND(ROUND(commit/1000*factor,2)*12,2)
    ws["A6"] = "Proposed annual debt service"
    ws["B6"] = f"=ROUND(ROUND(B5/1000*{PRICING_FACTOR_PER_1000_MO},2)*12,2)"
    ws["C6"] = f"=ROUND(ROUND(C5/1000*{PRICING_FACTOR_PER_1000_MO},2)*12,2)"
    ws["B6"].number_format = money_fmt
    ws["C6"].number_format = money_fmt
    ws["D6"] = f"Factor {PRICING_FACTOR_PER_1000_MO}"
    border_row(ws, 6, 5)

    ws["A7"] = "Existing bank annual DS"
    ws["B7"] = money(EXISTING_ANNUAL_DS)
    ws["C7"] = "=B7"
    ws["B7"].number_format = money_fmt
    ws["C7"].number_format = money_fmt
    border_row(ws, 7, 5)

    ws["A8"] = "Total annual debt service"
    ws["B8"] = "=ROUND(B6+B7,2)"
    ws["C8"] = "=ROUND(C6+C7,2)"
    ws["B8"].number_format = money_fmt
    ws["C8"].number_format = money_fmt
    border_row(ws, 8, 5)

    ws["A9"] = "Adjusted EBITDA"
    ws["B9"] = "=EBITDA_Build!B11"
    ws["C9"] = "=B9"
    ws["B9"].number_format = money_fmt
    ws["C9"].number_format = money_fmt
    border_row(ws, 9, 5)

    ws["A10"] = "DSCR"
    ws["B10"] = "=ROUND(B9/B8,4)"
    ws["C10"] = "=ROUND(C9/C8,4)"
    ws["B10"].number_format = ratio_fmt
    ws["C10"].number_format = ratio_fmt
    ws["D10"] = MIN_DSCR
    ws["E10"] = "Full ask fails 1.25; sized passes"
    ws["B10"].fill = fail_fill
    ws["C10"].fill = pass_fill
    border_row(ws, 10, 5)

    ws["A11"] = "Total debt (UPB + commitment)"
    ws["B11"] = f"=ROUND({EXISTING_DEBT_BAL}+B5,2)"
    ws["C11"] = f"=ROUND({EXISTING_DEBT_BAL}+C5,2)"
    ws["B11"].number_format = money_fmt
    ws["C11"].number_format = money_fmt
    border_row(ws, 11, 5)

    ws["A12"] = "Leverage (Total debt / Adj EBITDA)"
    ws["B12"] = "=ROUND(B11/B9,4)"
    ws["C12"] = "=ROUND(C11/C9,4)"
    ws["B12"].number_format = ratio_fmt
    ws["C12"].number_format = ratio_fmt
    ws["D12"] = MAX_LEVERAGE
    ws["E12"] = "Sized passes 3.50x"
    ws["C12"].fill = pass_fill
    border_row(ws, 12, 5)

    ws["A14"] = "Path trigger notes"
    ws["A14"].font = section_font
    ws["A15"] = (
        f"Full-ask DSCR {DSCR_FULL} is below {MIN_DSCR} but at/above {COND_DSCR_FLOOR}; "
        f"full-ask LTV {LTV_FULL} exceeds {MAX_LTV}. After sizing to binding max, "
        f"DSCR {DSCR_SIZED} and LTV {LTV_SIZED} clear Approve thresholds → Conditional path."
    )
    ws["A15"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A15:E17")
    autosize(ws)

    # ------------------------------------------------------------------
    # Exceptions
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Exceptions")
    ws["A1"] = "Exception and conflict log"
    ws["A1"].font = title_font
    headers = ["Item", "Source conflict", "Governing treatment", "Impact"]
    for i, h in enumerate(headers, 1):
        ws.cell(3, i, h)
    style_header(ws, 3, 4)
    exceptions = [
        (
            "RM Approve as requested",
            f"{FIN_XLSX} RM_Notes / Application",
            f"{MEMO_TXT} §1 — memo supersedes; path not Approve at full ask",
            "Blocks full-ask Approve",
        ),
        (
            "Management EBITDA",
            f"{FIN_XLSX} Income / RM_Notes",
            f"{MEMO_TXT} §3 — strip gain; cap owner add-back; disallow rent add-back",
            f"Adj EBITDA {ADJ_EBITDA:,.2f} vs MGMT {MGMT_EBITDA:,.2f}",
        ),
        (
            "Related-party building pledged",
            f"{COLLATERAL_CSV} BLD-ROCK-01",
            f"{MEMO_TXT} §6 — ineligible collateral",
            "Removed from eligible NOLV / bankable",
        ),
        (
            "Soft costs 100% advance claim",
            f"{COLLATERAL_CSV} SOFT-4418",
            f"{MEMO_TXT} §6 — soft costs 0% advance",
            "Excluded from bankable",
        ),
        (
            "Fixture RM advance 70%",
            f"{COLLATERAL_CSV} FIX-4418-10",
            f"{MEMO_TXT} §6 — fixtures 50%",
            "Bankable uses 50%",
        ),
        (
            "Leased forklift pledged",
            f"{COLLATERAL_CSV} FL-LEASE-02",
            f"{MEMO_TXT} §6 — lessor title ineligible",
            "Excluded",
        ),
        (
            "No personal guaranty offered",
            f"{FIN_XLSX} Application",
            f"{MEMO_TXT} §8 — Conditional requires unlimited PG",
            "Condition on recommendation",
        ),
        (
            "84-month tenor requested",
            f"{FIN_XLSX} Application",
            f"{MEMO_TXT} §8 — Conditional shortens to {CONDITIONS_TENOR_MO} months",
            "Condition on recommendation",
        ),
    ]
    for r, row in enumerate(exceptions, 4):
        for c, val in enumerate(row, 1):
            ws.cell(r, c, val).font = body
        border_row(ws, r, 4)
    autosize(ws)

    # ------------------------------------------------------------------
    # Credit_Decision
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Credit_Decision")
    ws["A1"] = f"{ENTITY} — credit path decision"
    ws["A1"].font = title_font
    ws["A2"] = f"Committee {COMMITTEE_DATE} | Application {APPLICATION_ID}"

    ws["A4"] = "Path"
    ws["B4"] = PATH
    ws["B4"].fill = warn_fill
    ws["B4"].font = section_font
    border_row(ws, 4, 2)

    ws["A5"] = "Recommended commitment (USD)"
    ws["B5"] = "=Collateral_LTV!B18"
    ws["B5"].number_format = money_fmt
    ws["B5"].fill = pass_fill
    border_row(ws, 5, 2)

    ws["A6"] = "Requested commitment (not approved)"
    ws["B6"] = money(REQUESTED)
    ws["B6"].number_format = money_fmt
    ws["B6"].fill = fail_fill
    border_row(ws, 6, 2)

    ws["A7"] = "DSCR at recommended commitment"
    ws["B7"] = "=DSCR_Leverage!C10"
    ws["B7"].number_format = ratio_fmt
    border_row(ws, 7, 2)

    ws["A8"] = "LTV at recommended commitment"
    ws["B8"] = "=Collateral_LTV!B21"
    ws["B8"].number_format = ratio_fmt
    border_row(ws, 8, 2)

    ws["A9"] = "Leverage at recommended commitment"
    ws["B9"] = "=DSCR_Leverage!C12"
    ws["B9"].number_format = ratio_fmt
    border_row(ws, 9, 2)

    ws["A11"] = "Conditions (Conditional path)"
    ws["A11"].font = section_font
    conditions = [
        f"1. Cap commitment at binding maximum (Collateral_LTV binding max; {COMMITMENT:,.2f}).",
        f"2. Tenor {CONDITIONS_TENOR_MO} months (not {REQUESTED_TENOR_MO}).",
        "3. Unlimited personal guaranty of the principal owner.",
        "4. Quarterly DSCR covenant at 1.25x on Adjusted EBITDA per memo §3.",
        "5. Pledge limited to eligible CNC and fixtures; exclude related-party realty, leases, soft costs, obsolete mill.",
        "6. Present to committee (Conditional sizing path) even though sized amount is under officer limit.",
    ]
    for i, ctext in enumerate(conditions, 12):
        ws.cell(i, 1, ctext).font = body
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)

    ws["A19"] = "Explicit do-nots"
    ws["A19"].font = section_font
    donots = [
        "Do not approve the full requested commitment.",
        "Do not treat Whitaker Realty LLC building or leasehold improvements as eligible collateral.",
        "Do not measure DSCR on management EBITDA that retains the parcel gain or related-party rent add-back.",
        "Do not Decline solely because the request exceeds bankable collateral — size to binding max under Conditional.",
    ]
    for i, t in enumerate(donots, 20):
        ws.cell(i, 1, t).font = body
        ws.cell(i, 1).fill = fail_fill
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
    autosize(ws)

    # ------------------------------------------------------------------
    # Notes
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Notes")
    ws["A1"] = "Committee notes and source citations"
    ws["A1"].font = title_font
    notes = [
        f"Governing hierarchy: {MEMO_TXT} §1 controls over RM Approve-as-requested language on {FIN_XLSX} RM_Notes.",
        f"Adjusted EBITDA {ADJ_EBITDA:,.2f} = reported {REPORTED_EBITDA:,.2f} − gain {ONE_TIME_GAIN:,.2f} + allowed owner {OWNER_ADD_ALLOWED:,.2f} ({FIN_XLSX}; {MEMO_TXT} §3).",
        f"Eligible NOLV {ELIGIBLE_NOLV:,.2f}; bankable {BANKABLE:,.2f}; 75% cap {MAX_BY_LTV:,.2f}; binding commitment {COMMITMENT:,.2f} ({COLLATERAL_CSV}; {MEMO_TXT} §6–7).",
        f"Full-ask DSCR {DSCR_FULL} and LTV {LTV_FULL} fail Approve tests; sized DSCR {DSCR_SIZED}, LTV {LTV_SIZED}, leverage {LEVERAGE_SIZED} pass ({MEMO_TXT} §5, §8).",
        f"Path selected: {PATH} with commitment {COMMITMENT:,.2f} and {CONDITIONS_TENOR_MO}-month tenor plus PG and DSCR covenant.",
        "Citations above appear earlier on this Notes tab than the committee path recommendation that follows.",
        f"Committee path recommendation: {PATH} at ${COMMITMENT:,.2f} with conditions listed on Credit_Decision — not Approve at ${REQUESTED:,.2f}, and not Decline.",
    ]
    for i, t in enumerate(notes, 3):
        ws.cell(i, 1, t).font = body
        ws.cell(i, 1).alignment = Alignment(wrap_text=True)
        ws.row_dimensions[i].height = 36
    ws.column_dimensions["A"].width = 110

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    sanitize_xlsx(OUT)

    cache = {
        "EBITDA_Build": {
            "B11": ADJ_EBITDA,
        },
        "Collateral_LTV": {
            "B15": ELIGIBLE_NOLV,
            "B16": BANKABLE,
            "B17": MAX_BY_LTV,
            "B18": COMMITMENT,
            "B20": LTV_FULL,
            "B21": LTV_SIZED,
        },
        "DSCR_Leverage": {
            "C5": COMMITMENT,
            "B6": PROPOSED_DS_FULL,
            "C6": PROPOSED_DS_SIZED,
            "B8": TOTAL_DS_FULL,
            "C8": TOTAL_DS_SIZED,
            "B9": ADJ_EBITDA,
            "C9": ADJ_EBITDA,
            "B10": DSCR_FULL,
            "C10": DSCR_SIZED,
            "B11": money(EXISTING_DEBT_BAL + REQUESTED),
            "C11": TOTAL_DEBT_SIZED,
            "B12": round(money(EXISTING_DEBT_BAL + REQUESTED) / ADJ_EBITDA, 4),
            "C12": LEVERAGE_SIZED,
        },
        "Credit_Decision": {
            "B5": COMMITMENT,
            "B7": DSCR_SIZED,
            "B8": LTV_SIZED,
            "B9": LEVERAGE_SIZED,
        },
    }
    cache_xlsx_formula_values(OUT, cache)
    print(f"Wrote {OUT}")
    print(f"Path={PATH} Commitment={COMMITMENT} AdjEBITDA={ADJ_EBITDA} DSCR_sized={DSCR_SIZED}")


if __name__ == "__main__":
    main()
