#!/usr/bin/env python3
"""Generate inputs for Ashcroft Foods single-claim coverage determination."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ashcroft_constants import (
    BI_CLAIMED,
    BI_WAITING_HOURS,
    CLAIM_FILE,
    CLAIM_ID,
    COMMITTEE_DATE,
    DAYS_SINCE_SERVICE,
    ENTITY,
    EQUIPMENT_COMPONENT_TOTAL,
    EQUIPMENT_FOOTING_VARIANCE,
    EQUIPMENT_LINES,
    EQUIPMENT_PRIOR_MISSTATED_SUBTOTAL,
    EQUIPMENT_REPAIR,
    EVAL_DATE,
    GRACE_DAYS,
    LAST_SERVICE,
    LOCATION,
    LOSS_DATE,
    LOTS,
    MEMO_FILE,
    OEM_SERVICE_DAYS,
    OUTAGE_HOURS,
    POLICY,
    POLICY_FILE,
    REPORT_DATE,
    SALVAGE_LOT,
    SALVAGE_RATE,
    SLUG,
    SPOILAGE_DEDUCTIBLE,
    SUPERSEDED_DEDUCTIBLE,
    VENDOR_CANCEL_DATE,
)
from sanitize_office import sanitize_xlsx

TASK_INPUTS = Path(__file__).resolve().parent.parent / "tasks" / SLUG / "inputs"
header_font = Font(bold=True, name="Calibri", size=11)
body_font = Font(name="Calibri", size=10)
thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def money(x: float) -> float:
    return round(float(x), 2)


def style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(1, c)
        cell.font = header_font
        cell.border = thin


def autosize(ws, min_w: int = 10, max_w: int = 28) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min(max(max(len(str(c.value or "")) for c in col) + 2, min_w), max_w)
        ws.column_dimensions[letter].width = width


def write_claim_file(path: Path) -> None:
    wb = Workbook()

    # --- Loss Notice ---
    ws = wb.active
    ws.title = "Loss Notice"
    rows = [
        ("Field", "Value"),
        ("Claim Number", CLAIM_ID),
        ("Policy Number", POLICY),
        ("Named Insured", f"{ENTITY} LLC"),
        ("Loss Location", LOCATION),
        ("Date of Loss", LOSS_DATE),
        ("Date Reported", REPORT_DATE),
        ("Cause Narrative", "Ammonia compressor #3 hard-faulted overnight; Zone B product temp rose above safe hold"),
        ("Reported Outage Hours", OUTAGE_HOURS),
        ("Temporary Cooling Restored", "2025-11-19 06:40 local"),
        ("Claimant Contact", "Marla Quince, Plant Controller"),
        ("Broker", "Hollis & Pike Retail Risk"),
        ("Carrier Desk", "Midwest Property Claims — Columbus"),
        ("As-Of Evaluation", EVAL_DATE),
        ("Committee Target", COMMITTEE_DATE),
        ("Claimant Spoilage Demand (sell)", money(sum(money(r[2] * r[4]) for r in LOTS if r[6] != "VOID"))),
        ("Claimant Equipment Demand", EQUIPMENT_REPAIR),
        ("Claimant BI Demand", BI_CLAIMED),
        ("Prior Case Reserve (placeholder)", 0),
        ("FNOL Source", "Agent portal + phone"),
        ("Occupancy", "Frozen food distribution warehouse"),
        ("Construction", "Tilt-up concrete, insulated panels"),
        ("Protection Class", "4"),
        ("Sprinklered", "Yes — dry system Zone B"),
    ]
    for r in rows:
        ws.append(r)
    style_header(ws, 2)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=2):
        for c in row:
            c.font = body_font
            c.border = thin
    autosize(ws, 14, 55)

    # --- Inventory ---
    ws = wb.create_sheet("Inventory")
    headers = [
        "lot_id",
        "description",
        "zone",
        "qty_cases",
        "unit_cost",
        "unit_selling_price",
        "line_cost",
        "line_selling",
        "status",
        "temp_excursion_flag",
        "notes",
    ]
    ws.append(headers)
    style_header(ws, len(headers))
    for lot in LOTS:
        lot_id, desc, qty, ucost, usell, lcost, status, _cov = lot
        zone = "B" if not lot_id.startswith("PKG") else "Staging"
        flag = "Y" if status != "VOID" and lot_id != "LOT-B-218-DUP" else ("N" if status == "VOID" else "Y")
        note = ""
        if lot_id == "PKG-991":
            note = "Packaging stored adjacent to Zone B dock"
        elif lot_id == SALVAGE_LOT:
            note = "QA holds ~30% recoverable to secondary channel if released"
        elif status == "Duplicate":
            note = "Rekeyed 2025-11-21; original LOT-B-218 remains"
        elif status == "VOID":
            note = "Withdrawn — never loaded to Zone B on loss night"
        ws.append(
            [
                lot_id,
                desc,
                zone,
                qty,
                money(ucost),
                money(usell),
                money(lcost),
                money(qty * usell),
                status,
                flag,
                note,
            ]
        )
    # Extra depth: unaffected Zone A reference lots (not part of this claim demand)
    zone_a = [
        ("LOT-A-112", "Butter solids 25lb", "A", 90, 88.40, 124.00, "Open", "N", "Unaffected zone — reference only"),
        ("LOT-A-118", "Cheese shreds 20lb", "A", 140, 96.15, 138.50, "Open", "N", "Unaffected zone — reference only"),
        ("LOT-A-124", "Pizza toppings mix 30lb", "A", 75, 102.88, 149.20, "Open", "N", "Unaffected zone — reference only"),
        ("LOT-A-131", "Breaded chicken 20lb", "A", 210, 91.33, 132.75, "Open", "N", "Unaffected zone — reference only"),
        ("LOT-A-140", "Breakfast sausage 10lb", "A", 188, 79.62, 115.40, "Open", "N", "Unaffected zone — reference only"),
        ("LOT-A-155", "Onion rings 15lb", "A", 96, 71.25, 104.80, "Open", "N", "Unaffected zone — reference only"),
        ("LOT-A-162", "Potato wedges 20lb", "A", 130, 64.18, 93.50, "Open", "N", "Unaffected zone — reference only"),
        ("LOT-A-170", "Garlic bread sticks", "A", 155, 58.90, 86.25, "Open", "N", "Unaffected zone — reference only"),
    ]
    for r in zone_a:
        lid, desc, zone, qty, uc, us, st, flag, note = r
        ws.append([lid, desc, zone, qty, money(uc), money(us), money(qty * uc), money(qty * us), st, flag, note])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for c in row:
            c.font = body_font
            c.border = thin
    autosize(ws)

    # --- Equipment ---
    ws = wb.create_sheet("Equipment")
    ws.append(["line_id", "description", "vendor", "quote_date", "amount", "status", "notes"])
    style_header(ws, 7)
    equip_rows = [
        ("EQ-01", EQUIPMENT_LINES[0][1], "North River Refrigeration", "2025-11-22", EQUIPMENT_LINES[0][2], "Quoted", "OEM core + labor"),
        ("EQ-02", EQUIPMENT_LINES[1][1], "North River Refrigeration", "2025-11-22", EQUIPMENT_LINES[1][2], "Quoted", "Includes purge"),
        ("EQ-03", EQUIPMENT_LINES[2][1], "North River Refrigeration", "2025-11-22", EQUIPMENT_LINES[2][2], "Quoted", "Zone B rack"),
        ("EQ-04", EQUIPMENT_LINES[3][1], "North River Refrigeration", "2025-11-22", EQUIPMENT_LINES[3][2], "Quoted", "Weekend callout"),
        ("EQ-05", EQUIPMENT_LINES[4][1], "Metro Calibration LLC", "2025-11-24", EQUIPMENT_LINES[4][2], "Quoted", "Post-loss verification"),
        (
            "EQ-SUB",
            "Equipment subtotal (insured demand)",
            "—",
            "2025-11-24",
            EQUIPMENT_REPAIR,
            "Demand",
            f"Sum of EQ-01..EQ-05 (footed {EQUIPMENT_COMPONENT_TOTAL:,.2f}; "
            f"prior worksheet {EQUIPMENT_PRIOR_MISSTATED_SUBTOTAL:,.2f} overstated by "
            f"{EQUIPMENT_FOOTING_VARIANCE:,.2f} — see Adjuster Note N-09)",
        ),
        ("EQ-ALT", "Used compressor swap option", "Great Lakes Cold Parts", "2025-11-25", 47880.00, "Alternate", "Not selected by insured"),
        ("EQ-HIST", "Prior 2024 compressor #2 service", "North River Refrigeration", "2024-09-18", 4120.75, "History", "Unrelated prior work"),
        ("EQ-HIST2", "2025-Q2 oil analysis all compressors", "North River Refrigeration", "2025-05-02", 980.00, "History", "Showed #3 trending metal"),
    ]
    for r in equip_rows:
        ws.append(list(r[:4]) + [money(r[4]), r[5], r[6]])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=7):
        for c in row:
            c.font = body_font
            c.border = thin
    autosize(ws)

    # --- BI Claim ---
    ws = wb.create_sheet("BI Claim")
    ws.append(["metric", "value", "unit", "notes"])
    style_header(ws, 4)
    bi_rows = [
        ("Daily contribution estimate", 4802.86, "USD", "Controller worksheet — not audited"),
        ("Days claimed", 14, "days", "Insured counted calendar days until full rack stable"),
        ("Hours of elevated product temp", OUTAGE_HOURS, "hours", "From logger download"),
        ("Policy BI waiting period (from schedule)", BI_WAITING_HOURS, "hours", "See policy schedule"),
        ("Gross BI demand", BI_CLAIMED, "USD", "Daily × 14"),
        ("Extra expense generators", 3180.00, "USD", "Portable chiller rental — also claimed under BI"),
        ("Payroll continuance claimed", 0.00, "USD", "Not requested"),
        ("Sales rebound week of 11/24", 1.04, "index", "Vs prior 4-week avg"),
        ("Alternate supplier fill-in margin loss", 2210.40, "USD", "Embedded in daily estimate"),
        ("Claimant BI narrative", "Shutdown", "text", "Insured labels event as multi-day interruption"),
    ]
    for r in bi_rows:
        val = money(r[1]) if isinstance(r[1], float) else r[1]
        ws.append([r[0], val, r[2], r[3]])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=4):
        for c in row:
            c.font = body_font
            c.border = thin
            c.alignment = Alignment(wrap_text=True)
    autosize(ws, 12, 40)

    # --- Adjuster Notes ---
    ws = wb.create_sheet("Adjuster Notes")
    ws.append(["note_id", "author", "note_date", "as_of_status", "body"])
    style_header(ws, 5)
    notes = [
        (
            "N-01",
            "T. Belmont (field)",
            "2025-11-20",
            "Superseded",
            "Initial scene: compressor #3 offline, Zone B warm. Insured wants selling-price inventory "
            "and full equipment. Recommend provisional full sell-price case pending docs.",
        ),
        (
            "N-02",
            "T. Belmont (field)",
            "2025-11-21",
            "Active",
            f"OEM service interval for compressor #3: {OEM_SERVICE_DAYS} days. "
            f"Maintenance log shows last completed OEM oil/filter service on {LAST_SERVICE}. "
            f"Vendor ticket {VENDOR_CANCEL_DATE} cancelled by North River for parts delay. "
            "Oil analysis from May flagged metal in #3. Logger confirms ~14 hours elevated product temp.",
        ),
        (
            "N-03",
            "S. Okonkwo (desk)",
            "2025-11-26",
            "Active",
            "Policy review started. Spoilage endorsement present. Mechanical Breakdown exclusion "
            "on equipment. Do not treat Belmont N-01 sell-price recommendation as governing.",
        ),
        (
            "N-04",
            "S. Okonkwo (desk)",
            "2025-12-02",
            "Active",
            "Inventory rekey created LOT-B-218-DUP. VOID lot LOT-B-099 never in Zone B. "
            "Packaging PKG-991 sitting in staging — confirm stock definition before paying.",
        ),
        (
            "N-05",
            "QA Ashcroft",
            "2025-11-23",
            "Active",
            f"{SALVAGE_LOT} turkey may move to secondary processor at roughly {int(SALVAGE_RATE*100)}% "
            "of invoice if released under hold protocol.",
        ),
        (
            "N-06",
            "Broker email excerpt",
            "2025-11-25",
            "Advocacy",
            "Broker argues superseded $10k spoilage deductible should apply and BI waiting period "
            "should be waived because customers were reallocated for two weeks.",
        ),
        (
            "N-07",
            "Carrier counsel brief",
            "2025-12-03",
            "Active",
            "Coverage memo forthcoming from SIU/coverage unit. Committee packet due "
            f"{COMMITTEE_DATE}. Authority for payment proposals above 250000 requires manager.",
        ),
        (
            "N-08",
            "Plant walkthrough",
            "2025-11-20",
            "Active",
            "Zone A unaffected. Dry sprinkler heads show no discharge. No ammonia exposure to product "
            "confirmed by industrial hygienist swab — loss is temperature only.",
        ),
        (
            "N-09",
            "S. Okonkwo (desk)",
            "2025-12-04",
            "Active",
            f"Equipment demand footing check: insured worksheet EQ-SUB showed "
            f"{EQUIPMENT_PRIOR_MISSTATED_SUBTOTAL:,.2f}, but EQ-01 through EQ-05 quote lines sum to "
            f"{EQUIPMENT_COMPONENT_TOTAL:,.2f} (variance {EQUIPMENT_FOOTING_VARIANCE:,.2f}). "
            f"Corrected EQ-SUB on the Equipment tab to the component sum. Variance is a worksheet "
            "addition error, not an additional quoted repair line. Coverage analysis still pending.",
        ),
    ]
    for r in notes:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=5):
        for c in row:
            c.font = body_font
            c.border = thin
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["E"].width = 70
    autosize(ws, 10, 22)

    # --- Payments / Misc ---
    ws = wb.create_sheet("Payments")
    ws.append(["txn_id", "date", "type", "amount", "status", "memo"])
    style_header(ws, 6)
    pay_rows = [
        ("P-000", "2025-11-19", "Reserve set", 0.00, "System", "FNOL auto-zero pending coverage"),
        ("P-001", "2025-11-22", "Expense", 1840.00, "Paid", "Independent adjuster day rate"),
        ("P-002", "2025-11-24", "Expense", 620.50, "Paid", "Logger download / lab courier"),
        ("P-003", "2025-11-26", "Expense", 275.00, "Paid", "Hygienist swab panel"),
        ("P-004", "2025-12-01", "Void check", 0.00, "VOID", "Miskeyed advance — never issued"),
        ("P-005", "2025-12-04", "Expense", 410.00, "Pending", "Coverage counsel consult hour block"),
    ]
    for r in pay_rows:
        ws.append([r[0], r[1], r[2], money(r[3]), r[4], r[5]])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=6):
        for c in row:
            c.font = body_font
            c.border = thin
    autosize(ws)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    sanitize_xlsx(path)


def write_policy_csv(path: Path) -> None:
    headers = [
        "row_id",
        "policy_number",
        "coverage_part",
        "description",
        "limit_amount",
        "deductible",
        "waiting_period_hours",
        "form_or_endorsement",
        "status",
        "effective_date",
        "notes",
    ]
    rows = [
        ("PS-01", POLICY, "Building", "Building — Dayton warehouse", 8400000, 25000, 0, "CP 00 10", "Active", "2025-03-01", "Replacement cost"),
        ("PS-02", POLICY, "BPP", "Business personal property", 3250000, 10000, 0, "CP 00 10", "Active", "2025-03-01", "Includes stock"),
        ("PS-03", POLICY, "Spoilage", "Spoilage coverage — refrigeration breakdown", 750000, SPOILAGE_DEDUCTIBLE, 0, "CP 04 40", "Active", "2025-03-01", "On-premises mechanical refrigeration failure"),
        ("PS-04", POLICY, "Spoilage", "Spoilage deductible — SUPERSEDED mid-term row", 750000, SUPERSEDED_DEDUCTIBLE, 0, "CP 04 40", "Superseded", "2025-03-01", "Replaced by endorsement revision 2025-06-15; status Superseded"),
        ("PS-05", POLICY, "Exclusion", "Mechanical Breakdown — equipment", 0, 0, 0, "Causes of Loss Special", "Active", "2025-03-01", "Bars direct equipment repair/replacement"),
        ("PS-06", POLICY, "Business Income", "Business Income and Extra Expense", 500000, 0, BI_WAITING_HOURS, "CP 00 30", "Active", "2025-03-01", "72-hour waiting period"),
        ("PS-07", POLICY, "Ordinance", "Ordinance or Law Coverage A", 500000, 25000, 0, "CP 04 05", "Active", "2025-03-01", "Not implicated"),
        ("PS-08", POLICY, "Debris", "Debris Removal", 250000, 10000, 0, "CP 00 10", "Active", "2025-03-01", "Not claimed"),
        ("PS-09", POLICY, "Spoilage", "Spoilage valuation clause", 750000, SPOILAGE_DEDUCTIBLE, 0, "CP 04 40 §Valuation", "Active", "2025-06-15", "Stock at invoice / replacement cost — not selling price"),
        ("PS-10", POLICY, "Condition", "Protective safeguards — refrigeration maintenance", 0, 0, 0, "IL 04 15 / schedule", "Active", "2025-03-01", "OEM service intervals required"),
        ("PS-11", POLICY, "BPP", "Peak season BPP increase", 400000, 10000, 0, "CP 12 30", "Active", "2025-10-01", "Oct–Dec; separate from spoilage sublimit"),
        ("PS-12", POLICY, "Liability", "Premises liability (reference only)", 2000000, 0, 0, "CG 00 01", "Active", "2025-03-01", "Not this claim"),
        ("PS-13", POLICY, "Auto", "Hired/non-owned auto (reference)", 1000000, 0, 0, "CA 00 01", "Active", "2025-03-01", "Not this claim"),
        ("PS-14", POLICY, "Spoilage", "Power / utility service interruption spoilage", 100000, 5000, 0, "CP 04 40 opt", "Not purchased", "2025-03-01", "Off-premises utility — N/A; loss is on-prem compressor"),
        ("PS-15", POLICY, "BI", "Extended period of indemnity", 60, 0, BI_WAITING_HOURS, "CP 15 05", "Active", "2025-03-01", "Days after operations resume; waiting period still applies first"),
        ("PS-16", POLICY, "Deductible", "Wind/hail deductible separate", 0, 50000, 0, "Schedule", "Active", "2025-03-01", "Not implicated"),
        ("PS-17", POLICY, "Spoilage", "Spoilage sublimit aggregate", 750000, SPOILAGE_DEDUCTIBLE, 0, "CP 04 40", "Active", "2025-06-15", "Per occurrence within BPP"),
        ("PS-18", POLICY, "Condition", "Cooperation / proof of loss", 0, 0, 0, "Common Policy Conditions", "Active", "2025-03-01", "Standard"),
        ("PS-19", POLICY, "Endorsement", "Ammonia contamination exclusion (product)", 0, 0, 0, "Manuscript", "Active", "2025-03-01", "Not triggered — no ammonia product contact"),
        ("PS-20", POLICY, "Endorsement", "Delayed maintenance / protective safeguards", 0, 0, 0, "IL 04 15 schedule notes", "Active", "2025-03-01", "See coverage investigation memo for committee practice"),
        ("PS-21", POLICY, "BPP", "Stock declaration — frozen foods", 2100000, 10000, 0, "Statement of Values", "Active", "2025-03-01", "Subset of BPP"),
        ("PS-22", POLICY, "Building", "Sign coverage", 25000, 1000, 0, "CP 14 40", "Active", "2025-03-01", "Not claimed"),
        ("PS-23", POLICY, "BI", "Ordinary payroll limitation 60 days", 0, 0, BI_WAITING_HOURS, "CP 15 10", "Active", "2025-03-01", "Payroll not claimed"),
        ("PS-24", POLICY, "Spoilage", "Spoilage — selling price NOT elected", 0, 0, 0, "CP 04 40 options", "Not elected", "2025-03-01", "Selling price valuation rider not on policy"),
        ("PS-25", POLICY, "Exclusion", "Wear and tear / inherent vice", 0, 0, 0, "Causes of Loss Special", "Active", "2025-03-01", "Background; refrigeration breakdown addressed via spoilage form"),
        ("PS-26", POLICY, "Contact", "Carrier claims office", 0, 0, 0, "Declarations", "Active", "2025-03-01", "Columbus property desk"),
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for r in rows:
            w.writerow(r)


def write_memo(path: Path) -> None:
    # Governing rules and hierarchy only — no finished coverage dispositions,
    # no named exclusion list, no complete reserve recipe, no final dollars.
    text = f"""ASHCROFT FOODS LLC — COVERAGE INVESTIGATION MEMO
Claim: {CLAIM_ID} | Policy: {POLICY} | Loss date: {LOSS_DATE}
Prepared for file-review committee {COMMITTEE_DATE} | Evaluation as of {EVAL_DATE}
Coverage unit: Midwest Property Claims — Columbus

Purpose and source hierarchy
This memo establishes the governing rules and evidence hierarchy for CLM-CP-2025-8841.
It does not replace independent application of those rules to the claim file and policy
schedule. Use all three packet sources. Where sources conflict on valuation basis,
deductible selection, maintenance/protective-safeguards treatment, salvage treatment,
stock eligibility, or business-income waiting-period application, control in this order:
(1) this memo’s rules; (2) Active rows on ashcroft-foods_policy_schedule.csv;
(3) contemporaneous claim-file evidence (Inventory, Equipment, BI Claim, Payments,
Active adjuster notes); (4) superseded schedule rows and superseded field notes
(including N-01) are non-governing for committee purposes. Broker advocacy favoring
superseded schedule amounts or waiting-period waivers is advocacy only and does not
control.

Loss facts (investigation summary — not coverage conclusions)
Ammonia refrigeration compressor #3 serving Zone B failed overnight on {LOSS_DATE}.
Product temperature loggers show approximately {OUTAGE_HOURS} hours of elevated
temperatures before temporary cooling restored the room on the morning of {REPORT_DATE}.
Industrial hygiene swabs found no ammonia contact with product. Zone A remained within
specification. The insured demands spoilage of frozen stock at selling price, repair or
replacement of compressor #3 and related equipment, and business income for a multi-day
customer reallocation period. Logger hours, inventory statuses, equipment quotes, BI
worksheet metrics, maintenance dates, and payment/expense rows live in the claim file;
coverage parts, deductibles, waiting periods, and endorsement statuses live on the
policy schedule.

Governing forms and endorsements
Confirm Active schedule rows for: Spoilage Coverage endorsement CP 04 40 (on-premises
mechanical refrigeration failure triggering temperature change to stock); Causes of Loss
Special Form Mechanical Breakdown exclusion applicable to direct equipment repair or
replacement; Business Income and Extra Expense CP 00 30 with its scheduled waiting
period; CP 04 40 valuation clause; protective-safeguards / refrigeration maintenance
conditions; and whether a selling-price valuation option was elected. Off-premises
utility spoilage options that were not purchased do not respond to an on-premises
compressor failure.

Spoilage endorsement — stock trigger (rule)
CP 04 40 responds to spoilage of “stock” caused by a change in temperature resulting
from mechanical breakdown of refrigeration equipment on the described premises, subject
to the spoilage sublimit, the governing deductible, the valuation clause, and protective-
safeguards conditions. The Mechanical Breakdown exclusion bars direct repair or
replacement of the failed equipment itself. That equipment exclusion does not, by itself,
nullify a separate Spoilage endorsement response for stock when the endorsement’s
triggering language is met. Whether stock, equipment, and business income each respond
on this file must be determined by applying these rules to the claim-file facts and
Active schedule terms — this memo does not pre-state those dispositions.

Equipment coverage conditions (rule)
Equipment rebuild, ammonia recharge, controls, crane premiums, and related mechanical
repair quotes are evaluated under the Mechanical Breakdown exclusion and against whether
any equipment floater or boiler & machinery coverage appears as Active on the schedule
for compressor #3. Absence of such a floater on the schedule is material. Do not infer
equipment indemnity solely because stock may be eligible under a separate endorsement.

Business income waiting period (rule)
The Business Income form’s waiting period is the Active schedule waiting_period_hours
value for CP 00 30 (and related BI endorsements that retain that waiting period). The
waiting period is measured against the duration of the covered interruption supported by
logger / operations evidence, not against calendar days of customer reallocation after
cooling is restored. Customer reallocation over subsequent calendar days does not restart
or waive the waiting period. Extended period of indemnity, if present, still requires the
waiting period to be satisfied first. Compare claim-file outage evidence to the schedule
waiting period before posting any BI indemnity reserve.

Definition of qualifying stock (rule)
For this desk under CP 04 40, “stock” means perishable frozen food product held for sale
or distribution. Packaging materials, corrugated cartons, dunnage, and staging supplies
are not stock even when stored adjacent to a freezer zone. Eligible stock lines are Zone B
inventory rows with status Open or Open-Salvage that meet this stock definition. Lines
with status VOID are ineligible. Duplicate rekey lines are not separately eligible; retain
only the original lot. Zone A lots and other unaffected-zone reference rows are outside
this loss population. Apply these filters to the Inventory tab; do not copy claimant
selling-price demand totals as the covered population.

Valuation basis (rule)
The CP 04 40 valuation clause and the schedule election status for selling-price valuation
require invoice / replacement cost for stock when selling price was not elected. Field notes
recommending provisional selling-price payment are non-governing when they conflict with
this rule. Use line_cost (invoice) for eligible stock, not unit_selling_price aggregates,
unless an Active schedule election expressly authorizes selling-price valuation.

Salvage treatment (rule)
Lots marked Open-Salvage remain partially recoverable. Credit salvage against that lot’s
invoice cost at the recoverable percentage documented in the claim file (QA / inventory
notes). Salvage credits apply only to covered stock indemnity — do not net salvage against
equipment or business-income demands. Apply salvage to the covered stock loss before
maintenance coinsurance and before the spoilage deductible.

Delayed maintenance / protective safeguards (rule)
OEM service intervals for compressor oil and filter service are required under the
protective-safeguards schedule. The claim file records the last completed service date,
OEM interval days, and any vendor cancellation near the loss. Desk practice for committee
files: when service is overdue beyond a {GRACE_DAYS}-day scheduling grace measured from
the OEM interval, and a vendor cancellation is documented within 14 days of the loss,
stock coverage under CP 04 40 remains available but a 25% coinsurance reduction applies
(pay 75% of the otherwise covered stock loss after salvage). Do not void the entire
spoilage claim solely on the maintenance condition when that documented vendor-cancellation
safe harbor is met. Compute overdue days and confirm cancellation timing from the claim
file before applying or declining this adjustment.

Deductible hierarchy (rule)
The governing spoilage deductible is the deductible on the Active CP 04 40 spoilage
endorsement / sublimit rows on the policy schedule. Schedule rows with status Superseded
(including mid-term deductible revisions later replaced) are non-governing. Building,
BPP, wind/hail, or other unrelated deductibles do not substitute for the spoilage
deductible on a covered spoilage indemnity. Apply the governing spoilage deductible to
covered spoilage after valuation, salvage credit, and any applicable maintenance
coinsurance.

Indemnity reserve construction (rules — not a finished recipe)
Any indemnity case reserve for accepted stock coverage must be formula-driven from the
cleaned eligible inventory population and the adjustments above so that changing an
eligible lot’s included amount changes the reserve and changing an excluded, VOID, or
duplicate lot does not. Keep ALAE / expense payments from the Payments tab separate from
the indemnity case reserve; do not roll adjuster, lab, or counsel expense into the
indemnity recommendation. Payment proposals above $250,000 require manager approval;
show whether the recommended indemnity reserve sits under or over that threshold.

Committee workbook expectations
Document Stage 1 coverage analysis (stock, equipment, BI) with source citations before
Stage 2 financial build. Record material conflicts (stale notes vs memo; superseded vs
Active schedule; selling-price demand vs valuation clause) with governing-source treatment.
Show inventory reconciliation and valuation reconciliation. Perform a simple sensitivity
check on covered vs excluded inventory. Derive dispositions and dollars from the rules and
evidence; do not treat this memo as a completed coverage opinion.

Documents reviewed
ashcroft-foods_claim_file.xlsx (Loss Notice, Inventory, Equipment, BI Claim, Adjuster Notes,
Payments); ashcroft-foods_policy_schedule.csv; refrigeration maintenance log excerpts
supplied by the insured; Zone B logger export summary; North River cancellation email
referenced in the claim-file notes.

Distribution
File-review committee; Columbus property desk; broker copy after committee (coverage
position only — after independent determination).
"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.strip() + "\n", encoding="utf-8")


def main() -> None:
    TASK_INPUTS.mkdir(parents=True, exist_ok=True)
    write_claim_file(TASK_INPUTS / CLAIM_FILE)
    write_policy_csv(TASK_INPUTS / POLICY_FILE)
    write_memo(TASK_INPUTS / MEMO_FILE)
    words = len((TASK_INPUTS / MEMO_FILE).read_text(encoding="utf-8").split())
    print(f"Wrote inputs to {TASK_INPUTS}")
    print(f"Memo word count: {words}")
    from ashcroft_constants import CASE_RESERVE, COVERED_STOCK_COST, AFTER_COINSURANCE, SALVAGE_AMOUNT, NET_AFTER_SALVAGE

    print(
        f"Golden targets: covered={COVERED_STOCK_COST} salvage={SALVAGE_AMOUNT} "
        f"net={NET_AFTER_SALVAGE} after_coin={AFTER_COINSURANCE} case={CASE_RESERVE}"
    )


if __name__ == "__main__":
    main()
