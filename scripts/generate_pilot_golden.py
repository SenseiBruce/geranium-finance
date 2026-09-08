#!/usr/bin/env python3
"""Generate pilot golden workbook for harborview-q1-consolidation."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

TASK = Path(__file__).resolve().parent.parent / "tasks" / "harborview-q1-consolidation" / "golden"
OUT = TASK / "consolidated_review_harborview_media_group.xlsx"

FX_RATE = 0.74
LATTICE_PRORATION = 59 / 90
LATTICE_GROSS = 1_063_627.65
CASCADE_CAD = 611_842.27
BEACON_USD = 440_566.01


def main() -> None:
    TASK.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Assumptions"
    ws["A1"] = "Reporting period"
    ws["B1"] = "2025-01-01 to 2025-03-31"
    ws["A2"] = "CAD/USD rate"
    ws["B2"] = FX_RATE
    ws["A3"] = "Lattice acquisition date"
    ws["B3"] = "2025-02-01"
    ws["A4"] = "Lattice proration factor"
    ws["B4"] = f"={LATTICE_PRORATION}"
    ws["A5"] = "Lattice ownership days in Q1"
    ws["B5"] = 59
    ws["A6"] = "Q1 calendar days"
    ws["B6"] = 90

    ws2 = wb.create_sheet("Source_Data")
    ws2.append(["Brand", "Gross USD/CAD", "Currency", "Notes"])
    ws2.append(["Beacon Media", BEACON_USD, "USD", "Full quarter"])
    ws2.append(["Cascade Creative", CASCADE_CAD, "CAD", "Convert at period rate"])
    ws2.append(["Lattice Digital", LATTICE_GROSS, "USD", "Acquired Feb 1"])

    ws3 = wb.create_sheet("Rollup")
    ws3["A1"] = "Brand"
    ws3["B1"] = "Gross"
    ws3["C1"] = "Proration"
    ws3["D1"] = "Consolidated USD"
    ws3["A2"] = "Beacon Media"
    ws3["B2"] = BEACON_USD
    ws3["C2"] = 1
    ws3["D2"] = "=B2*C2"
    ws3["A3"] = "Cascade Creative"
    ws3["B3"] = CASCADE_CAD
    ws3["C3"] = 1
    ws3["D3"] = "=B3*Assumptions!B2"
    ws3["A4"] = "Lattice Digital"
    ws3["B4"] = LATTICE_GROSS
    ws3["C4"] = "=Assumptions!B4"
    ws3["D4"] = "=B4*C4"
    ws3["A6"] = "Total consolidated Q1 2025"
    ws3["D6"] = "=SUM(D2:D4)"

    ws4 = wb.create_sheet("Reconciliation_Log")
    ws4.append(["Issue", "Action", "Source"])
    ws4.append(["Void/cancelled rows", "Excluded from revenue", "harborview_media_group_brand_ledger.csv"])
    ws4.append(["Post-period rows (HM-3007)", "Excluded", "harborview_media_group_brand_ledger.csv"])
    ws4.append(["Prior-year row HM-4001", "Excluded from Q1", "harborview_media_group_brand_ledger.csv"])
    ws4.append(["CAD label", "Converted at 0.74", "cascade_creative_fx_and_period_reference.txt"])
    ws4.append(["Lattice mid-period acquisition", "Prorated from Feb 1", "lattice_digital_acquisition_memo.txt"])

    ws5 = wb.create_sheet("Recommendation")
    ws5["A1"] = "Board recommendation"
    ws5["A2"] = (
        "Report consolidated Q1 2025 revenue after CAD conversion, transaction exclusions, "
        "and Lattice acquisition-date proration."
    )
    ws5["A3"] = "Consolidated total"
    ws5["B3"] = "=Rollup!D6"

    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
