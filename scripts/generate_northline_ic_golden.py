#!/usr/bin/env python3
"""Generate golden intercompany settlement workbook for Northline."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

from northline_ic_constants import (
    AP_C_441_CAD,
    AP_C_441_FX_DIFF,
    AP_C_441_USD_BOOKED,
    AP_C_441_USD_MEMO,
    AP_C_442_CAD,
    AP_C_442_USD_MEMO,
    AP_C_455_ORPHAN_CAD,
    AP_CSV,
    AP_M_205_MXN,
    AP_M_205_USD_MEMO,
    AP_M_210_MXN,
    AP_M_210_USD_MEMO,
    APPENDIX_A_MIN_USD,
    AR_8688,
    AR_8766,
    AR_8790,
    AR_8812,
    AR_8820_VOID,
    AR_8841,
    AR_8855,
    AR_CSV,
    AS_OF,
    AUTHOR,
    CAD_PER_USD,
    DELIVERABLE,
    ENTITY,
    FX_DIFF_CANADA,
    FX_DIFF_MEXICO,
    IN_TRANSIT_TOTAL,
    MATCHED_AR_CANADA,
    MATCHED_AR_MEXICO,
    MEMO_TXT,
    MXN_PER_USD,
    ORPHAN_USD_AT_MEMO,
    SETTLE_AP_CANADA_USD,
    SETTLE_AP_MEXICO_USD,
    SETTLE_AR_CANADA,
    SETTLE_AR_MEXICO,
    SETTLE_AR_TOTAL,
    SLUG,
    STALE_CAD_PER_USD,
    WIRE_DATE,
)
from sanitize_office import cache_xlsx_formula_values, sanitize_xlsx


def _set_workbook_author(path: Path, author: str) -> None:
    """Set creator / lastModifiedBy and clear OpenPyXL Application tell."""
    import zipfile
    from tempfile import NamedTemporaryFile
    from xml.etree import ElementTree as ET

    ns_cp = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
    ns_dc = "http://purl.org/dc/elements/1.1/"
    ns_dcterms = "http://purl.org/dc/terms/"
    ns_ep = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"

    with zipfile.ZipFile(path, "r") as zin:
        names = set(zin.namelist())
        parts = {n: zin.read(n) for n in names}

    if "docProps/core.xml" in parts:
        root = ET.fromstring(parts["docProps/core.xml"])
        creator = root.find(f"{{{ns_dc}}}creator")
        if creator is None:
            creator = ET.SubElement(root, f"{{{ns_dc}}}creator")
        creator.text = author
        last_mod = root.find(f"{{{ns_cp}}}lastModifiedBy")
        if last_mod is None:
            last_mod = ET.SubElement(root, f"{{{ns_cp}}}lastModifiedBy")
        last_mod.text = author
        ET.register_namespace("cp", ns_cp)
        ET.register_namespace("dc", ns_dc)
        ET.register_namespace("dcterms", ns_dcterms)
        parts["docProps/core.xml"] = ET.tostring(
            root, encoding="utf-8", xml_declaration=True
        )

    if "docProps/app.xml" in parts:
        root = ET.fromstring(parts["docProps/app.xml"])
        app = root.find(f"{{{ns_ep}}}Application")
        if app is None:
            app = ET.SubElement(root, f"{{{ns_ep}}}Application")
        app.text = "Microsoft Excel"
        ET.register_namespace("", ns_ep)
        parts["docProps/app.xml"] = ET.tostring(
            root, encoding="utf-8", xml_declaration=True
        )

    with NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        tmp_path = Path(tmp.name)
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(
        tmp_path, "w", compression=zipfile.ZIP_DEFLATED
    ) as zout:
        for info in zin.infolist():
            zout.writestr(info, parts[info.filename])
    tmp_path.replace(path)

OUT = Path(__file__).resolve().parent.parent / "tasks" / SLUG / "golden" / DELIVERABLE
header_fill = PatternFill("solid", fgColor="37474F")
header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
title_font = Font(bold=True, size=13, name="Calibri")
section_font = Font(bold=True, size=11, name="Calibri")
thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
pass_fill = PatternFill("solid", fgColor="C8E6C9")
excl_fill = PatternFill("solid", fgColor="FFECB3")
money_fmt = "#,##0.00"
rate_fmt = "0.0000"


def style_header(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin


def border_row(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        ws.cell(row, c).border = thin


def main() -> None:
    wb = Workbook()

    # ------------------------------------------------------------------
    # AR_Settlement — Appendix A population + void/provisional exclusions
    # ------------------------------------------------------------------
    ws = wb.active
    ws.title = "AR_Settlement"
    ws["A1"] = f"{ENTITY}: Appendix A AR settlement"
    ws["A1"].font = title_font
    ws["A2"] = f"As of {AS_OF} | Wire cycle {WIRE_DATE} | Governing: {MEMO_TXT}"
    headers = [
        "Invoice ID",
        "Counterparty",
        "Invoice date",
        "Currency",
        "Local / source USD",
        "USD amount",
        "Status",
        "Include",
        "Reason",
        "Source",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header(ws, 4, 10)

    # Rows 5-10 = Appendix A (discovered via memo materiality); 11 = void; Include drives SUMIFS
    mat = f"${APPENDIX_A_MIN_USD:,.2f}"
    ar_rows = [
        # invoice, cp, date, ccy, amount, status, include, reason
        ("INV-AR-8841", "Northline Canada Ltd", "2026-08-27", "USD", AR_8841, "InTransit", "Yes",
         f"Appendix A (≥{mat} + InTransit); FOB ship_date on/before Aug 31; AP not booked yet",
         f"{AR_CSV} + {MEMO_TXT} §2–3"),
        ("INV-AR-8790", "Northline Canada Ltd", "2026-08-11", "USD", AR_8790, "Open", "Yes",
         f"Appendix A (≥{mat} + Open); matched AP present", f"{AR_CSV} + {MEMO_TXT} §2"),
        ("INV-AR-8766", "Northline Canada Ltd", "2026-07-15", "USD", AR_8766, "Open", "Yes",
         f"Appendix A (≥{mat} + Open); management fee", f"{AR_CSV} + {MEMO_TXT} §2"),
        ("INV-AR-8812", "Northline Mexico SA de CV", "2026-08-14", "USD", AR_8812, "Open", "Yes",
         f"Appendix A (≥{mat} + Open); matched AP present", f"{AR_CSV} + {MEMO_TXT} §2"),
        ("INV-AR-8855", "Northline Mexico SA de CV", "2026-08-28", "USD", AR_8855, "InTransit", "Yes",
         f"Appendix A (≥{mat} + InTransit); FOB ship_date on/before Aug 31; AP not booked yet",
         f"{AR_CSV} + {MEMO_TXT} §2–3"),
        ("INV-AR-8688", "Northline Mexico SA de CV", "2026-08-04", "USD", AR_8688, "Open", "Yes",
         f"Appendix A (≥{mat} + Open); matched AP present", f"{AR_CSV} + {MEMO_TXT} §2"),
        ("INV-AR-8820", "Northline Canada Ltd", "2026-08-19", "USD", AR_8820_VOID, "Void", "No",
         "Void status — exclude from settlement (discovered on AR status_flag)", f"{AR_CSV} + {MEMO_TXT} §6"),
    ]

    for r, row in enumerate(ar_rows, 5):
        inv, cp, idate, ccy, amt, status, include, reason, source = row
        ws.cell(r, 1, inv)
        ws.cell(r, 2, cp)
        ws.cell(r, 3, idate)
        ws.cell(r, 4, ccy)
        ws.cell(r, 5, amt).number_format = money_fmt
        # USD amount equals source USD for parent invoices (formula link)
        ws.cell(r, 6, f"=E{r}").number_format = money_fmt
        ws.cell(r, 7, status)
        ws.cell(r, 8, include)
        ws.cell(r, 9, reason)
        ws.cell(r, 10, source)
        border_row(ws, r, 10)
        if include == "No":
            ws.cell(r, 8).fill = excl_fill

    # Totals via SUMIFS on Include=Yes by counterparty
    ws["A13"] = "Canada Appendix A AR (Include=Yes)"
    ws["F13"] = '=SUMIFS(F5:F11,B5:B11,"Northline Canada Ltd",H5:H11,"Yes")'
    ws["F13"].number_format = money_fmt
    ws["F13"].font = section_font
    ws["A14"] = "Mexico Appendix A AR (Include=Yes)"
    ws["F14"] = '=SUMIFS(F5:F11,B5:B11,"Northline Mexico SA de CV",H5:H11,"Yes")'
    ws["F14"].number_format = money_fmt
    ws["F14"].font = section_font
    ws["A15"] = "Total Appendix A AR settlement USD"
    ws["F15"] = "=F13+F14"
    ws["F15"].number_format = money_fmt
    ws["F15"].font = section_font
    ws["A16"] = (
        f"Appendix A rule: Open/InTransit + amount_usd ≥ {mat} "
        f"(memo §2). Control: Canada {SETTLE_AR_CANADA:,.2f} | Mexico {SETTLE_AR_MEXICO:,.2f}"
    )
    ws["A17"] = (
        f"Discovered InTransit kept per FOB §3: INV-AR-8841, INV-AR-8855 | "
        f"Void excluded §6: INV-AR-8820 | Source: {AR_CSV} + {MEMO_TXT}"
    )

    for col, w in enumerate([14, 28, 12, 8, 14, 12, 10, 8, 62, 42], 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    # ------------------------------------------------------------------
    # AP_FX_TieOut — matched AP at memo FX + stale-rate revalue
    # ------------------------------------------------------------------
    ws2 = wb.create_sheet("AP_FX_TieOut")
    ws2["A1"] = "Affiliate AP FX tie-out (memo August 31 rates)"
    ws2["A1"].font = title_font
    ws2["A2"] = (
        f"Memo CAD {CAD_PER_USD} / MXN {MXN_PER_USD} per USD | "
        f"AP-C-441 booked {STALE_CAD_PER_USD} revalued at memo | Governing: {MEMO_TXT} §4–5"
    )
    ap_headers = [
        "AP ID",
        "Parent invoice",
        "Counterparty",
        "Local amount",
        "CCY",
        "Booked FX",
        "Memo FX",
        "USD @ booked",
        "USD @ memo",
        "Parent AR USD",
        "FX difference",
        "Include",
        "Tie-out status",
        "Notes",
    ]
    for i, h in enumerate(ap_headers, 1):
        ws2.cell(4, i, h)
    style_header(ws2, 4, 14)

    # Row layout:
    # 5 AP-C-441, 6 AP-C-442, 7 AP-C-448 void, 8 AP-C-455 orphan,
    # 9 AP-M-210, 10 AP-M-205
    ap_data = [
        # ap, parent, cp, local, ccy, booked, memo, parent_ar, include, status, notes
        ("AP-C-441", "INV-AR-8790", "Northline Canada Ltd", AP_C_441_CAD, "CAD",
         STALE_CAD_PER_USD, CAD_PER_USD, AR_8790, "Yes", "Revalued",
         "Stale July booking rate; memo §5 requires revalue at 1.3724"),
        ("AP-C-442", "INV-AR-8766", "Northline Canada Ltd", AP_C_442_CAD, "CAD",
         CAD_PER_USD, CAD_PER_USD, AR_8766, "Yes", "Tied",
         "Matched INV-AR-8766 at memo rate"),
        ("AP-C-448", "INV-AR-8790", "Northline Canada Ltd", AP_C_441_CAD, "CAD",
         STALE_CAD_PER_USD, CAD_PER_USD, AR_8790, "No", "Excluded — void duplicate",
         "Void duplicate of AP-C-441; do not settle"),
        ("AP-C-455", "", "Northline Canada Ltd", AP_C_455_ORPHAN_CAD, "CAD",
         CAD_PER_USD, CAD_PER_USD, 0.0, "No", "Excluded — orphan",
         "No parent_invoice_ref; local reclass only"),
        ("AP-M-210", "INV-AR-8812", "Northline Mexico SA de CV", AP_M_210_MXN, "MXN",
         MXN_PER_USD, MXN_PER_USD, AR_8812, "Yes", "Tied",
         "Matched INV-AR-8812"),
        ("AP-M-205", "INV-AR-8688", "Northline Mexico SA de CV", AP_M_205_MXN, "MXN",
         MXN_PER_USD, MXN_PER_USD, AR_8688, "Yes", "Tied",
         "Matched INV-AR-8688"),
    ]
    for r, row in enumerate(ap_data, 5):
        ap, parent, cp, local, ccy, booked, memo, parent_ar, include, status, notes = row
        ws2.cell(r, 1, ap)
        ws2.cell(r, 2, parent)
        ws2.cell(r, 3, cp)
        ws2.cell(r, 4, local).number_format = money_fmt
        ws2.cell(r, 5, ccy)
        ws2.cell(r, 6, booked).number_format = rate_fmt
        ws2.cell(r, 7, memo).number_format = rate_fmt
        # USD @ booked = local / booked FX
        ws2.cell(r, 8, f"=IF(F{r}=0,0,D{r}/F{r})").number_format = money_fmt
        # USD @ memo = local / memo FX
        ws2.cell(r, 9, f"=IF(G{r}=0,0,D{r}/G{r})").number_format = money_fmt
        ws2.cell(r, 10, parent_ar).number_format = money_fmt
        # FX difference = parent AR − memo USD (matched only; zero when excluded/orphan)
        ws2.cell(r, 11, f'=IF(L{r}="Yes",J{r}-I{r},0)').number_format = money_fmt
        ws2.cell(r, 12, include)
        ws2.cell(r, 13, status)
        ws2.cell(r, 14, notes)
        border_row(ws2, r, 14)
        if include == "No":
            ws2.cell(r, 12).fill = excl_fill

    # In-transit note rows (no AP)
    ws2["A12"] = "(none)"
    ws2["B12"] = "INV-AR-8841"
    ws2["C12"] = "Northline Canada Ltd"
    ws2["L12"] = "n/a"
    ws2["M12"] = "No AP — InTransit"
    ws2["N12"] = f"FOB shipping point; AP expected on receipt — {MEMO_TXT} §3"
    border_row(ws2, 12, 14)

    ws2["A13"] = "(none)"
    ws2["B13"] = "INV-AR-8855"
    ws2["C13"] = "Northline Mexico SA de CV"
    ws2["L13"] = "n/a"
    ws2["M13"] = "No AP — InTransit"
    ws2["N13"] = f"FOB shipping point; AP expected on receipt — {MEMO_TXT} §3"
    border_row(ws2, 13, 14)

    ws2["A15"] = "Canada matched AP USD @ memo (Include=Yes)"
    ws2["I15"] = '=SUMIFS(I5:I10,C5:C10,"Northline Canada Ltd",L5:L10,"Yes")'
    ws2["I15"].number_format = money_fmt
    ws2["A16"] = "Mexico matched AP USD @ memo (Include=Yes)"
    ws2["I16"] = '=SUMIFS(I5:I10,C5:C10,"Northline Mexico SA de CV",L5:L10,"Yes")'
    ws2["I16"].number_format = money_fmt
    ws2["A17"] = "Canada FX difference on matched pairs"
    ws2["K17"] = '=SUMIFS(K5:K10,C5:C10,"Northline Canada Ltd",L5:L10,"Yes")'
    ws2["K17"].number_format = money_fmt
    ws2["A18"] = "AP-C-441 FX difference (parent AR − memo USD)"
    ws2["K18"] = "=K5"
    ws2["K18"].number_format = money_fmt
    ws2["A19"] = "Memo governs: do not use booked 1.3680 for AP-C-441 tie-out"

    for col, w in enumerate([10, 14, 26, 12, 6, 10, 10, 12, 12, 12, 12, 8, 22, 48], 1):
        ws2.column_dimensions[get_column_letter(col)].width = w

    # ------------------------------------------------------------------
    # Exceptions — seven memo §6 documentation fields per discovered row
    # ------------------------------------------------------------------
    ws4 = wb.create_sheet("Exceptions")
    ws4["A1"] = "Exception log: discovered from AR/AP ledgers under memo rules"
    ws4["A1"].font = title_font
    ws4["A2"] = (
        f"Governing memo: {MEMO_TXT} (rules only — exception IDs discovered in {AR_CSV} / {AP_CSV})"
    )
    ws4["A3"] = (
        "Memo §6 documentation fields: record ID | counterparty | exception type | "
        "source fact | governing memo rule | treatment | settlement impact"
    )
    for i, h in enumerate(
        [
            "Record ID",
            "Counterparty",
            "Exception type",
            "Source fact",
            "Governing memo rule",
            "Treatment",
            "Settlement impact",
            "Settlement include",
        ],
        1,
    ):
        ws4.cell(4, i, h)
    style_header(ws4, 4, 8)

    exceptions = [
        (
            "INV-AR-8841",
            "Northline Canada Ltd",
            "InTransit / cut-off",
            "status_flag=InTransit; ship_date 2026-08-28; no AP parent_invoice_ref",
            f"{MEMO_TXT} §2–3 / §6 test 7 (FOB; missing AP expected, not orphan)",
            "Include in Canada AR settlement; no AP expected for Sep 3 wire",
            f"Include ${AR_8841:,.2f} in Canada wire",
            "Yes",
        ),
        (
            "INV-AR-8855",
            "Northline Mexico SA de CV",
            "InTransit / cut-off",
            "status_flag=InTransit; ship_date 2026-08-29; no AP parent_invoice_ref",
            f"{MEMO_TXT} §2–3 / §6 test 7 (FOB; missing AP expected, not orphan)",
            "Include in Mexico AR settlement; no AP expected for Sep 3 wire",
            f"Include ${AR_8855:,.2f} in Mexico wire",
            "Yes",
        ),
        (
            "INV-AR-8820",
            "Northline Canada Ltd",
            "Void AR",
            "status_flag=Void on AR subledger",
            f"{MEMO_TXT} §6 test 1 (void AR never settled)",
            "Exclude from AR settlement and wire",
            f"Exclude ${AR_8820_VOID:,.2f} from Canada wire",
            "No",
        ),
        (
            "AP-C-448",
            "Northline Canada Ltd",
            "Void / duplicate AP",
            "status_flag=Void; same parent_invoice_ref/amount as Open AP-C-441",
            f"{MEMO_TXT} §6 tests 2–3 (void AP excluded; duplicate retains Open payable)",
            "Exclude from Canada AP tie-out and wire",
            "No wire; omit from matched AP USD",
            "No",
        ),
        (
            "AP-C-455",
            "Northline Canada Ltd",
            "Orphan AP",
            "status_flag=Open; blank parent_invoice_ref",
            f"{MEMO_TXT} §6 test 4 (orphan AP excluded / local reclass)",
            "Exclude from Canada settlement net; route to entity controllers",
            f"Do not wire ~${ORPHAN_USD_AT_MEMO:,.2f} memo USD",
            "No",
        ),
        (
            "AP-C-441",
            "Northline Canada Ltd",
            "Stale FX booking",
            f"fx_rate_used={STALE_CAD_PER_USD} vs memo CAD {CAD_PER_USD}; parent INV-AR-8790",
            f"{MEMO_TXT} §4–5 (revalue stale fx_rate_used; parent USD governs wire)",
            f"Revalue at memo {CAD_PER_USD} for tie-out; wire remains parent AR USD",
            f"FX difference ${AP_C_441_FX_DIFF:,.2f}; wire still ${AR_8790:,.2f}",
            "Yes",
        ),
    ]

    for r, row in enumerate(exceptions, 5):
        for c, v in enumerate(row, 1):
            ws4.cell(r, c, v)
        border_row(ws4, r, 8)
        if row[7] == "No":
            ws4.cell(r, 6).fill = excl_fill
            ws4.cell(r, 8).fill = excl_fill

    ws4["A12"] = "AP-C-441 FX difference (USD)"
    ws4["B12"] = "=AP_FX_TieOut!K5"
    ws4["B12"].number_format = money_fmt
    ws4["C12"] = "Parent INV-AR-8790 USD minus AP revalued at memo CAD rate"
    ws4["A13"] = "Orphan CAD at memo (info only - not wired)"
    ws4["B13"] = ORPHAN_USD_AT_MEMO
    ws4["B13"].number_format = money_fmt

    for col, w in enumerate([14, 26, 18, 58, 48, 52, 36, 14], 1):
        ws4.column_dimensions[get_column_letter(col)].width = w

    # ------------------------------------------------------------------
    # Section6_Tests — seven independent memo §6 exception criteria
    # ------------------------------------------------------------------
    ws6t = wb.create_sheet("Section6_Tests")
    ws6t["A1"] = "Memo §6 exception criteria — independent tests"
    ws6t["A1"].font = title_font
    ws6t["A2"] = (
        f"Source: {MEMO_TXT} §6 | Evaluated against {AR_CSV} / {AP_CSV} fields "
        "(not collapsed into a four-category summary)"
    )
    for i, h in enumerate(
        [
            "Test #",
            "Section 6 criterion",
            "Source field(s)",
            "Observed condition",
            "Result",
            "Treatment",
            "Record ID",
        ],
        1,
    ):
        ws6t.cell(4, i, h)
    style_header(ws6t, 4, 7)

    section6_tests = [
        (
            1,
            "Void AR: status_flag Void excluded from settlement and Sep 3 wire",
            "AR status_flag",
            "INV-AR-8820 status_flag=Void",
            "FAIL (exception)",
            "Exclude from AR settlement / wire",
            "INV-AR-8820",
        ),
        (
            2,
            "Void AP: status_flag Void excluded from AP FX tie-out and not wired",
            "AP status_flag",
            "AP-C-448 status_flag=Void",
            "FAIL (exception)",
            "Exclude from AP tie-out / wire",
            "AP-C-448",
        ),
        (
            3,
            "Duplicate AP: same parent_invoice_ref with Void sibling — exclude Void, retain Open",
            "AP parent_invoice_ref; status_flag",
            "AP-C-448 and AP-C-441 both reference INV-AR-8790; AP-C-448 Void / AP-C-441 Open",
            "FAIL (exception on Void)",
            "Exclude Void duplicate; retain Open AP-C-441 for tie-out",
            "AP-C-448",
        ),
        (
            4,
            "Orphan AP: Open AP with blank parent_invoice_ref excluded; local reclass",
            "AP status_flag; parent_invoice_ref",
            "AP-C-455 Open with blank parent_invoice_ref",
            "FAIL (exception)",
            "Exclude from settlement net; route to entity controllers",
            "AP-C-455",
        ),
        (
            5,
            "Matched AP only: FX tie-out includes only Open (non-void, non-orphan) AP",
            "AP status_flag; Include flag",
            "AP-C-441/442 and AP-M-210/205 Open+Include=Yes; void/orphan Include=No",
            "PASS",
            "Matched Open AP enter tie-out; void/orphan do not",
            "AP-C-441; AP-C-442; AP-M-210; AP-M-205",
        ),
        (
            6,
            "Matched AP only: parent_invoice_ref equals an Appendix A AR included in settlement",
            "AP parent_invoice_ref; AR Include",
            "Matched refs INV-AR-8790/8766/8812/8688 are Appendix A Include=Yes",
            "PASS",
            "Only Appendix A parent refs enter matched AP set",
            "AP-C-441; AP-C-442; AP-M-210; AP-M-205",
        ),
        (
            7,
            "InTransit missing AP is expected and is not an orphan",
            "AR status_flag; AP parent_invoice_ref absence",
            "INV-AR-8841 / INV-AR-8855 InTransit; no matching AP parent_invoice_ref",
            "PASS (expected gap)",
            "Keep InTransit AR in settlement; do not classify as orphan AP",
            "INV-AR-8841; INV-AR-8855",
        ),
    ]
    for r, row in enumerate(section6_tests, 5):
        for c, v in enumerate(row, 1):
            ws6t.cell(r, c, v)
        border_row(ws6t, r, 7)
        if str(row[4]).startswith("FAIL"):
            ws6t.cell(r, 5).fill = excl_fill
        else:
            ws6t.cell(r, 5).fill = pass_fill

    ws6t["A13"] = "§6 tests evaluated (count)"
    ws6t["B13"] = "=COUNTA(A5:A11)"
    ws6t["A14"] = "Required §6 test count"
    ws6t["B14"] = 7
    ws6t["A15"] = "All seven §6 tests present"
    ws6t["B15"] = '=IF(B13=B14,"PASS","FAIL")'
    ws6t["B15"].fill = pass_fill

    for col, w in enumerate([8, 72, 36, 70, 18, 52, 36], 1):
        ws6t.column_dimensions[get_column_letter(col)].width = w

    # ------------------------------------------------------------------
    # Consistency — row classification vs exception log vs settlement
    # ------------------------------------------------------------------
    wsc = wb.create_sheet("Consistency")
    wsc["A1"] = "Exception self-consistency control"
    wsc["A1"].font = title_font
    wsc["A2"] = (
        "Compares AR/AP Include flags, Exceptions treatment, and settlement inclusion. "
        "FAIL if a row is both settlement-included and exception-excluded, or vice versa."
    )
    for i, h in enumerate(
        [
            "Record ID",
            "Source sheet Include",
            "Exceptions settlement include",
            "Aligned?",
            "Control note",
        ],
        1,
    ):
        wsc.cell(4, i, h)
    style_header(wsc, 4, 5)

    # Keyed checks (ID-based, not fragile source-CSV row positions)
    consistency_rows = [
        ("INV-AR-8841", "AR_Settlement!H5", "Exceptions!H5", "InTransit kept in settlement"),
        ("INV-AR-8855", "AR_Settlement!H9", "Exceptions!H6", "InTransit kept in settlement"),
        ("INV-AR-8820", "AR_Settlement!H11", "Exceptions!H7", "Void AR excluded both places"),
        ("AP-C-448", "AP_FX_TieOut!L7", "Exceptions!H8", "Void duplicate excluded both places"),
        ("AP-C-455", "AP_FX_TieOut!L8", "Exceptions!H9", "Orphan excluded both places"),
        ("AP-C-441", "AP_FX_TieOut!L5", "Exceptions!H10", "Matched stale FX retained in tie-out"),
    ]
    for r, (rid, src_ref, exc_ref, note) in enumerate(consistency_rows, 5):
        wsc.cell(r, 1, rid)
        wsc.cell(r, 2, f"={src_ref}")
        wsc.cell(r, 3, f"={exc_ref}")
        wsc.cell(r, 4, f'=IF(B{r}=C{r},"PASS","FAIL")')
        wsc.cell(r, 5, note)
        border_row(wsc, r, 5)

    wsc["A12"] = "Contradiction count (FAIL rows)"
    wsc["B12"] = '=COUNTIF(D5:D10,"FAIL")'
    wsc["A13"] = "Self-consistency control"
    wsc["B13"] = '=IF(B12=0,"PASS","FAIL")'
    wsc["B13"].font = section_font
    wsc["B13"].fill = pass_fill
    wsc["A14"] = (
        "PASS requires: no Include=Yes row treated as settlement-excluded in Exceptions, "
        "and no exception-excluded row still Include=Yes on AR/AP sheets."
    )

    for col, w in enumerate([14, 22, 26, 12, 42], 1):
        wsc.column_dimensions[get_column_letter(col)].width = w

    # ------------------------------------------------------------------
    # Counterparty_Netting
    # ------------------------------------------------------------------
    ws3 = wb.create_sheet("Counterparty_Netting")
    ws3["A1"] = "Counterparty netting: September 3 wire basis"
    ws3["A1"].font = title_font
    ws3["A2"] = "Parent invoice USD governs wires; converted AP is tie-out only (memo §5)"
    for i, h in enumerate(
        [
            "Counterparty",
            "Valid AR USD",
            "Matched AP USD @ memo",
            "FX adjustment",
            "Net settlement USD",
            "Direction",
            "Exception impact",
        ],
        1,
    ):
        ws3.cell(4, i, h)
    style_header(ws3, 4, 7)

    ws3["A5"] = "Northline Canada Ltd"
    ws3["B5"] = "=AR_Settlement!F13"
    ws3["C5"] = "=AP_FX_TieOut!I15"
    ws3["D5"] = "=AP_FX_TieOut!K17"
    ws3["E5"] = "=B5"  # parent AR governs
    ws3["F5"] = "Affiliate pays parent"
    ws3["G5"] = "Void AP-C-448 and orphan AP-C-455 excluded; InTransit INV-AR-8841 included"
    for col in range(2, 6):
        ws3.cell(5, col).number_format = money_fmt
    border_row(ws3, 5, 7)

    ws3["A6"] = "Northline Mexico SA de CV"
    ws3["B6"] = "=AR_Settlement!F14"
    ws3["C6"] = "=AP_FX_TieOut!I16"
    # Matched Mexico AR = total Mexico AR − InTransit INV-AR-8855 (AR_Settlement!F9)
    ws3["D6"] = "=(B6-AR_Settlement!F9)-C6"
    ws3["E6"] = "=B6"
    ws3["F6"] = "Affiliate pays parent"
    ws3["G6"] = "InTransit INV-AR-8855 included; no void/orphan on Mexico matched set"
    for col in range(2, 6):
        ws3.cell(6, col).number_format = money_fmt
    border_row(ws3, 6, 7)

    ws3["A8"] = "Total net settlement / wire instruction USD"
    ws3["E8"] = "=E5+E6"
    ws3["E8"].number_format = money_fmt
    ws3["E8"].font = section_font
    ws3["A9"] = "Do not wire void INV-AR-8820, void AP-C-448, or orphan AP-C-455."
    ws3["A10"] = f"Control totals: Canada {SETTLE_AR_CANADA:,.2f} | Mexico {SETTLE_AR_MEXICO:,.2f} | Total {SETTLE_AR_TOTAL:,.2f}"

    for col, w in enumerate([28, 14, 18, 14, 16, 20, 55], 1):
        ws3.column_dimensions[get_column_letter(col)].width = w

    # ------------------------------------------------------------------
    # Recommendation
    # ------------------------------------------------------------------
    ws5 = wb.create_sheet("Recommendation")
    ws5["A1"] = f"{ENTITY}: September 3 intercompany settlement recommendation"
    ws5["A1"].font = title_font
    ws5["A2"] = f"Balances as of {AS_OF} | Wire date {WIRE_DATE} | Citation: {MEMO_TXT}"

    ws5["A4"] = "Ready to net / wire"
    ws5["A4"].font = section_font
    ws5["A5"] = "Northline Canada Ltd pays parent (Appendix A AR, incl. InTransit INV-AR-8841)"
    ws5["B5"] = "=Counterparty_Netting!E5"
    ws5["B5"].number_format = money_fmt
    ws5["C5"] = "Ready to wire"
    ws5["C5"].fill = pass_fill
    ws5["A6"] = "Northline Mexico SA de CV pays parent (Appendix A AR, incl. InTransit INV-AR-8855)"
    ws5["B6"] = "=Counterparty_Netting!E6"
    ws5["B6"].number_format = money_fmt
    ws5["C6"] = "Ready to wire"
    ws5["C6"].fill = pass_fill
    ws5["A7"] = "Total USD to collect from affiliates"
    ws5["B7"] = "=B5+B6"
    ws5["B7"].number_format = money_fmt
    ws5["B7"].font = section_font

    ws5["A9"] = "Exception handling: do not wire"
    ws5["A9"].font = section_font
    ws5["A10"] = "INV-AR-8820 (Void): exclude from Canada wire"
    ws5["A11"] = "AP-C-448 (Void duplicate): exclude from AP tie-out and settlement"
    ws5["A12"] = "AP-C-455 (Orphan): exclude from Canada net; send to entity controllers for local AP reclass"

    ws5["A14"] = "Investigation / tie-out only"
    ws5["A14"].font = section_font
    ws5["A15"] = (
        f"AP-C-441 booked at CAD {STALE_CAD_PER_USD}; revalue at memo {CAD_PER_USD}. "
        f"FX difference vs INV-AR-8790 is captured on AP_FX_TieOut (approx ${AP_C_441_FX_DIFF:,.2f}). "
        "Wire the parent invoice USD, not the converted AP USD."
    )
    ws5["A15"].alignment = Alignment(wrap_text=True)
    ws5.merge_cells("A15:D15")
    ws5.row_dimensions[15].height = 48

    ws5["A17"] = "Controller recommendation"
    ws5["A17"].font = section_font
    ws5["A18"] = (
        f"Proceed with the {WIRE_DATE} sweep collecting Appendix A AR from Canada and Mexico "
        f"per {MEMO_TXT}. Keep InTransit INV-AR-8841 and INV-AR-8855 in the population (FOB shipping point). "
        "Use August 31 memo FX only to tie out matched AP. Do not instruct treasury to wire voids or the orphan payable. "
        "Parent invoice USD amounts govern the wires."
    )
    ws5["A18"].alignment = Alignment(wrap_text=True)
    ws5.merge_cells("A18:D18")
    ws5.row_dimensions[18].height = 72

    ws5.column_dimensions["A"].width = 78
    ws5.column_dimensions["B"].width = 14
    ws5.column_dimensions["C"].width = 10

    # ------------------------------------------------------------------
    # Notes
    # ------------------------------------------------------------------
    ws_notes = wb.create_sheet("Notes")
    ws_notes["A1"] = "Working notes / methodology"
    ws_notes["A1"].font = title_font
    notes = [
        (
            f"1. {MEMO_TXT} governs where the AR and AP ledgers conflict, "
            "including the applicable cut-off, FX, and exception treatment."
        ),
        f"2. Sources: {AR_CSV}, {AP_CSV}, {MEMO_TXT}.",
        (
            f"3. Appendix A discovered from AR attributes (memo §2): Open or InTransit, "
            f"Canada/Mexico counterparty, amount_usd ≥ ${APPENDIX_A_MIN_USD:,.2f} "
            "→ INV-AR-8841/8790/8766/8812/8855/8688."
        ),
        (
            "4. FOB shipping-point (memo §3): InTransit Appendix A rows INV-AR-8841 and "
            "INV-AR-8855 stay in settlement even though affiliate AP is not booked yet."
        ),
        (
            f"5. August 31 FX (memo §4–5): CAD {CAD_PER_USD}, MXN {MXN_PER_USD} per USD. "
            f"Stale booking AP-C-441 (fx_rate_used {STALE_CAD_PER_USD}) revalued at memo; "
            "parent invoice USD still governs the wire."
        ),
        "6. Tie-out conversion: USD = local amount / (foreign units per 1 USD).",
        (
            "7. Memo §6 seven tests applied independently on Section6_Tests: "
            "void AR; void AP; duplicate AP; orphan AP; matched Open-only; "
            "matched parent_invoice_ref=Appendix A; InTransit missing AP not orphan."
        ),
        "8. Consistency sheet confirms Include flags align with Exceptions settlement treatment.",
        "9. Totals and FX differences are formula-linked to the AR_Settlement and AP_FX_TieOut rows.",
    ]
    for i, line in enumerate(notes, 3):
        ws_notes.cell(i, 1, line)
        ws_notes.cell(i, 1).alignment = Alignment(wrap_text=True)
    ws_notes.column_dimensions["A"].width = 120
    for i in range(3, 13):
        ws_notes.row_dimensions[i].height = 32

    # Author metadata before save (sanitize clears tool creators only)
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    sanitize_xlsx(OUT)

    formula_cache = {
        "AR_Settlement": {
            "F5": AR_8841,
            "F6": AR_8790,
            "F7": AR_8766,
            "F8": AR_8812,
            "F9": AR_8855,
            "F10": AR_8688,
            "F11": AR_8820_VOID,
            "F13": SETTLE_AR_CANADA,
            "F14": SETTLE_AR_MEXICO,
            "F15": SETTLE_AR_TOTAL,
        },
        "AP_FX_TieOut": {
            "H5": AP_C_441_USD_BOOKED,
            "I5": AP_C_441_USD_MEMO,
            "K5": AP_C_441_FX_DIFF,
            "H6": AP_C_442_USD_MEMO,
            "I6": AP_C_442_USD_MEMO,
            "K6": round(AR_8766 - AP_C_442_USD_MEMO, 2),
            "H7": AP_C_441_USD_BOOKED,
            "I7": AP_C_441_USD_MEMO,
            "K7": 0.0,
            "H8": ORPHAN_USD_AT_MEMO,
            "I8": ORPHAN_USD_AT_MEMO,
            "K8": 0.0,
            "H9": AP_M_210_USD_MEMO,
            "I9": AP_M_210_USD_MEMO,
            "K9": round(AR_8812 - AP_M_210_USD_MEMO, 2),
            "H10": AP_M_205_USD_MEMO,
            "I10": AP_M_205_USD_MEMO,
            "K10": round(AR_8688 - AP_M_205_USD_MEMO, 2),
            "I15": SETTLE_AP_CANADA_USD,
            "I16": SETTLE_AP_MEXICO_USD,
            "K17": FX_DIFF_CANADA,
            "K18": AP_C_441_FX_DIFF,
        },
        "Exceptions": {
            "B12": AP_C_441_FX_DIFF,
        },
        "Section6_Tests": {
            "B13": 7,
            "B15": "PASS",
        },
        "Consistency": {
            "B5": "Yes",
            "C5": "Yes",
            "D5": "PASS",
            "B6": "Yes",
            "C6": "Yes",
            "D6": "PASS",
            "B7": "No",
            "C7": "No",
            "D7": "PASS",
            "B8": "No",
            "C8": "No",
            "D8": "PASS",
            "B9": "No",
            "C9": "No",
            "D9": "PASS",
            "B10": "Yes",
            "C10": "Yes",
            "D10": "PASS",
            "B12": 0,
            "B13": "PASS",
        },
        "Counterparty_Netting": {
            "B5": SETTLE_AR_CANADA,
            "C5": SETTLE_AP_CANADA_USD,
            "D5": FX_DIFF_CANADA,
            "E5": SETTLE_AR_CANADA,
            "B6": SETTLE_AR_MEXICO,
            "C6": SETTLE_AP_MEXICO_USD,
            "D6": FX_DIFF_MEXICO,
            "E6": SETTLE_AR_MEXICO,
            "E8": SETTLE_AR_TOTAL,
        },
        "Recommendation": {
            "B5": SETTLE_AR_CANADA,
            "B6": SETTLE_AR_MEXICO,
            "B7": SETTLE_AR_TOTAL,
        },
    }
    cache_xlsx_formula_values(OUT, formula_cache)
    # Re-apply author after sanitize / formula cache rewrite (both touch the zip)
    _set_workbook_author(OUT, AUTHOR)
    print(f"Wrote {OUT}")
    print(f"  Canada {SETTLE_AR_CANADA} Mexico {SETTLE_AR_MEXICO} Total {SETTLE_AR_TOTAL}")
    print(f"  AP-C-441 FX diff {AP_C_441_FX_DIFF} | Canada FX {FX_DIFF_CANADA}")
    print(f"  AP USD C {SETTLE_AP_CANADA_USD} M {SETTLE_AP_MEXICO_USD}")
    print(f"  In-transit {IN_TRANSIT_TOTAL} | matched CA {MATCHED_AR_CANADA} MX {MATCHED_AR_MEXICO}")
    print(f"  Author: {AUTHOR}")


if __name__ == "__main__":
    main()
