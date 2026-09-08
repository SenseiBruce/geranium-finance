#!/usr/bin/env python3
"""Generate golden workbook for granite-ridge-oct-renewal-desk."""

from __future__ import annotations

import csv
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

from sanitize_office import cache_xlsx_formula_values, sanitize_xlsx

TASK = Path(__file__).resolve().parent.parent / "tasks" / "granite-ridge-oct-renewal-desk"
INPUTS = TASK / "inputs"
OUT = TASK / "golden" / "renewal_underwriting_review.xlsx"

EFFECTIVE = date(2025, 10, 1)
INSPECTION_GRACE_DAYS = 30
LOSS_LAST = 120
ACCOUNT_LAST = 40

QUOTE_RATIONALE_VARIANTS = [
    "No Section 3-6 referral triggers after reconciling granite_ridge_loss_run_36mo.csv and granite_ridge_location_schedule.csv.",
    "Loss ratio and location checks clear under granite_ridge_routing_notes.txt with no pending broker items.",
    "Account summary ties to location detail; no inspection, vacancy, or subsidized-housing flags require referral.",
    "Attritional loss history within appetite after reconciling granite_ridge_loss_run_36mo.csv; no referral triggers active.",
    "Renewal-ready at summary level confirmed once location schedule and loss run were reconciled.",
]

STUDENT_HOUSING_ACCOUNTS = {"GR-2007", "GR-2065"}
STUDENT_HOUSING_NOTE = (
    "Section 10: GRP-APPETITE review must be submitted before bind "
    "(Campus Edge student-housing account; clearance not on file)"
)

HEADER_FILL = PatternFill("solid", fgColor="434343")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Calibri", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, size=14)
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def style_header_row(ws, row: int = 1) -> None:
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    ws.row_dimensions[row].height = 28


def style_body(ws, start_row: int, end_row: int, wrap_cols: set[int] | None = None) -> None:
    wrap_cols = wrap_cols or set()
    for r in range(start_row, end_row + 1):
        for c, cell in enumerate(ws[r], start=1):
            if cell.value is None and c not in wrap_cols:
                continue
            cell.font = BODY_FONT
            cell.border = THIN
            if c in wrap_cols:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, 45)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=False)


def set_widths(ws, widths: dict[str, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def desk_print_setup(ws, last_col: str, last_row: int, *, landscape: bool = True) -> None:
    ws.freeze_panes = "A2"
    if last_row >= 2:
        ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    ws.print_title_rows = "1:1"
    ws.print_area = f"A1:{last_col}{last_row}"
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.sheet_view.showGridLines = False


def format_workbook(
    wb,
    n_accounts: int,
    n_losses: int,
    n_exceptions: int,
    n_recon: int,
    n_citations: int,
) -> None:
    ws = wb["Account_Source"]
    style_header_row(ws)
    style_body(ws, 2, n_accounts + 1)
    set_widths(ws, {"A": 12, "B": 34, "C": 16, "D": 14})
    desk_print_setup(ws, "D", n_accounts + 1)

    ws = wb["Loss_Run"]
    style_header_row(ws)
    style_body(ws, 2, n_losses + 1)
    set_widths(ws, {"A": 12, "B": 12, "C": 14, "D": 13, "E": 14, "F": 12, "G": 14})
    desk_print_setup(ws, "G", n_losses + 1)

    ws = wb["Account_Routing"]
    style_header_row(ws)
    style_body(ws, 2, n_accounts + 1, wrap_cols={7})
    set_widths(ws, {"A": 12, "B": 34, "C": 15, "D": 13, "E": 16, "F": 12, "G": 62})
    desk_print_setup(ws, "G", n_accounts + 1)

    ws = wb["Location_Exceptions"]
    last = max(n_exceptions + 1, 2)
    style_header_row(ws)
    if n_exceptions:
        style_body(ws, 2, last, wrap_cols={4})
    set_widths(ws, {"A": 16, "B": 12, "C": 36, "D": 58, "E": 36})
    desk_print_setup(ws, "E", last)

    ws = wb["Loss_Reconciliation"]
    last = max(n_recon + 1, 2)
    style_header_row(ws)
    if n_recon:
        style_body(ws, 2, last, wrap_cols={5})
    set_widths(ws, {"A": 12, "B": 14, "C": 18, "D": 10, "E": 52})
    desk_print_setup(ws, "E", last)

    ws = wb["Rules_Citations"]
    last = max(n_citations + 1, 2)
    style_header_row(ws)
    style_body(ws, 2, last, wrap_cols={1})
    set_widths(ws, {"A": 48, "B": 12, "C": 36})
    desk_print_setup(ws, "C", last, landscape=False)

    ws = wb["Summary"]
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")
    for r in range(2, 9):
        for col in ("A", "B"):
            cell = ws[f"{col}{r}"]
            if cell.value is None and r == 7:
                continue
            cell.font = BODY_FONT
            cell.border = THIN
            cell.alignment = Alignment(vertical="center")
        if ws[f"A{r}"].value:
            ws[f"A{r}"].font = Font(name="Calibri", bold=True, size=11)
    set_widths(ws, {"A": 32, "B": 14})
    ws.freeze_panes = "A2"
    ws.print_area = "A1:B8"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.5, bottom=0.5)
    ws.sheet_view.showGridLines = False

    order = [
        "Summary",
        "Account_Routing",
        "Location_Exceptions",
        "Loss_Reconciliation",
        "Rules_Citations",
        "Account_Source",
        "Loss_Run",
    ]
    for idx, name in enumerate(order):
        wb.move_sheet(name, offset=idx - wb.sheetnames.index(name))


def load_csv(name: str) -> list[dict[str, str]]:
    with (INPUTS / name).open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def parse_date(value: str) -> date:
    y, m, d = value.split("-")
    return date(int(y), int(m), int(d))


def attritional_incurred(losses: list[dict[str, str]]) -> float:
    return sum(
        float(x["paid_amount"]) + float(x["reserved_amount"])
        for x in losses
        if not x["cat_code"].strip()
    )


def loc_incurred(losses: list[dict[str, str]]) -> float:
    return attritional_incurred(losses)


def inspection_expired(inspection_date: str) -> bool:
    insp = parse_date(inspection_date)
    return (EFFECTIVE - insp).days > INSPECTION_GRACE_DAYS


def quote_rationale(acct_id: str, acct_losses: list[dict[str, str]] | None = None) -> str:
    cat_codes = sorted({
        x["cat_code"].strip()
        for x in (acct_losses or [])
        if x.get("cat_code", "").strip()
    })
    if cat_codes:
        codes = ", ".join(cat_codes)
        return (
            f"Section 9: excluded catastrophe-coded loss {codes} from attritional LR "
            f"per granite_ridge_routing_notes.txt; remaining attritional history within appetite "
            f"with no Section 3-6 referral triggers."
        )
    idx = int(acct_id.split("-")[1]) % len(QUOTE_RATIONALE_VARIANTS)
    return QUOTE_RATIONALE_VARIANTS[idx]


def route_account(
    acct: dict[str, str],
    locs: list[dict[str, str]],
    losses: list[dict[str, str]],
) -> tuple[str, str, float, float, list[str]]:
    premium = float(acct["expiring_premium"])
    incurred = attritional_incurred(losses)
    lr = incurred / premium * 100 if premium else 0.0
    list_lr = float(acct["account_loss_ratio_36mo"])
    tiv_acct = float(acct["account_tiv"])
    tiv_loc = sum(float(x["building_tiv"]) for x in locs)
    tiv_var_pct = abs(tiv_acct - tiv_loc) / tiv_acct * 100 if tiv_acct else 0.0
    units_acct = int(acct["unit_count"])
    units_loc = sum(int(x["unit_count"]) for x in locs)
    unit_var_pct = abs(units_acct - units_loc) / units_acct * 100 if units_acct else 0.0

    triggers: list[str] = []
    rationale_parts: list[str] = []

    if lr >= 70:
        triggers.append("decline_lr")
        rationale_parts.append(f"Section 3 decline: recomputed LR {lr:.1f}% from granite_ridge_loss_run_36mo.csv")
    elif lr >= 45:
        triggers.append("refer_lr")
        rationale_parts.append(f"Section 3 refer: recomputed LR {lr:.1f}% from granite_ridge_loss_run_36mo.csv")

    for loc in locs:
        loc_losses = [x for x in losses if x["location_id"] == loc["location_id"]]
        loc_total = loc_incurred(loc_losses)
        if loc_total >= 250_000:
            triggers.append("refer_loc")
            rationale_parts.append(
                f"Section 3 refer: {loc['location_id']} incurred ${loc_total:,.0f} per granite_ridge_loss_run_36mo.csv"
            )
        if inspection_expired(loc["sprinkler_inspection_date"]) or inspection_expired(
            loc["fire_alarm_inspection_date"]
        ):
            triggers.append("refer_insp")
            rationale_parts.append(
                f"Section 5 refer: {loc['location_id']} life-safety inspection past grace per granite_ridge_location_schedule.csv"
            )
        if loc["pool_liability_flag"] == "pending_fence_po":
            triggers.append("refer_pool")
            rationale_parts.append(
                f"Section 5 refer: {loc['location_id']} pool fence remediation pending per granite_ridge_location_schedule.csv"
            )
        if int(loc["subsidized_unit_count"]) > 0:
            triggers.append("subsidized")
            rationale_parts.append(
                f"Section 6 refer: {loc['location_id']} subsidized units require endorsement documentation"
            )
        if loc["vacancy_flag"] == "Y" and loc["broker_confirmation"] == "no":
            triggers.append("ask_broker")
            rationale_parts.append(
                f"Section 6 Ask Broker: {loc['location_id']} vacancy without broker confirmation"
            )
        # Section 6 Refer applies only when occupancy_pct *implies full occupancy*.
        # No numeric full-occupancy threshold is stated in the routing notes; do not invent one.
        # Vacancy with broker_confirmation=yes and non-full occupancy is not a Section 6 Refer.

    if tiv_var_pct > 3:
        triggers.append("refer_tiv")
        rationale_parts.append(
            f"Section 4 refer: TIV variance {tiv_var_pct:.1f}% between granite_ridge_renewal_accounts.csv and granite_ridge_location_schedule.csv"
        )
    if unit_var_pct > 5:
        triggers.append("refer_units")
        rationale_parts.append(
            f"Section 4 refer: unit count variance {unit_var_pct:.1f}% between granite_ridge_renewal_accounts.csv and granite_ridge_location_schedule.csv"
        )

    # Section 10: Campus Edge / student-housing GRP-APPETITE review is a mandatory
    # pre-bind referral trigger (Section 2: Quote requires no mandatory referral active).
    if acct["account_id"] in STUDENT_HOUSING_ACCOUNTS or acct.get("named_insured", "").strip() == (
        "Campus Edge Housing Co."
    ):
        triggers.append("refer_appetite")
        rationale_parts.append(
            "Section 10 refer: Campus Edge student-housing account — GRP-APPETITE review "
            "must be submitted before bind per granite_ridge_routing_notes.txt"
        )

    if "decline_lr" in triggers:
        disposition = "Decline"
    elif "ask_broker" in triggers:
        disposition = "Ask Broker"
    elif any(
        t in triggers
        for t in (
            "refer_lr",
            "refer_loc",
            "refer_insp",
            "refer_pool",
            "refer_tiv",
            "refer_units",
            "subsidized",
            "refer_appetite",
        )
    ):
        disposition = "Refer"
    else:
        disposition = "Quote"
        rationale_parts = [quote_rationale(acct["account_id"], losses)]

    rationale = "; ".join(dict.fromkeys(rationale_parts))
    return disposition, rationale, lr, list_lr, triggers


def main() -> None:
    accounts = load_csv("granite_ridge_renewal_accounts.csv")
    locations = load_csv("granite_ridge_location_schedule.csv")
    losses = load_csv("granite_ridge_loss_run_36mo.csv")

    loc_by_acct: dict[str, list[dict[str, str]]] = defaultdict(list)
    for loc in locations:
        loc_by_acct[loc["account_id"]].append(loc)

    loss_by_acct: dict[str, list[dict[str, str]]] = defaultdict(list)
    for loss in losses:
        loss_by_acct[loss["account_id"]].append(loss)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws_src = wb.active
    ws_src.title = "Account_Source"
    ws_src.append(["account_id", "named_insured", "expiring_premium", "list_loss_ratio"])
    for acct in accounts:
        ws_src.append([
            acct["account_id"],
            acct["named_insured"],
            float(acct["expiring_premium"]),
            float(acct["account_loss_ratio_36mo"]),
        ])

    ws_loss = wb.create_sheet("Loss_Run")
    ws_loss.append([
        "loss_id", "account_id", "location_id", "paid_amount", "reserved_amount",
        "cat_code", "total_incurred",
    ])
    loss_row = 2
    for loss in losses:
        ws_loss.append([
            loss["loss_id"],
            loss["account_id"],
            loss["location_id"],
            float(loss["paid_amount"]),
            float(loss["reserved_amount"]),
            loss["cat_code"],
            None,
        ])
        ws_loss[f"G{loss_row}"] = f"=D{loss_row}+E{loss_row}"
        loss_row += 1

    ws = wb.create_sheet("Account_Routing")
    ws.append([
        "account_id", "named_insured", "expiring_premium", "list_loss_ratio",
        "recomputed_loss_ratio", "disposition", "rationale",
    ])

    routing_rows: list[tuple[str, str]] = []
    loss_recon_accounts: list[str] = []
    location_exceptions: list[list] = []
    formula_cache: dict[str, dict[str, float]] = {
        "Account_Routing": {},
        "Loss_Run": {},
        "Loss_Reconciliation": {},
        "Summary": {},
    }

    for idx, acct in enumerate(accounts, start=2):
        acct_id = acct["account_id"]
        locs = loc_by_acct[acct_id]
        acct_losses = loss_by_acct[acct_id]
        disposition, rationale, lr, list_lr, triggers = route_account(acct, locs, acct_losses)
        routing_rows.append((acct_id, disposition))

        ws.append([acct_id, acct["named_insured"], None, None, None, disposition, rationale])
        ws[f"C{idx}"] = f"=Account_Source!C{idx}"
        ws[f"D{idx}"] = f"=Account_Source!D{idx}"
        lr_formula = (
            f'=IF(C{idx}=0,"",ROUND(SUMIFS(Loss_Run!$G$2:$G${LOSS_LAST},'
            f'Loss_Run!$B$2:$B${LOSS_LAST},$A{idx},Loss_Run!$F$2:$F${LOSS_LAST},"")/C{idx}*100,1))'
        )
        ws[f"E{idx}"] = lr_formula
        formula_cache["Account_Routing"][f"C{idx}"] = float(acct["expiring_premium"])
        formula_cache["Account_Routing"][f"D{idx}"] = float(acct["account_loss_ratio_36mo"])
        formula_cache["Account_Routing"][f"E{idx}"] = round(lr, 1)

        if abs(lr - list_lr) > 5:
            loss_recon_accounts.append(acct_id)

        for loc in locs:
            if int(loc["subsidized_unit_count"]) > 0:
                location_exceptions.append([
                    loc["location_id"],
                    acct_id,
                    "Subsidized housing endorsement required",
                    f"{loc['subsidized_unit_count']} subsidized units",
                    "granite_ridge_location_schedule.csv",
                ])
            if inspection_expired(loc["sprinkler_inspection_date"]) or inspection_expired(
                loc["fire_alarm_inspection_date"]
            ):
                location_exceptions.append([
                    loc["location_id"],
                    acct_id,
                    "Expired life-safety inspection",
                    f"sprinkler {loc['sprinkler_inspection_date']}",
                    "granite_ridge_location_schedule.csv",
                ])
            if loc["pool_liability_flag"] == "pending_fence_po":
                location_exceptions.append([
                    loc["location_id"],
                    acct_id,
                    "Pending pool remediation",
                    loc["pool_liability_flag"],
                    "granite_ridge_location_schedule.csv",
                ])
            if loc["vacancy_flag"] == "Y" and loc["broker_confirmation"] == "no":
                location_exceptions.append([
                    loc["location_id"],
                    acct_id,
                    "Vacancy pending broker update",
                    "vacancy_flag=Y",
                    "granite_ridge_location_schedule.csv",
                ])
        if acct_id in STUDENT_HOUSING_ACCOUNTS:
            location_exceptions.append([
                acct_id,
                acct_id,
                "Student housing appetite note",
                STUDENT_HOUSING_NOTE,
                "granite_ridge_routing_notes.txt",
            ])
        if "refer_tiv" in triggers:
            tiv_acct = float(acct["account_tiv"])
            tiv_loc = sum(float(x["building_tiv"]) for x in locs)
            # Signed variance = location sum minus account figure (positive when locations exceed account).
            tiv_var = tiv_loc - tiv_acct
            sign = "+" if tiv_var >= 0 else "-"
            location_exceptions.append([
                acct_id,
                acct_id,
                "Account TIV does not tie to locations",
                f"variance {sign}${abs(tiv_var):,.0f}",
                "granite_ridge_renewal_accounts.csv; granite_ridge_location_schedule.csv",
            ])
        if "refer_units" in triggers:
            units_acct = int(acct["unit_count"])
            units_loc = sum(int(x["unit_count"]) for x in locs)
            location_exceptions.append([
                acct_id,
                acct_id,
                "Account unit count does not tie to locations",
                f"variance {units_loc - units_acct} units",
                "granite_ridge_renewal_accounts.csv; granite_ridge_location_schedule.csv",
            ])

    for loss_idx in range(2, loss_row):
        paid = float(ws_loss[f"D{loss_idx}"].value)
        reserved = float(ws_loss[f"E{loss_idx}"].value)
        formula_cache["Loss_Run"][f"G{loss_idx}"] = round(paid + reserved, 2)

    ws2 = wb.create_sheet("Location_Exceptions")
    ws2.append(["location_or_account", "account_id", "exception_type", "detail", "source"])
    for row in location_exceptions:
        ws2.append(row)

    ws3 = wb.create_sheet("Loss_Reconciliation")
    ws3.append(["account_id", "list_loss_ratio", "recomputed_loss_ratio", "delta", "source"])
    recon_row = 2
    for acct_id in loss_recon_accounts:
        src_row = next(i for i, a in enumerate(accounts, start=2) if a["account_id"] == acct_id)
        ws3.append([acct_id, None, None, None, "granite_ridge_loss_run_36mo.csv vs granite_ridge_renewal_accounts.csv"])
        ws3[f"B{recon_row}"] = f"=Account_Routing!D{src_row}"
        ws3[f"C{recon_row}"] = f"=Account_Routing!E{src_row}"
        ws3[f"D{recon_row}"] = f"=C{recon_row}-B{recon_row}"
        list_lr = float(next(a for a in accounts if a["account_id"] == acct_id)["account_loss_ratio_36mo"])
        recomputed = formula_cache["Account_Routing"][f"E{src_row}"]
        formula_cache["Loss_Reconciliation"][f"B{recon_row}"] = list_lr
        formula_cache["Loss_Reconciliation"][f"C{recon_row}"] = recomputed
        formula_cache["Loss_Reconciliation"][f"D{recon_row}"] = round(recomputed - list_lr, 1)
        recon_row += 1

    ws4 = wb.create_sheet("Rules_Citations")
    ws4.append(["rule_topic", "section", "source_file"])
    citations = [
        ("Loss ratio decline/refer thresholds", "Section 3", "granite_ridge_routing_notes.txt"),
        ("Single-location incurred referral", "Section 3", "granite_ridge_routing_notes.txt"),
        ("Unit count reconciliation tolerance", "Section 4", "granite_ridge_routing_notes.txt"),
        ("TIV reconciliation tolerance", "Section 4", "granite_ridge_routing_notes.txt"),
        ("Life-safety inspection lapse referral", "Section 5", "granite_ridge_routing_notes.txt"),
        ("Pool remediation referral", "Section 5", "granite_ridge_routing_notes.txt"),
        ("Subsidized housing endorsement requirement", "Section 6", "granite_ridge_routing_notes.txt"),
        ("Vacancy broker confirmation", "Section 6", "granite_ridge_routing_notes.txt"),
        ("Cat loss exclusion from attritional LR", "Section 9", "granite_ridge_routing_notes.txt"),
        ("Student housing appetite note — submit GRP-APPETITE before bind", "Section 10", "granite_ridge_routing_notes.txt"),
    ]
    for row in citations:
        ws4.append(row)

    ws5 = wb.create_sheet("Summary")
    ws5["A1"] = "October renewal batch disposition summary"
    ws5["A2"] = "Accounts reviewed"
    ws5["B2"] = f"=COUNTA(Account_Routing!A2:A{ACCOUNT_LAST})"
    ws5["A3"] = "Quote"
    ws5["B3"] = '=COUNTIF(Account_Routing!F2:F100,"Quote")'
    ws5["A4"] = "Refer"
    ws5["B4"] = '=COUNTIF(Account_Routing!F2:F100,"Refer")'
    ws5["A5"] = "Decline"
    ws5["B5"] = '=COUNTIF(Account_Routing!F2:F100,"Decline")'
    ws5["A6"] = "Ask Broker"
    ws5["B6"] = '=COUNTIF(Account_Routing!F2:F100,"Ask Broker")'
    ws5["A8"] = "Subsidized locations flagged"
    ws5["B8"] = '=COUNTIF(Location_Exceptions!C2:C200,"Subsidized housing endorsement required")'

    counts = Counter(d for _, d in routing_rows)
    subsidized_count = sum(
        1 for row in location_exceptions if row[2] == "Subsidized housing endorsement required"
    )
    formula_cache["Summary"] = {
        "B2": len(accounts),
        "B3": counts["Quote"],
        "B4": counts["Refer"],
        "B5": counts["Decline"],
        "B6": counts["Ask Broker"],
        "B8": subsidized_count,
    }

    format_workbook(
        wb,
        n_accounts=len(accounts),
        n_losses=len(losses),
        n_exceptions=len(location_exceptions),
        n_recon=len(loss_recon_accounts),
        n_citations=len(citations),
    )

    wb.save(OUT)
    cache_xlsx_formula_values(OUT, formula_cache)
    sanitize_xlsx(OUT)

    print(f"Wrote {OUT}")
    print(f"Accounts: {len(accounts)}")
    print("Dispositions:", counts)
    for acct_id in [
        "GR-2007",
        "GR-2018",
        "GR-2033",
        "GR-2047",
        "GR-2065",
        "GR-2071",
        "GR-2074",
        "GR-2089",
        "GR-2113",
        "GR-2162",
        "GR-2195",
    ]:
        match = next((d for a, d in routing_rows if a == acct_id), None)
        lr = formula_cache["Account_Routing"].get(
            f"E{next(i for i, a in enumerate(accounts, start=2) if a['account_id'] == acct_id)}"
        )
        print(f"  {acct_id}: {match} (LR {lr})")


if __name__ == "__main__":
    main()
