#!/usr/bin/env python3
"""Generate source files for the Ironclad Telematics valuation task."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, numbers
from openpyxl.utils import get_column_letter

from ironclad_constants import (
    BRIEF_RUN_RATE,
    COUNSEL_POOL_PCT,
    FD_PRE,
    LEAD_INVESTOR,
    NOTE_A_CAP,
    NOTE_A_DISCOUNT,
    NOTE_A_HOLDER,
    NOTE_A_PRINCIPAL,
    NOTE_B_CAP,
    NOTE_B_HOLDER,
    NOTE_B_PRINCIPAL,
    PILOT_ARR,
    PILOT_CUSTOMER,
    POOL_TARGET_PCT,
    POST_MONEY,
    PRE_MONEY,
    PRIMARY,
    ROLLUP_2024,
    SEED_INVESTOR,
    SEED_PRICE,
    SEED_SHARES,
)
from sanitize_office import cache_xlsx_formula_values, sanitize_xlsx

TASK = Path(__file__).resolve().parent.parent / "tasks" / "ironclad-telematics-valuation" / "inputs"
PERIODS = [
    ("FY24-M01", "2023-Q2"),
    ("FY24-M04", "2023-Q3"),
    ("FY24-M07", "2023-Q4"),
    ("FY24-M10", "2024-Q1"),
    ("FY25-M01", "2024-Q2"),
    ("FY25-M04", "2024-Q3"),
    ("FY25-M07", "2024-Q4"),
]
# Gross retention is cohort-level observed (not a graded driver). Keep fleet-SaaS
# range with modest reversals — avoid a clean monotonic ladder across vintages.
VINTAGES = [
    ("IC-2019-01", 430_000, 44, 0.931, 31_000, 17_500),
    ("IC-2020-02", 510_000, 53, 0.924, 38_000, 19_700),
    ("IC-2020-04", 620_000, 61, 0.938, 47_000, 22_300),
    ("IC-2021-01", 735_000, 74, 0.947, 55_000, 26_100),
    ("IC-2021-03", 865_000, 88, 0.941, 69_000, 31_800),
    ("IC-2022-01", 1_020_000, 101, 0.953, 82_000, 35_600),
    ("IC-2022-04", 1_180_000, 116, 0.949, 96_000, 39_400),
    ("IC-2023-02", 1_345_000, 132, 0.958, 112_000, 43_900),
    ("IC-2023-04", 1_510_000, 148, 0.955, 129_000, 48_200),
    ("IC-MUNI-24", PILOT_ARR, 286, 1.0, 0, 0),
]

# Irregular QoQ operating shapes (not constant increments). Indexed by cohort family.
# Values are multipliers vs the vintage's base expansion/churn (period 0 = 1.0).
_START_SHAPES = (
    (1.000, 1.052, 1.097, 1.168, 1.209, 1.286, 1.331),
    (1.000, 1.038, 1.109, 1.151, 1.237, 1.271, 1.342),
    (1.000, 1.061, 1.104, 1.147, 1.224, 1.298, 1.348),
)
_EXP_SHAPES = (
    (1.000, 1.031, 1.088, 1.112, 1.171, 1.198, 1.262),
    (1.000, 1.046, 1.069, 1.141, 1.158, 1.234, 1.251),
    (1.000, 1.019, 1.084, 1.101, 1.186, 1.209, 1.268),
)
_CHURN_SHAPES = (
    (1.000, 1.022, 1.068, 1.091, 1.139, 1.152, 1.211),
    (1.000, 1.048, 1.061, 1.118, 1.134, 1.191, 1.208),
    (1.000, 1.015, 1.074, 1.096, 1.128, 1.185, 1.229),
)
# Vehicle deltas vs opening count — batches, flat quarters, modest dips.
_VEH_DELTAS = (
    (0, 5, 2, 6, 1, 4, 3),
    (0, 3, 7, 2, 5, 0, 4),
    (0, 4, 1, 3, 6, 2, 5),
    (0, 2, 5, 1, 4, -1, 3),
)

# Cap-table employee common and advisor option holders (plausible private-company names).
EMPLOYEE_COMMON = [
    ("Priya Nandakumar", 250_000),
    ("James Okafor", 180_000),
    ("Elena Varga", 145_000),
    ("Marcus Chen", 125_000),
    ("Sofia Lindqvist", 110_000),
    ("Derek Alvarez", 95_000),
    ("Hannah Whitfield", 85_000),
    ("Rajiv Mehta", 70_000),
]
ADVISOR_GRANTS = [
    ("Naomi Kessler", 75_000, 1.05),
    ("Victor Lang", 60_000, 1.18),
    ("Aisha Rahman", 55_000, 1.26),
    ("Peter Holtz", 50_000, 1.41),
    ("Camille Fortin", 45_000, 1.52),
    ("Owen Bristow", 40_000, 1.59),
]
# Historical grant strikes — earlier vintages lower, later higher, uneven gaps.
OPTION_TRANCHES = [
    (650_000, 1.31),
    (420_000, 1.48),
    (310_000, 1.67),
    (180_000, 1.81),
    (95_000, 2.03),
]

ROLLUP_VEHICLES = [1642, 1789, 1948, 2117]
ROLLUP_NRR = [107.4, 108.6, 108.2, 109.5]


def write_cohort_xlsx(path: Path) -> None:
    wb = Workbook()
    bold = Font(bold=True)
    detail = wb.active
    detail.title = "Cohort Detail"
    detail.append(
        [
            "cohort_id",
            "fiscal_month_label",
            "calendar_quarter",
            "starting_arr",
            "vehicle_count",
            "gross_retention_pct",
            "expansion_arr",
            "churn_arr",
            "ending_arr",
            "revenue_tag",
            "notes",
        ]
    )
    for cell in detail[1]:
        cell.font = bold

    duplicate_source: list[object] | None = None
    detail_row_by_key: dict[tuple[str, str], int] = {}
    for cohort_idx, (cohort_id, base, vehicles, retention, expansion, churn) in enumerate(VINTAGES):
        retention_pct = round(retention * 100, 1)
        shape = cohort_idx % len(_START_SHAPES)
        start_shape = _START_SHAPES[shape]
        exp_shape = _EXP_SHAPES[shape]
        churn_shape = _CHURN_SHAPES[shape]
        veh_deltas = _VEH_DELTAS[cohort_idx % len(_VEH_DELTAS)]
        # Small cohort-specific level shift so families do not share identical series.
        level_bump = 1.0 + 0.004 * (cohort_idx % 5) - 0.003 * ((cohort_idx + 2) % 3)
        prior_ending: float | None = None
        for period_idx, (fiscal, quarter) in enumerate(PERIODS):
            if prior_ending is None:
                start = round(base * start_shape[period_idx] * level_bump, 2)
            else:
                # Opening ARR near prior ending with irregular booking lag / early renewals.
                lag = (-0.012, 0.008, -0.004, 0.011, -0.007, 0.005, 0.003)[period_idx]
                start = round(prior_ending * (1.0 + lag), 2)
            if cohort_id == "IC-MUNI-24":
                # Pilot: award-year base with modest mid-pilot vehicle/rate movement only.
                start = round(base * (1.0 + 0.018 * period_idx + (0.006 if period_idx in (2, 5) else -0.002 * (period_idx % 3))), 2)
                exp = 0.0
                churn_amt = 0.0
                veh = vehicles + (0, 12, 8, 21, 15, 9, 18)[period_idx]
            else:
                exp = round(expansion * exp_shape[period_idx] * level_bump, 2)
                churn_amt = round(churn * churn_shape[period_idx] * level_bump, 2)
                veh = max(1, vehicles + veh_deltas[period_idx] + (cohort_idx % 3) - 1)
            tag = "pilot_non_recurring" if cohort_id == "IC-MUNI-24" else "recurring_fleet"
            notes = ""
            if cohort_id == "IC-MUNI-24":
                notes = (
                    f"{PILOT_CUSTOMER} 12-month deployment; procurement has not approved renewal. "
                    f"Exclude from run-rate under sponsor brief Footnote 1 using the brief's "
                    f"${PILOT_ARR:,.0f} award-year annualized exclusion (not this row's ending_arr). "
                    "Current-period ending_arr reflects mid-pilot vehicle adds and rate schedule."
                )
            ending = round(start * retention + exp - churn_amt, 2)
            prior_ending = ending
            row = [
                cohort_id,
                fiscal,
                quarter,
                start,
                veh,
                retention_pct,
                exp,
                churn_amt,
                ending,
                tag,
                notes,
            ]
            detail.append(row)
            excel_row = detail.max_row
            detail_row_by_key[(cohort_id, quarter)] = excel_row
            # Force clean percentage display (avoid binary float serialization noise).
            detail.cell(excel_row, 6).number_format = "0.0"
            detail.cell(excel_row, 6).value = float(f"{retention_pct:.1f}")
            if cohort_id == "IC-2022-04" and quarter == "2024-Q3":
                duplicate_source = row

    assert duplicate_source is not None
    duplicate = duplicate_source.copy()
    duplicate[0] = "IC-2023-02-XP"
    duplicate[10] = (
        "Migration crosswalk row; expansion dollars duplicate IC-2022-04 for the same fleet amendment."
    )
    detail.append(duplicate)
    detail_row_by_key[("IC-2023-02-XP", "2024-Q3")] = detail.max_row
    last_detail_row = detail.max_row

    # Recognition Bridge: defines how Fleet Rollup.fleet_arr relates to Cohort Detail.
    bridge = wb.create_sheet("Recognition Bridge")
    bridge.append(
        [
            "calendar_quarter",
            "gross_cohort_ending_arr",
            "less_duplicate_crosswalk",
            "less_recognition_timing_adj",
            "fleet_arr",
            "cohort_vehicle_count",
            "plus_unassigned_onboarding_vehicles",
            "active_vehicles",
            "bridge_check",
            "methodology_note",
        ]
    )
    for cell in bridge[1]:
        cell.font = bold

    quarters_2024 = list(ROLLUP_2024.keys())
    for i, quarter in enumerate(quarters_2024):
        row_num = i + 2
        fleet_arr = ROLLUP_2024[quarter]
        active_vehicles = ROLLUP_VEHICLES[i]
        # Gross cohort ending ARR and vehicle counts via SUMIF from Cohort Detail.
        bridge.cell(row_num, 1, quarter)
        bridge.cell(
            row_num,
            2,
            f"=SUMIF('Cohort Detail'!C$2:C${last_detail_row},A{row_num},"
            f"'Cohort Detail'!I$2:I${last_detail_row})",
        )
        bridge.cell(
            row_num,
            3,
            f"=SUMIFS('Cohort Detail'!I$2:I${last_detail_row},"
            f"'Cohort Detail'!C$2:C${last_detail_row},A{row_num},"
            f"'Cohort Detail'!A$2:A${last_detail_row},\"IC-2023-02-XP\")",
        )
        # recognition_timing_adj is the reconciling plug that makes the bridge identity hold:
        # fleet_arr = gross_cohort_ending_arr - duplicate - recognition_timing_adj
        # Stored as a formula once fleet_arr is known (column E).
        bridge.cell(row_num, 4, f"=B{row_num}-C{row_num}-E{row_num}")
        bridge.cell(row_num, 5, fleet_arr)
        bridge.cell(
            row_num,
            6,
            f"=SUMIF('Cohort Detail'!C$2:C${last_detail_row},A{row_num},"
            f"'Cohort Detail'!E$2:E${last_detail_row})",
        )
        bridge.cell(row_num, 7, f"=H{row_num}-F{row_num}")
        bridge.cell(row_num, 8, active_vehicles)
        bridge.cell(row_num, 9, f'=IF(ABS(B{row_num}-C{row_num}-D{row_num}-E{row_num})<0.01,"ties","BREAK")')
        bridge.cell(
            row_num,
            10,
            (
                "fleet_arr is finance-recognized subscription ARR for board TTM. "
                "Cohort Detail ending_arr is gross contracted ARR by billing cohort. "
                "recognition_timing_adj removes implementation deferrals and multi-year "
                "contract billings not yet in the subscription ledger. "
                "active_vehicles includes onboarding units not yet tagged to a cohort_id."
            ),
        )
        for col in (2, 3, 4, 5):
            bridge.cell(row_num, col).number_format = numbers.FORMAT_NUMBER_COMMA_SEPARATED1

    # Pre-compute bridge cache values (openpyxl does not evaluate formulas).
    # Collect ending_arr / vehicles by quarter from the detail sheet we just wrote.
    from collections import defaultdict

    gross_by_q: dict[str, float] = defaultdict(float)
    veh_by_q: dict[str, float] = defaultdict(float)
    dup_by_q: dict[str, float] = defaultdict(float)
    for r in range(2, last_detail_row + 1):
        q = detail.cell(r, 3).value
        cid = detail.cell(r, 1).value
        end = float(detail.cell(r, 9).value)
        veh = float(detail.cell(r, 5).value)
        gross_by_q[str(q)] += end
        veh_by_q[str(q)] += veh
        if cid == "IC-2023-02-XP":
            dup_by_q[str(q)] += end

    bridge_cache: dict[str, float] = {}
    for i, quarter in enumerate(quarters_2024):
        row_num = i + 2
        fleet_arr = float(ROLLUP_2024[quarter])
        gross = round(gross_by_q[quarter], 2)
        dup = round(dup_by_q[quarter], 2)
        timing = round(gross - dup - fleet_arr, 2)
        veh = float(veh_by_q[quarter])
        active = float(ROLLUP_VEHICLES[i])
        unassigned = round(active - veh, 2)
        bridge_cache[f"B{row_num}"] = gross
        bridge_cache[f"C{row_num}"] = dup
        bridge_cache[f"D{row_num}"] = timing
        bridge_cache[f"F{row_num}"] = veh
        bridge_cache[f"G{row_num}"] = unassigned

    rollup = wb.create_sheet("Fleet Rollup")
    rollup.append(
        [
            "calendar_quarter",
            "fleet_arr",
            "active_vehicles",
            "reported_nrr_pct",
            "source_note",
        ]
    )
    for cell in rollup[1]:
        cell.font = bold
    for i, quarter in enumerate(quarters_2024):
        row_num = i + 2
        rollup.cell(row_num, 1, quarter)
        # Literal finance-export figures (board TTM source). Bridge defines the identity.
        rollup.cell(row_num, 2, ROLLUP_2024[quarter])
        rollup.cell(row_num, 3, ROLLUP_VEHICLES[i])
        rollup.cell(row_num, 4, ROLLUP_NRR[i])
        rollup.cell(
            row_num,
            5,
            (
                "Finance calendar-quarter rollup (finance-recognized ARR). "
                "See Recognition Bridge + Definitions for Cohort Detail → fleet_arr identity. "
                "Includes municipal pilot in recognized ARR; exclude per brief Footnote 1."
            ),
        )
        rollup.cell(row_num, 2).number_format = numbers.FORMAT_NUMBER_COMMA_SEPARATED1
        rollup.cell(row_num, 4).number_format = "0.0"

    defs = wb.create_sheet("Definitions")
    defs.append(["field", "definition"])
    for cell in defs[1]:
        cell.font = bold
    definitions = [
        ("cohort_id", "Fleet cohort identifier from the billing migration crosswalk."),
        ("fiscal_month_label", "Ironclad fiscal-month label; FY25-M07 maps to calendar 2024-Q4."),
        ("calendar_quarter", "Calendar period used for sponsor valuation analysis."),
        ("starting_arr", "Annualized recurring subscription value at period opening."),
        ("vehicle_count", "Connected billable vehicles at period end for the cohort."),
        ("gross_retention_pct", "Opening ARR retained before expansion (one decimal place)."),
        ("expansion_arr", "Annualized in-cohort seat, module, and vehicle expansion."),
        ("churn_arr", "Annualized contraction and cancellation value."),
        (
            "ending_arr",
            "Gross contracted ARR: starting_arr × gross_retention_pct/100 + expansion_arr − churn_arr. "
            "This is the billing-cohort view, not the finance-recognized board rollup.",
        ),
        (
            "fleet_arr",
            "Finance-recognized subscription ARR for the calendar quarter used in board TTM. "
            "Identity (Recognition Bridge): fleet_arr = SUMIF(Cohort Detail ending_arr by calendar_quarter) "
            "− duplicate IC-2023-02-XP ending_arr − recognition_timing_adj. "
            "Not equal to the raw sum of Cohort Detail ending_arr.",
        ),
        (
            "recognition_timing_adj",
            "Reconciling amount on Recognition Bridge: implementation deferrals and multi-year "
            "contract billings present in Cohort Detail contracted ARR but not yet in the "
            "subscription ledger used for Fleet Rollup.",
        ),
        (
            "active_vehicles",
            "Fleet Rollup vehicle count = SUMIF(Cohort Detail vehicle_count) + unassigned "
            "onboarding vehicles not yet tagged to a billing cohort_id (Recognition Bridge).",
        ),
        (
            "pilot_non_recurring",
            "Time-limited municipal deployment. Model exclusion uses growth_equity_diligence_brief.txt "
            f"Footnote 1 amount (${PILOT_ARR:,.0f} award-year annualized), which may differ from "
            "IC-MUNI-24 Cohort Detail ending_arr for the same quarters.",
        ),
        (
            "duplicate crosswalk",
            "IC-2023-02-XP repeats an IC-2022-04 fleet amendment and must not be counted twice.",
        ),
    ]
    for row in definitions:
        defs.append(row)

    for ws in (detail, bridge, rollup, defs):
        for column in ws.columns:
            letter = get_column_letter(column[0].column)
            max_len = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
            ws.column_dimensions[letter].width = max(12, max_len)

    wb.save(path)
    cache_xlsx_formula_values(path, {"Recognition Bridge": bridge_cache})
    sanitize_xlsx(path)


def cap_table_rows() -> list[list[object]]:
    rows: list[list[object]] = [
        ["Maya Desai", "Common", 5_100_000, "", "", "issued", "Co-founder and CEO"],
        ["Thomas Becker", "Common", 4_200_000, "", "", "issued", "Co-founder and CTO"],
        [SEED_INVESTOR, "Series Seed Preferred", SEED_SHARES, SEED_PRICE, 10_810_000, "issued", "2021 financing"],
        [
            "2021 Equity Plan (available)",
            "Options",
            1_250_000,
            "",
            "",
            "pool",
            f"COUNSEL NOTE: {COUNSEL_POOL_PCT:.0%} available pool shown in working capitalization extract.",
        ],
    ]
    for i, (shares, price) in enumerate(OPTION_TRANCHES, 1):
        rows.append(
            [
                f"Employee option tranche {i}",
                "Options",
                shares,
                price,
                "",
                "unvested",
                "2021/2023 grant vintages",
            ]
        )
    for name, shares in EMPLOYEE_COMMON:
        rows.append([name, "Common", shares, "", "", "issued", "Exercise ledger"])
    rows.append(
        [
            "Mobility channel partner",
            "Warrant",
            160_000,
            1.40,
            "",
            "partially exercised",
            "Original 240,000; 80,000 exercised in 2024; 160,000 remain outstanding.",
        ]
    )
    for name, shares, price in ADVISOR_GRANTS:
        rows.append([name, "Options", shares, price, "", "vested", "Board-approved advisor grant"])
    rows.append(["Employee stock purchase clearing", "Common", 50_000, "", "", "issued", "Q4 clearing balance"])
    rows.extend(
        [
            [
                NOTE_A_HOLDER,
                "Convertible Note A",
                0,
                "",
                NOTE_A_PRINCIPAL,
                "outstanding",
                f"${NOTE_A_CAP:,.0f} valuation cap; {NOTE_A_DISCOUNT:.0%} discount; converts before Note B per sponsor preference.",
            ],
            [
                NOTE_B_HOLDER,
                "Convertible Note B",
                0,
                "",
                NOTE_B_PRINCIPAL,
                "outstanding",
                f"${NOTE_B_CAP:,.0f} valuation cap; no discount.",
            ],
            [SEED_INVESTOR, "Growth round pro-rata (planned)", 0, "", "", "planned", "Sponsor permits full pro-rata participation."],
            [LEAD_INVESTOR, "Growth round lead (planned)", 0, "", "", "planned", "Remainder of primary financing."],
            [
                "Ironclad Telematics FD reconciliation",
                "Reference",
                FD_PRE,
                "",
                "",
                "subtotal",
                "Issued, options, and outstanding warrants above; excludes notes and planned round.",
            ],
        ]
    )
    equity_types = {"Common", "Series Seed Preferred", "Options", "Warrant"}
    assert sum(int(r[2]) for r in rows if r[1] in equity_types) == FD_PRE
    assert len(rows) == 30
    return rows


def write_cap_table(path: Path) -> None:
    header = ["holder", "security_type", "shares", "price", "investment_usd", "vesting_status", "notes"]
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        for row in cap_table_rows():
            out = list(row)
            # Normalize option/warrant prices to two decimal places (no IEEE tails in CSV text).
            if out[3] != "" and out[3] is not None:
                out[3] = f"{float(out[3]):.2f}"
            writer.writerow(out)


def write_brief(path: Path) -> None:
    seed_primary = PRIMARY * SEED_SHARES / FD_PRE
    lead_primary = PRIMARY - seed_primary
    text = f"""IRONCLAD TELEMATICS — GROWTH EQUITY DILIGENCE BRIEF
Prepared for the Board Finance Committee
Draft date: January 16, 2025
Lead sponsor: {LEAD_INVESTOR}

Transaction overview
{LEAD_INVESTOR} has proposed a ${PRIMARY:,.0f} primary growth-equity financing at a ${PRE_MONEY:,.0f} pre-money valuation. The stated post-money value before option-pool mechanics is ${POST_MONEY:,.0f}. Ironclad Telematics sells connected-vehicle workflow and compliance subscriptions to trucking fleets, service contractors, and municipal transit agencies. The round is intended to fund national account coverage, carrier integrations, and implementation capacity. The committee needs one defensible workbook rather than separate revenue and capitalization schedules.

The sponsor term sheet requests a {POOL_TARGET_PCT:.0%} fully diluted post-money option pool at closing. Counsel's capitalization extract still labels the available plan reserve as {COUNSEL_POOL_PCT:.0%}; that label is not a board approval of the closing target. Model the {POOL_TARGET_PCT:.0%} sponsor case as the primary pro forma and show the {COUNSEL_POOL_PCT:.0%} counsel reference as an alternate case. Identify the incremental shares required in each case and do not assume that the existing reserve already satisfies the sponsor target.

Primary allocation
Quorum Ventures may exercise its full pro-rata right based on {SEED_SHARES:,} Series Seed shares out of {FD_PRE:,} fully diluted pre-money shares. At that participation level, {SEED_INVESTOR} would purchase ${seed_primary:,.2f} of the primary and {LEAD_INVESTOR} would purchase the remaining ${lead_primary:,.2f}. The allocation is a modeling convention for the board packet and remains subject to executed subscription documents.

ARR bridge
Management's operating update describes a ${BRIEF_RUN_RATE:,.0f} ARR run-rate. That headline includes contracted implementation uplift, a time-limited municipal deployment, and January activations that are outside the calendar-2024 cohort rollup. The finance export is lower. Use the sum of calendar 2024 quarters on the Fleet Rollup tab as the TTM reference, then apply the exclusion in Footnote 1. Do not force the cohort file to equal management's headline. Fleet Rollup fleet_arr is finance-recognized ARR; Cohort Detail ending_arr is gross contracted ARR by billing cohort. The Recognition Bridge tab states the identity that ties the two measures.

Fiscal labels in Cohort Detail follow Ironclad's internal reporting calendar. FY25-M01, FY25-M04, and FY25-M07 are not calendar 2025 quarters; use the explicit calendar_quarter field when determining the trailing period. One fleet amendment was loaded under both IC-2022-04 and IC-2023-02-XP during the billing migration. The repeated expansion dollars should appear only once in retention and expansion analysis.

Footnote 1 — run-rate convention
The {PILOT_CUSTOMER} deployment is tagged pilot_non_recurring in ironclad_fleet_cohort_arr.xlsx. Its ${PILOT_ARR:,.0f} annualized value is the sponsor-approved award-year exclusion for valuation (one-year municipal pilot with no approved renewal). Cohort Detail row IC-MUNI-24 may show a higher current-period ending_arr after mid-pilot vehicle adds and rate schedule changes; that schedule difference is a documentation item, not a substitute for this footnote. Exclude ${PILOT_ARR:,.0f} from model ARR and from the revenue-multiple base. Retain the reported Fleet Rollup TTM as a visible bridge item so reviewers can reconcile the source.

Footnote 2 — period convention
For trend, growth, and retention calculations, calendar_quarter governs. Fiscal-month labels are operational labels only. Exclude the duplicated IC-2023-02-XP expansion record identified in Cohort Detail. No other cohort rows should be removed solely because the fiscal label appears to fall in another year.

Convertible notes
Two notes remain outstanding and must convert in the proposed financing. Note A is held by {NOTE_A_HOLDER}: ${NOTE_A_PRINCIPAL:,.0f} principal, a ${NOTE_A_CAP:,.0f} valuation cap, and a {NOTE_A_DISCOUNT:.0%} discount to the growth-round price. Use the more favorable of the cap price and discounted round price. Note B is held by {NOTE_B_HOLDER}: ${NOTE_B_PRINCIPAL:,.0f} principal, a ${NOTE_B_CAP:,.0f} cap, and no discount. Calculate each instrument against the fully diluted pre-money denominator stated in ironclad_cap_table.csv.

The sponsor prefers Note A to be displayed before Note B in the conversion schedule. That preference is presentation and closing-sequence guidance; it does not permit a changing denominator between the two contractual conversion calculations. Show both price tests for Note A and the cap-only price for Note B. Both notes must be in post-money ownership.

Forecast and valuation conventions
Build a five-year forecast from normalized model ARR with annual growth of 26%, 22%, 18%, 15%, and 12%. Use a 20% free-cash-flow margin as a steady-state underwriting proxy. The DCF cross-check should use a 13% to 15% discount-rate range and 3% terminal growth. These are valuation scenarios, not management guidance.

For the market cross-check, apply 4.5x to 6.0x to year-one NTM revenue. Use 4.0x on normalized TTM only for the downside case. The peer median currently sits near 5.2x NTM revenue. Keep multiple and DCF conclusions separate enough that reviewers can see the effect of growth assumptions instead of counting the same optimism twice.

Requested board output
Provide one Excel workbook tying fleet cohort retention and expansion to a formula-driven forecast, a DCF-style cross-check, market multiples, note conversion, primary issuance, and option-pool dilution. Include sensitivity to both fleet growth and net retention. Document the management run-rate difference, municipal pilot exclusion, Fleet Rollup versus Cohort Detail recognition bridge, duplicate expansion treatment, fiscal/calendar mapping, note order, and pool conflict with citations to growth_equity_diligence_brief.txt and the relevant source schedules.

Confidential — Ironclad Telematics board materials.
"""
    assert len(text.split()) >= 500
    path.write_text(text, encoding="utf-8")


def main() -> None:
    TASK.mkdir(parents=True, exist_ok=True)
    write_cohort_xlsx(TASK / "ironclad_fleet_cohort_arr.xlsx")
    write_cap_table(TASK / "ironclad_cap_table.csv")
    write_brief(TASK / "growth_equity_diligence_brief.txt")
    print(f"Wrote inputs to {TASK}")


if __name__ == "__main__":
    main()
