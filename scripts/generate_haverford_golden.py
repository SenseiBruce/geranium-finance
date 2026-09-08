#!/usr/bin/env python3
"""Generate golden flux close pack for Haverford Process Equipment."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from haverford_constants import (
    ACCOUNTS,
    AUG,
    BS_FLUX_ABS,
    CAPEX_REBUILD,
    CLOSE_STATUS,
    DELIVERABLE,
    ENTITY_LEGAL,
    FREIGHT_OPEN,
    INV_PHYSICAL,
    INV_TB,
    INV_WRITE_DOWN,
    JUL,
    MEMO_TXT,
    PERIOD,
    PNL_FLUX_ABS,
    PNL_FLUX_PCT,
    PREPAID_ANNUAL,
    PREPAID_AUG_AMORT,
    PTO_AJE,
    PTO_SUPPORT,
    PTO_TB,
    SLUG,
    SOFT_CLOSE,
    SUPPORT_CSV,
    TB_XLSX,
    VOID_JE_AMT,
    WARRANTY_AJE,
    WARRANTY_DETAIL_ROLL,
    WARRANTY_SOURCE_VARIANCE,
    WARRANTY_SUPPORT,
    WARRANTY_TB,
    flux,
)
from sanitize_office import cache_xlsx_formula_values, sanitize_xlsx

TASK_GOLDEN = Path(__file__).resolve().parent.parent / "tasks" / SLUG / "golden"

body = Font(name="Calibri", size=10)
header_font = Font(name="Calibri", size=10, bold=True)
title_font = Font(name="Calibri", size=12, bold=True)
thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
money = "#,##0.00"
pct = "0.0%"


def autosize(ws, max_w: int = 44) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min(max(len(str(c.value or "")) for c in col) + 2, max_w)
        ws.column_dimensions[letter].width = max(width, 10)


def needs_flux(acct: str, typ: str) -> bool:
    f = abs(flux(acct))
    prior = abs(JUL[acct])
    if typ == "PL":
        if f < PNL_FLUX_ABS:
            return False
        if prior == 0:
            return True
        return (f / prior) >= PNL_FLUX_PCT
    return f >= BS_FLUX_ABS


def main() -> None:
    wb = Workbook()

    # --- Cover ---
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = f"{ENTITY_LEGAL} — Month-end GL flux close pack"
    cover["A1"].font = title_font
    rows = [
        ("Period", PERIOD),
        ("Soft-close target", SOFT_CLOSE),
        ("Governing policy", MEMO_TXT),
        ("TB source", TB_XLSX),
        ("Support source", SUPPORT_CSV),
        ("Close readiness", CLOSE_STATUS),
        ("Hierarchy", "flux_threshold_and_aje_memo.txt governs over Plant_Notes clean-close language"),
    ]
    for i, (k, v) in enumerate(rows, 3):
        cover.cell(i, 1, k).font = header_font
        cover.cell(i, 2, v).font = body
    cover["A11"] = (
        "Prepared for corporate accounting soft close. Proposed AJEs must post before books are labeled Ready. "
        "Warranty AJE uses W-04 governing support; the $6,000 detail-roll vs W-04 gap is an open source exception."
    )
    autosize(cover)

    # --- Flux workpaper ---
    fx = wb.create_sheet("Flux_Workpaper")
    fx["A1"] = f"Material flux investigation — {PERIOD} vs prior"
    fx["A1"].font = title_font
    headers = [
        "Account",
        "Name",
        "Type",
        "July",
        "August",
        "Flux",
        "Flux_Pct_of_July",
        "Threshold_Hit",
        "Explanation",
        "Support_Ref",
    ]
    for c, h in enumerate(headers, 1):
        cell = fx.cell(3, c, h)
        cell.font = header_font
        cell.border = thin

    explanations = {
        "1000": (
            "Cash declined on scheduled AP pay-downs (JE-8871) and ordinary operating disbursements.",
            "JE_Activity_Notes JE-8871",
        ),
        "1100": (
            "AR rose on late-month shipments; INV-9918 billing-void status remains with original JE-8841 amounts still in the GL.",
            "JE-8841 / support S-01",
        ),
        "1200": (
            "Inventory TB vs CNT-AUG physical total (I-04/I-05); memo §4.4 physical-vs-book true-up applies.",
            "inventory I-04 / I-05",
        ),
        "1650": (
            "Accumulated depreciation increased by the monthly depreciation run JE-8860.",
            "JE-8860",
        ),
        "2000": (
            "AP increased on receipt timing vs pay-downs; aging support is informational only.",
            "AP-AGE / JE-8871",
        ),
        "5600": (
            "Maintenance includes CAP-8821 CNC spindle rebuild (useful life 5 years) posted to expense; memo §4.5 CapEx test applies.",
            "capex C-01 / C-02 / JE-8810",
        ),
    }

    flux_rows: list[tuple[str, str, str, float, float]] = []
    for acct, name, typ, _norm in ACCOUNTS:
        if not needs_flux(acct, typ):
            continue
        flux_rows.append((acct, name, typ, JUL[acct], AUG[acct]))

    flux_cache: dict[str, float | int | str] = {}
    r = 4
    for acct, name, typ, july, aug in flux_rows:
        fx.cell(r, 1, acct).font = body
        fx.cell(r, 2, name).font = body
        fx.cell(r, 3, typ).font = body
        fx.cell(r, 4, round(july, 2)).number_format = money
        fx.cell(r, 5, round(aug, 2)).number_format = money
        fx.cell(r, 6, f"=E{r}-D{r}").number_format = money
        fx.cell(r, 7, f"=IF(D{r}=0,1,ABS(F{r})/ABS(D{r}))").number_format = pct
        if typ == "BS":
            fx.cell(r, 8, f'=IF(ABS(F{r})>={BS_FLUX_ABS},"Yes","No")')
        else:
            fx.cell(
                r,
                8,
                f'=IF(AND(ABS(F{r})>={PNL_FLUX_ABS},G{r}>={PNL_FLUX_PCT}),"Yes","No")',
            )
        expl, ref = explanations.get(acct, ("See JE activity and support schedules.", "TB/support"))
        fx.cell(r, 9, expl).font = body
        fx.cell(r, 10, ref).font = body
        for c in range(1, 11):
            fx.cell(r, c).border = thin
            fx.cell(r, c).font = body

        f_amt = round(aug - july, 2)
        pct_amt = 1.0 if july == 0 else abs(f_amt) / abs(july)
        flux_cache[f"F{r}"] = f_amt
        flux_cache[f"G{r}"] = round(pct_amt, 10)
        flux_cache[f"H{r}"] = "Yes"
        r += 1
    fx_last = r - 1
    autosize(fx)

    # --- Support reconciliation (source facts + derived diffs) ---
    # Built before Proposed_AJEs so AJE amounts can formula-link here.
    rec = wb.create_sheet("Support_Reconcile")
    rec["A1"] = "TB vs support reconciliation"
    rec["A1"].font = title_font
    rh = ["Item", "TB_Amount", "Support_Amount", "Difference", "AJE_ID", "Disposition"]
    for c, h in enumerate(rh, 1):
        cell = rec.cell(3, c, h)
        cell.font = header_font
        cell.border = thin

    # Row map (fixed):
    # 4 warranty, 5 prepaid, 6 inventory, 7 freight, 8 PTO, 9 void, 10 CapEx
    recon_source = [
        ("Accrued Warranty 2200", WARRANTY_TB, WARRANTY_SUPPORT, "AJE-02", "True TB to W-04 governing support"),
        ("Prepaid Insurance amort", 0.0, PREPAID_ANNUAL, "AJE-03", "Monthly = annual/12; book missing Aug amort"),
        ("Inventory 1200", INV_TB, INV_PHYSICAL, "AJE-04", "Write down to I-04 physical total"),
        ("Accrued Freight 2210", 0.0, FREIGHT_OPEN, "AJE-05", "Accrue F-04 open total"),
        ("Accrued PTO 2220", PTO_TB, PTO_SUPPORT, "AJE-06", "True TB to T-03"),
        ("Voided INV-9918 in AR/Sales", VOID_JE_AMT, 0.0, "AJE-01", "Reverse billing-void amounts still in GL"),
        ("CAP-8821 in Maintenance", CAPEX_REBUILD, CAPEX_REBUILD, "AJE-07", "Reclass to PP&E 1600 per §4.5"),
    ]
    for i, (item, tb, sup, aje_id, disp) in enumerate(recon_source, 4):
        rec.cell(i, 1, item)
        rec.cell(i, 2, round(tb, 2)).number_format = money
        rec.cell(i, 3, round(sup, 2)).number_format = money
        if i == 5:
            # Prepaid: derived monthly amort from annual support (col C), not a TB-support subtract
            rec.cell(i, 4, f"=ROUND(C{i}/12,2)").number_format = money
        else:
            rec.cell(i, 4, f"=B{i}-C{i}").number_format = money
        rec.cell(i, 5, aje_id)
        rec.cell(i, 6, disp)
        for c in range(1, 7):
            rec.cell(i, c).font = body
            rec.cell(i, c).border = thin

    # Derived AJE amount bridge (live) — Proposed_AJEs links here
    rec["A12"] = "AJE amount bridge (live formulas from rows above)"
    rec["A12"].font = header_font
    bridge = [
        (13, "AJE-01 Void reverse", "=ABS(D9)"),
        (14, "AJE-02 Warranty true-up (W-04 basis)", "=ABS(D4)"),
        (15, "AJE-03 Prepaid amort", "=D5"),
        (16, "AJE-04 Inventory write-down", "=ABS(D6)"),
        (17, "AJE-05 Freight accrual", "=ABS(D7)"),
        (18, "AJE-06 PTO true-up", "=ABS(D8)"),
        (19, "AJE-07 CapEx reclass", "=ABS(C10)"),
    ]
    for row, label, formula in bridge:
        rec.cell(row, 1, label).font = body
        rec.cell(row, 2, formula).number_format = money
        rec.cell(row, 1).border = thin
        rec.cell(row, 2).border = thin

    # Warranty source exception disclosure (Path A — W-04 governing)
    rec["A21"] = "Warranty source exception (open support inconsistency)"
    rec["A21"].font = header_font
    rec["A22"] = "Detail roll W-01+W-02+W-03"
    rec["B22"] = round(WARRANTY_DETAIL_ROLL, 2)
    rec["B22"].number_format = money
    rec["A23"] = "W-04 governing support total"
    rec["B23"] = round(WARRANTY_SUPPORT, 2)
    rec["B23"].number_format = money
    rec["A24"] = "Variance (detail roll − W-04)"
    rec["B24"] = "=B22-B23"
    rec["B24"].number_format = money
    rec["A25"] = "Selected governing figure"
    rec["B25"] = "=B23"
    rec["B25"].number_format = money
    rec["C25"] = "W-04"
    rec["A26"] = "Selection reason"
    rec["B26"] = (
        "W-04 status=Support and notes='Governing support total' establish W-04 as the "
        "controlling support ending under memo §4.2; AJE-02 trues TB 2200 to W-04."
    )
    rec["B26"].alignment = Alignment(wrap_text=True)
    rec.merge_cells("B26:F28")
    rec["A29"] = "Open exception treatment"
    rec["B29"] = (
        f"Detail roll ${WARRANTY_DETAIL_ROLL:,.2f} vs W-04 ${WARRANTY_SUPPORT:,.2f} leaves an "
        f"unexplained ${WARRANTY_SOURCE_VARIANCE:,.2f} source gap. Documented as an open "
        "reconciliation exception — not a second AJE and not silently ignored."
    )
    rec["B29"].alignment = Alignment(wrap_text=True)
    rec.merge_cells("B29:F31")
    rec["A32"] = "Exception amount check (should equal variance)"
    rec["B32"] = f'=IF(ABS(B24-{WARRANTY_SOURCE_VARIANCE})<0.005,"PASS","FAIL")'
    autosize(rec)

    # --- Proposed AJEs (amounts formula-linked to Support_Reconcile bridge) ---
    aje = wb.create_sheet("Proposed_AJEs")
    aje["A1"] = f"Proposed adjusting journal entries — {PERIOD}"
    aje["A1"].font = title_font
    aje_headers = [
        "AJE_ID",
        "Line",
        "Account",
        "Account Name",
        "Debit",
        "Credit",
        "Amount_Check",
        "Rationale",
        "Support_Ref",
    ]
    for c, h in enumerate(aje_headers, 1):
        cell = aje.cell(3, c, h)
        cell.font = header_font
        cell.border = thin

    # (id, line, acct, name, debit_formula, credit_formula, rationale, ref)
    aje_lines = [
        (
            "AJE-01",
            1,
            "4100",
            "Product Sales",
            "=Support_Reconcile!B13",
            "0",
            "Reverse voided INV-9918 revenue still in GL",
            "sales S-01 / JE-8841",
        ),
        (
            "AJE-01",
            2,
            "1100",
            "Accounts Receivable",
            "0",
            "=Support_Reconcile!B13",
            "Reverse voided INV-9918 AR still in GL",
            "sales S-01 / JE-8841",
        ),
        (
            "AJE-02",
            1,
            "2200",
            "Accrued Warranty",
            "=Support_Reconcile!B14",
            "0",
            "True warranty liability to W-04 governing support (see $6k detail-roll exception)",
            "warranty W-04",
        ),
        (
            "AJE-02",
            2,
            "5200",
            "Warranty Expense",
            "0",
            "=Support_Reconcile!B14",
            "Offset warranty true-up to W-04",
            "warranty W-04",
        ),
        (
            "AJE-03",
            1,
            "5300",
            "Insurance Expense",
            "=Support_Reconcile!B15",
            "0",
            "Book August prepaid insurance amortization (annual/12)",
            "prepaid P-01 / memo §4.3",
        ),
        (
            "AJE-03",
            2,
            "1400",
            "Prepaid Insurance",
            "0",
            "=Support_Reconcile!B15",
            "Amortize prepaid insurance",
            "prepaid P-01 / memo §4.3",
        ),
        (
            "AJE-04",
            1,
            "5100",
            "Cost of Goods Sold",
            "=Support_Reconcile!B16",
            "0",
            "Write inventory book down to physical",
            "inventory I-04",
        ),
        (
            "AJE-04",
            2,
            "1200",
            "Inventory - Finished & WIP",
            "0",
            "=Support_Reconcile!B16",
            "Write inventory book down to physical",
            "inventory I-04",
        ),
        (
            "AJE-05",
            1,
            "5400",
            "Freight-In / Outbound",
            "=Support_Reconcile!B17",
            "0",
            "Accrue open freight support items",
            "freight F-04",
        ),
        (
            "AJE-05",
            2,
            "2210",
            "Accrued Freight",
            "0",
            "=Support_Reconcile!B17",
            "Accrue open freight support items",
            "freight F-04",
        ),
        (
            "AJE-06",
            1,
            "2220",
            "Accrued PTO",
            "=Support_Reconcile!B18",
            "0",
            "True PTO liability to HR support",
            "pto T-03",
        ),
        (
            "AJE-06",
            2,
            "5500",
            "Payroll Benefits / PTO",
            "0",
            "=Support_Reconcile!B18",
            "Offset PTO true-up",
            "pto T-03",
        ),
        (
            "AJE-07",
            1,
            "1600",
            "PP&E - Machinery",
            "=Support_Reconcile!B19",
            "0",
            "Capitalize CapEx-qualifying rebuild CAP-8821",
            "capex C-01 / JE-8810",
        ),
        (
            "AJE-07",
            2,
            "5600",
            "Maintenance & Repairs",
            "0",
            "=Support_Reconcile!B19",
            "Remove CapEx ticket from maintenance expense",
            "capex C-01 / JE-8810",
        ),
    ]

    for i, row in enumerate(aje_lines, 4):
        aje_id, line, acct, name, debit_f, credit_f, rationale, ref = row
        aje.cell(i, 1, aje_id)
        aje.cell(i, 2, line)
        aje.cell(i, 3, acct)
        aje.cell(i, 4, name)
        aje.cell(i, 5, debit_f).number_format = money
        aje.cell(i, 6, credit_f).number_format = money
        aje.cell(i, 7, f"=E{i}+F{i}").number_format = money
        aje.cell(i, 8, rationale)
        aje.cell(i, 9, ref)
        for c in range(1, 10):
            aje.cell(i, c).font = body
            aje.cell(i, c).border = thin

    last = 3 + len(aje_lines)
    tot_row = last + 2
    cred_row = last + 3
    bal_row = last + 4
    aje.cell(tot_row, 1, "Total debits").font = header_font
    aje.cell(tot_row, 5, f"=SUM(E4:E{last})").number_format = money
    aje.cell(cred_row, 1, "Total credits").font = header_font
    aje.cell(cred_row, 6, f"=SUM(F4:F{last})").number_format = money
    aje.cell(bal_row, 1, "Out of balance (should be 0)").font = header_font
    aje.cell(bal_row, 5, f"=E{tot_row}-F{cred_row}").number_format = money
    aje.cell(bal_row + 1, 1, "Balance check").font = header_font
    aje.cell(bal_row + 1, 5, f'=IF(E{bal_row}=0,"PASS","FAIL")')
    autosize(aje)

    # --- Close decision ---
    dec = wb.create_sheet("Close_Decision")
    dec["A1"] = "Close readiness decision"
    dec["A1"].font = title_font
    dec["A3"] = "Status"
    dec["B3"] = CLOSE_STATUS
    dec["A3"].font = header_font
    dec["B3"].font = body
    dec["A5"] = "Decision notes"
    dec["A5"].font = header_font
    dec["A6"] = (
        "Plant_Notes assert a clean close with no AJEs. Under flux_threshold_and_aje_memo.txt hierarchy, "
        "that note does not clear mandatory AJE categories. Seven AJEs are proposed (void reverse, warranty "
        "to W-04, prepaid amort, inventory write-down, freight accrual, PTO true-up, CapEx reclass). "
        f"The ${WARRANTY_SOURCE_VARIANCE:,.2f} warranty detail-roll vs W-04 gap is an open source exception "
        f"documented on Support_Reconcile; it does not block quantification of AJE-02 on the W-04 basis. "
        f"Soft close may proceed as {CLOSE_STATUS} once those entries post. Status is not Ready (AJEs remain) "
        "and not Hold close (amounts are fully quantified from TB + governing support)."
    )
    dec["A6"].alignment = Alignment(wrap_text=True)
    dec.merge_cells("A6:B12")
    dec["A14"] = "Path matrix check"
    dec["A14"].font = header_font
    dec["A15"] = "Ready?"
    dec["B15"] = "No — required AJEs outstanding"
    dec["A16"] = "Ready with AJEs?"
    dec["B16"] = "Yes — pack proposes all section 4 AJEs"
    dec["A17"] = "Hold close?"
    dec["B17"] = "No — quantification complete on governing support; $6k warranty source gap disclosed as exception"
    autosize(dec)

    # --- Values_Snapshot: live links only ---
    snap = wb.create_sheet("Values_Snapshot")
    snap["A1"] = "Rounded key amounts — live links to calculation sheets (not re-typed)"
    snap["A1"].font = title_font
    snap_rows = [
        ("VOID_JE_AMT", "=Support_Reconcile!B13"),
        ("WARRANTY_AJE", "=Support_Reconcile!B14"),
        ("PREPAID_AUG_AMORT", "=Support_Reconcile!B15"),
        ("INV_WRITE_DOWN", "=Support_Reconcile!B16"),
        ("FREIGHT_OPEN", "=Support_Reconcile!B17"),
        ("PTO_AJE", "=Support_Reconcile!B18"),
        ("CAPEX_REBUILD", "=Support_Reconcile!B19"),
        ("AJE_TOTAL_DEBITS", f"=Proposed_AJEs!E{tot_row}"),
        ("AJE_OUT_OF_BALANCE", f"=Proposed_AJEs!E{bal_row}"),
        ("WARRANTY_DETAIL_ROLL", "=Support_Reconcile!B22"),
        ("WARRANTY_W04", "=Support_Reconcile!B23"),
        ("WARRANTY_SOURCE_VARIANCE", "=Support_Reconcile!B24"),
        ("CLOSE_STATUS", "=Close_Decision!B3"),
    ]
    for i, (k, v) in enumerate(snap_rows, 3):
        snap.cell(i, 1, k).font = header_font
        snap.cell(i, 2, v)
        if k != "CLOSE_STATUS":
            snap.cell(i, 2).number_format = money
    autosize(snap)

    TASK_GOLDEN.mkdir(parents=True, exist_ok=True)
    out = TASK_GOLDEN / DELIVERABLE
    wb.save(out)
    sanitize_xlsx(out)

    aje_total = round(
        VOID_JE_AMT
        + WARRANTY_AJE
        + PREPAID_AUG_AMORT
        + INV_WRITE_DOWN
        + FREIGHT_OPEN
        + PTO_AJE
        + CAPEX_REBUILD,
        2,
    )
    # Cache formula results for data_only / digest readers
    formula_cache: dict[str, dict[str, float | int | str]] = {
        "Flux_Workpaper": flux_cache,
        "Support_Reconcile": {
            "D4": round(WARRANTY_TB - WARRANTY_SUPPORT, 2),
            "D5": PREPAID_AUG_AMORT,
            "D6": round(INV_TB - INV_PHYSICAL, 2),
            "D7": round(0.0 - FREIGHT_OPEN, 2),
            "D8": round(PTO_TB - PTO_SUPPORT, 2),
            "D9": round(VOID_JE_AMT - 0.0, 2),
            "D10": round(CAPEX_REBUILD - CAPEX_REBUILD, 2),
            "B13": VOID_JE_AMT,
            "B14": WARRANTY_AJE,
            "B15": PREPAID_AUG_AMORT,
            "B16": INV_WRITE_DOWN,
            "B17": FREIGHT_OPEN,
            "B18": PTO_AJE,
            "B19": CAPEX_REBUILD,
            "B24": WARRANTY_SOURCE_VARIANCE,
            "B25": WARRANTY_SUPPORT,
            "B32": "PASS",
        },
        "Proposed_AJEs": {},
        "Values_Snapshot": {
            "B3": VOID_JE_AMT,
            "B4": WARRANTY_AJE,
            "B5": PREPAID_AUG_AMORT,
            "B6": INV_WRITE_DOWN,
            "B7": FREIGHT_OPEN,
            "B8": PTO_AJE,
            "B9": CAPEX_REBUILD,
            "B10": aje_total,
            "B11": 0.0,
            "B12": WARRANTY_DETAIL_ROLL,
            "B13": WARRANTY_SUPPORT,
            "B14": WARRANTY_SOURCE_VARIANCE,
            "B15": CLOSE_STATUS,
        },
    }
    # Proposed_AJEs debit/credit/check/totals
    amounts = [
        (VOID_JE_AMT, 0.0),
        (0.0, VOID_JE_AMT),
        (WARRANTY_AJE, 0.0),
        (0.0, WARRANTY_AJE),
        (PREPAID_AUG_AMORT, 0.0),
        (0.0, PREPAID_AUG_AMORT),
        (INV_WRITE_DOWN, 0.0),
        (0.0, INV_WRITE_DOWN),
        (FREIGHT_OPEN, 0.0),
        (0.0, FREIGHT_OPEN),
        (PTO_AJE, 0.0),
        (0.0, PTO_AJE),
        (CAPEX_REBUILD, 0.0),
        (0.0, CAPEX_REBUILD),
    ]
    for i, (d, c) in enumerate(amounts, 4):
        formula_cache["Proposed_AJEs"][f"E{i}"] = d
        formula_cache["Proposed_AJEs"][f"F{i}"] = c
        formula_cache["Proposed_AJEs"][f"G{i}"] = round(d + c, 2)
    formula_cache["Proposed_AJEs"][f"E{tot_row}"] = aje_total
    formula_cache["Proposed_AJEs"][f"F{cred_row}"] = aje_total
    formula_cache["Proposed_AJEs"][f"E{bal_row}"] = 0.0
    formula_cache["Proposed_AJEs"][f"E{bal_row + 1}"] = "PASS"

    cache_xlsx_formula_values(out, formula_cache)

    # Stable sheet order for reviewers / prior submissions
    order = [
        "Cover",
        "Flux_Workpaper",
        "Proposed_AJEs",
        "Support_Reconcile",
        "Close_Decision",
        "Values_Snapshot",
    ]
    # Re-open after cache inject (XML rewrite) and reorder
    from openpyxl import load_workbook as _lw

    wb2 = _lw(out)
    existing = {ws.title: ws for ws in wb2._sheets}
    wb2._sheets = [existing[name] for name in order]
    wb2.save(out)
    sanitize_xlsx(out)
    # Re-apply cache after sanitize/reorder (sanitize may rewrite XML)
    cache_xlsx_formula_values(out, formula_cache)

    print(f"Wrote {out}")
    print(f"Flux rows: {fx_last - 3}; AJE total debits/credits: {aje_total}")
    print(
        f"Warranty: detail={WARRANTY_DETAIL_ROLL} W04={WARRANTY_SUPPORT} "
        f"var={WARRANTY_SOURCE_VARIANCE} AJE={WARRANTY_AJE}"
    )


if __name__ == "__main__":
    main()
