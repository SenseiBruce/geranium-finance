#!/usr/bin/env python3
"""Generate golden consolidation workbook for Vesper Culinary Brands Q3.

Architecture:
- Line_Detail holds raw ledger rows (typed) plus live inclusion / eligible formulas.
- Brand_Rollup brand grosses are SUMIFS over Line_Detail (not typed constants).
- Assumptions!B13 selects Kiln ownership method: 50/92 proration (base / board) or
  Transaction-date cutoff (conditional alternative only — approval evidence absent).
- Open follow-ups document missing Schedule 2.1 / bank-wire and post-close billing support.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from sanitize_office import cache_xlsx_formula_values, sanitize_xlsx

TASK = Path(__file__).resolve().parent.parent / "tasks" / "vesper-culinary-brands-consolidation"
INPUTS = TASK / "inputs"
OUT = TASK / "golden" / "consolidated_review_vesper_culinary_brands.xlsx"

Q3_START = "2025-07-01"
Q3_END = "2025-09-30"
KILN_CLOSE = "2025-08-12"
EUR_USD = 1.0874
KILN_OWNED_DAYS = 50
Q3_DAYS = 92
METHOD_PRORATION = "50/92 proration"
METHOD_CUTOFF = "Transaction-date cutoff"
# Golden base case remains 50/92; cutoff is modeled as conditional only.
SELECTED_METHOD = METHOD_PRORATION
HEADER_FONT = Font(bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
SUBTITLE_FONT = Font(bold=True, size=11, color="1F4E79")
SECTION_FILL = PatternFill("solid", fgColor="D6E3F0")
USD_FORMAT = '"$"#,##0.00'
EUR_FORMAT = '"€"#,##0.00'
TEXT_USD = "$#,##0.00"
RATE_FORMAT = "0.0000"
WORKBOOK_TITLE = "Vesper Culinary Brands — Consolidated Net Revenue Review"
PERIOD_LABEL = "Q3 2025 (July 1, 2025 – September 30, 2025)"
AS_OF_LABEL = "Analysis as of Q3 2025 close (early October 2025)"
APPROVAL_EVIDENCE = "Approval evidence in supplied packet: not present"
CUTOFF_LABEL = (
    "Conditional transaction-date cutoff scenario — approval evidence not present in packet"
)
OWNERSHIP_MODELING_BASIS = "August 12, 2025 — per acquisition memo"
OWNERSHIP_EVIDENCE_STATUS = "Open - supporting records not provided"
OWNERSHIP_SUPPORTING_EVIDENCE = (
    "Not independently confirmed from Schedule 2.1 or bank wire confirmation; "
    "those records are not in the packet."
)
POST_CLOSE_BILLING_SUPPORT = (
    "Open follow-up - shipment/service evidence not provided"
)
OWNERSHIP_FOLLOWUP_NOTE = (
    "The August 12, 2025 Kiln ownership date is used for the current analysis based on "
    "the acquisition memo, but confirmation against Schedule 2.1 and the bank wire "
    "confirmation remains outstanding because those records are not included in the "
    "supplied packet."
)
BILLING_FOLLOWUP_NOTE = (
    "Post-close Kiln billings should be tied to shipment or service evidence where "
    "practical; the underlying supporting records are not included in the supplied "
    "packet, so this remains an open follow-up item."
)


def write_title_block(ws: Worksheet, sheet_subtitle: str, *, cols: int = 2) -> int:
    """Write controller-ready title rows; returns first data row index."""
    ws["A1"] = WORKBOOK_TITLE
    ws["A1"].font = TITLE_FONT
    ws["A2"] = PERIOD_LABEL
    ws["A2"].font = SUBTITLE_FONT
    ws["A3"] = f"Sheet: {sheet_subtitle} | {AS_OF_LABEL}"
    ws["A3"].font = Font(italic=True, size=10, color="595959")
    for col in range(1, cols + 1):
        ws.cell(1, col).fill = SECTION_FILL
        ws.cell(2, col).fill = SECTION_FILL
    return 5


def fmt_usd(ws: Worksheet, *coords: str) -> None:
    for coord in coords:
        ws[coord].number_format = USD_FORMAT


def fmt_eur(ws: Worksheet, *coords: str) -> None:
    for coord in coords:
        ws[coord].number_format = EUR_FORMAT


def style_table(ws: Worksheet, header_row: int, last_row: int, last_col: int, widths: list[float]) -> None:
    """Freeze header, autofilter, and set readable column widths."""
    ws.freeze_panes = f"A{header_row + 1}"
    if last_row >= header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{last_row}"
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


EXCLUDE_STATUSES = {"void", "cancelled", "duplicate", "out_of_period"}


def load_ledger() -> list[dict]:
    with (INPUTS / "vesper_culinary_brands_brand_ledger.csv").open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def classify_status_period(row: dict) -> tuple[str, bool]:
    """Status + period inclusion only (ownership handled by method selector)."""
    status = row["status"].strip().lower()
    date = row["transaction_date"]
    if status in EXCLUDE_STATUSES:
        return f"status_{status}", False
    if date < Q3_START:
        return "prior_period", False
    if date > Q3_END:
        return "post_period", False
    return "", True


def round2(x: float) -> float:
    return round(x, 2)


def main() -> None:
    rows = load_ledger()
    formula_cache: dict[str, dict[str, float | str]] = {}

    wb = Workbook()

    # --- Assumptions (central parameters) ---
    wa = wb.active
    wa.title = "Assumptions"
    r0 = write_title_block(wa, "Assumptions", cols=2)

    wa.cell(r0, 1, "Item").font = HEADER_FONT
    wa.cell(r0, 2, "Value").font = HEADER_FONT

    wa.cell(r0 + 1, 1, "Q3 period start")
    wa.cell(r0 + 1, 2, Q3_START)
    wa.cell(r0 + 2, 1, "Q3 period end")
    wa.cell(r0 + 2, 2, Q3_END)
    wa.cell(r0 + 3, 1, "EUR/USD board close rate")
    wa.cell(r0 + 3, 2, EUR_USD)
    wa[f"B{r0 + 3}"].number_format = RATE_FORMAT
    wa.cell(r0 + 4, 1, "Kiln Spice Works ownership start (modeling date)")
    wa.cell(r0 + 4, 2, KILN_CLOSE)
    wa.cell(r0 + 5, 1, "Kiln ownership days in Q3 (numerator)")
    wa.cell(r0 + 5, 2, KILN_OWNED_DAYS)
    wa.cell(r0 + 6, 1, "Q3 calendar days (denominator)")
    wa.cell(r0 + 6, 2, Q3_DAYS)
    wa.cell(r0 + 7, 1, "Kiln proration factor (50/92)")
    wa.cell(r0 + 7, 2, f"=B{r0 + 5}/B{r0 + 6}")
    wa.cell(r0 + 8, 1, "Selected final Kiln ownership method (board)")
    wa.cell(r0 + 8, 2, SELECTED_METHOD)
    wa.cell(r0 + 9, 1, "Supported methods (memo)")
    wa.cell(r0 + 9, 2, f"{METHOD_PRORATION} (governing) | {METHOD_CUTOFF} (conditional)")
    wa.cell(r0 + 10, 1, "Effective Kiln factor applied to selected board figure")
    wa.cell(r0 + 10, 2, f'=IF(B{r0 + 8}="{METHOD_CUTOFF}",1,B{r0 + 7})')
    wa.cell(r0 + 11, 1, "Governing FX source")
    wa.cell(
        r0 + 11,
        2,
        "brine_barrel_fx_and_period_reference.txt (1.0874; ignore spot aside 1.1025)",
    )
    wa.cell(r0 + 12, 1, "Brine FX rounding convention")
    wa.cell(
        r0 + 12,
        2,
        (
            "Brine EUR revenue is summed in native EUR before conversion to USD; "
            "the converted total is rounded to cents. Line-level USD rounding is not "
            "used for the consolidated Brine total."
        ),
    )
    wa.cell(r0 + 13, 1, "Governing acquisition / 50/92 source")
    wa.cell(
        r0 + 13,
        2,
        "50/92 ownership/proration method — governed by kiln_spice_acquisition_memo.txt",
    )
    wa.cell(r0 + 14, 1, "Reporting period label")
    wa.cell(r0 + 14, 2, f"{Q3_START} to {Q3_END}")
    wa.cell(r0 + 15, 1, "Acquisition transaction document")
    wa.cell(r0 + 15, 2, "Stock purchase agreement (signed 2025-08-06; close 2025-08-12)")
    wa.cell(r0 + 16, 1, "As-of date for this exercise")
    wa.cell(r0 + 16, 2, "Early October 2025 (historical Q3 2025 close)")
    wa.cell(r0 + 17, 1, "Approval evidence in supplied packet (cutoff method)")
    wa.cell(r0 + 17, 2, APPROVAL_EVIDENCE)
    wa.cell(r0 + 18, 1, "Cutoff scenario label")
    wa.cell(r0 + 18, 2, CUTOFF_LABEL)
    wa.cell(r0 + 19, 1, "Modeling basis (Kiln ownership date)")
    wa.cell(r0 + 19, 2, OWNERSHIP_MODELING_BASIS)
    wa.cell(r0 + 20, 1, "Ownership Date Evidence Status")
    wa.cell(r0 + 20, 2, OWNERSHIP_EVIDENCE_STATUS)
    wa.cell(r0 + 21, 1, "Supporting evidence status (ownership date)")
    wa.cell(r0 + 21, 2, OWNERSHIP_SUPPORTING_EVIDENCE)
    wa.cell(r0 + 22, 1, "Post-Close Billing Support")
    wa.cell(r0 + 22, 2, POST_CLOSE_BILLING_SUPPORT)

    a_period_start = f"Assumptions!$B${r0 + 1}"
    a_period_end = f"Assumptions!$B${r0 + 2}"
    a_fx = f"Assumptions!$B${r0 + 3}"
    a_kiln_close = f"Assumptions!$B${r0 + 4}"
    a_proration = f"Assumptions!$B${r0 + 7}"
    a_method = f"Assumptions!$B${r0 + 8}"
    a_eff_factor = f"Assumptions!$B${r0 + 10}"
    a_fx_plain = f"Assumptions!B{r0 + 3}"
    a_method_plain = f"Assumptions!B{r0 + 8}"
    a_proration_plain = f"Assumptions!B{r0 + 7}"
    a_approval_plain = f"Assumptions!B{r0 + 17}"
    a_cutoff_label_plain = f"Assumptions!B{r0 + 18}"
    a_ownership_model_plain = f"Assumptions!B{r0 + 19}"
    a_ownership_evidence_plain = f"Assumptions!B{r0 + 20}"
    a_ownership_support_plain = f"Assumptions!B{r0 + 21}"
    a_billing_support_plain = f"Assumptions!B{r0 + 22}"

    formula_cache["Assumptions"] = {
        f"B{r0 + 7}": round(KILN_OWNED_DAYS / Q3_DAYS, 10),
        f"B{r0 + 10}": round(KILN_OWNED_DAYS / Q3_DAYS, 10)
        if SELECTED_METHOD == METHOD_PRORATION
        else 1.0,
    }

    wa.column_dimensions["A"].width = 52
    wa.column_dimensions["B"].width = 88
    wa.freeze_panes = "A6"

    # --- Line_Detail ---
    wl = wb.create_sheet("Line_Detail")
    headers = [
        "transaction_id",
        "brand",
        "transaction_date",
        "description",
        "amount_native",
        "currency",
        "status",
        "include_status_period",
        "exclusion_reason",
        "eligible_native",
        "kiln_post_close_eligible",
        "usd_if_status_period_included",
    ]
    for col, h in enumerate(headers, 1):
        cell = wl.cell(1, col, h)
        cell.font = HEADER_FONT

    hearth_gross = 0.0
    brine_eur_gross = 0.0
    kiln_full_gross = 0.0
    kiln_cutoff_gross = 0.0
    exclusions: list[tuple[str, str, str, float, str]] = []
    kiln_line_rows: list[int] = []

    for i, row in enumerate(rows, start=2):
        reason, include = classify_status_period(row)
        amt = float(row["amount"])
        currency = row["currency"]
        brand = row["brand"]
        date = row["transaction_date"]

        wl.cell(i, 1, row["transaction_id"])
        wl.cell(i, 2, brand)
        wl.cell(i, 3, date)
        wl.cell(i, 4, row["description"])
        wl.cell(i, 5, amt)
        wl.cell(i, 6, currency)
        wl.cell(i, 7, row["status"])

        wl.cell(
            i,
            8,
            (
                f'=IF(OR(G{i}="void",G{i}="cancelled",G{i}="duplicate",G{i}="out_of_period"),'
                f'"no",IF(OR(C{i}<{a_period_start},C{i}>{a_period_end}),"no","yes"))'
            ),
        )
        wl.cell(i, 9, reason)
        wl.cell(i, 10, f'=IF(H{i}="yes",E{i},0)')
        wl.cell(
            i,
            11,
            (
                f'=IF(AND(H{i}="yes",B{i}="Kiln Spice Works",'
                f"C{i}>={a_kiln_close}),E{i},0)"
            ),
        )
        if currency == "EUR":
            wl.cell(i, 12, f'=IF(H{i}="yes",ROUND(E{i}*{a_fx},2),0)')
            usd_val = round2(amt * EUR_USD) if include else 0.0
        else:
            wl.cell(i, 12, f'=IF(H{i}="yes",E{i},0)')
            usd_val = round2(amt) if include else 0.0

        if currency == "EUR":
            wl.cell(i, 5).number_format = EUR_FORMAT
            wl.cell(i, 10).number_format = EUR_FORMAT
            wl.cell(i, 12).number_format = USD_FORMAT
        else:
            wl.cell(i, 5).number_format = USD_FORMAT
            wl.cell(i, 10).number_format = USD_FORMAT
            wl.cell(i, 11).number_format = USD_FORMAT
            wl.cell(i, 12).number_format = USD_FORMAT

        eligible = amt if include else 0.0
        kiln_pc = amt if (include and brand == "Kiln Spice Works" and date >= KILN_CLOSE) else 0.0

        formula_cache.setdefault("Line_Detail", {})[f"J{i}"] = eligible
        formula_cache["Line_Detail"][f"K{i}"] = kiln_pc
        formula_cache["Line_Detail"][f"L{i}"] = usd_val

        if brand == "Kiln Spice Works":
            kiln_line_rows.append(i)

        if include:
            if brand == "Hearth Kitchen Co":
                hearth_gross = hearth_gross + amt
            elif brand == "Brine & Barrel EU":
                brine_eur_gross = brine_eur_gross + amt
            elif brand == "Kiln Spice Works":
                kiln_full_gross = kiln_full_gross + amt
                if date >= KILN_CLOSE:
                    kiln_cutoff_gross = kiln_cutoff_gross + amt
        else:
            exclusions.append(
                (row["transaction_id"], brand, reason or row["status"], amt, currency)
            )

    hearth_gross = round2(hearth_gross)
    brine_eur_gross = round2(brine_eur_gross)
    kiln_full_gross = round2(kiln_full_gross)
    kiln_cutoff_gross = round2(kiln_cutoff_gross)

    last_line = 1 + len(rows)
    last_row = last_line

    kiln_prorated = round2(kiln_full_gross * KILN_OWNED_DAYS / Q3_DAYS)
    kiln_base = (
        kiln_prorated if SELECTED_METHOD == METHOD_PRORATION else round2(kiln_cutoff_gross)
    )
    kiln_native_base = (
        kiln_full_gross if SELECTED_METHOD == METHOD_PRORATION else kiln_cutoff_gross
    )
    brine_usd = round2(brine_eur_gross * EUR_USD)
    brine_line_round_control = round2(
        sum(
            round2(float(r["amount"]) * EUR_USD)
            for r in rows
            if r["brand"] == "Brine & Barrel EU" and classify_status_period(r)[1]
        )
    )
    hearth_usd = round2(hearth_gross)
    total = round2(hearth_usd + brine_usd + kiln_base)

    kiln_alt = round2(kiln_cutoff_gross)
    total_proration = round2(hearth_usd + brine_usd + kiln_prorated)
    total_alt = round2(hearth_usd + brine_usd + kiln_alt)
    factor_cache = round(KILN_OWNED_DAYS / Q3_DAYS, 10)

    style_table(
        wl,
        1,
        last_row,
        12,
        [14, 20, 14, 38, 14, 10, 12, 18, 18, 14, 18, 18],
    )

    # --- Brand_Rollup ---
    wr = wb.create_sheet("Brand_Rollup")
    r0 = write_title_block(wr, "Brand_Rollup", cols=5)
    wr.cell(r0, 1, "Brand").font = HEADER_FONT
    wr.cell(r0, 2, "Native gross (method-aware)").font = HEADER_FONT
    wr.cell(r0, 3, "Currency").font = HEADER_FONT
    wr.cell(r0, 4, "FX / ownership factor").font = HEADER_FONT
    wr.cell(r0, 5, "Consolidated USD (selected final method)").font = HEADER_FONT

    wr.cell(r0 + 1, 1, "Hearth Kitchen Co")
    wr.cell(
        r0 + 1,
        2,
        f'=SUMIF(Line_Detail!$B$2:$B${last_row},A{r0 + 1},Line_Detail!$J$2:$J${last_row})',
    )
    wr.cell(r0 + 1, 3, "USD")
    wr.cell(r0 + 1, 4, 1)
    wr.cell(r0 + 1, 5, f"=ROUND(B{r0 + 1}*D{r0 + 1},2)")

    wr.cell(r0 + 2, 1, "Brine & Barrel EU")
    # Live chain: Line_Detail → Reconciliation!D12 → F12 → Brand_Rollup
    wr.cell(r0 + 2, 2, "=Reconciliation!D12")
    wr.cell(r0 + 2, 3, "EUR")
    wr.cell(r0 + 2, 4, f"={a_fx_plain}")
    wr.cell(r0 + 2, 5, "=Reconciliation!F12")

    wr.cell(r0 + 3, 1, "Kiln Spice Works")
    wr.cell(
        r0 + 3,
        2,
        (
            f'=IF({a_method}="{METHOD_CUTOFF}",'
            f"SUMIF(Line_Detail!$B$2:$B${last_row},A{r0 + 3},Line_Detail!$K$2:$K${last_row}),"
            f"Reconciliation!D13)"
        ),
    )
    wr.cell(r0 + 3, 3, "USD")
    wr.cell(r0 + 3, 4, f"={a_eff_factor}")
    wr.cell(r0 + 3, 5, f"=ROUND(B{r0 + 3}*D{r0 + 3},2)")

    wr.cell(r0 + 5, 1, "Total consolidated Q3 2025 (selected final method)").font = HEADER_FONT
    wr.cell(r0 + 5, 5, f"=ROUND(SUM(E{r0 + 1}:E{r0 + 3}),2)")

    wr.cell(r0 + 7, 1, "Selected final Kiln ownership method")
    wr.cell(r0 + 7, 2, f"={a_method_plain}")
    wr.cell(r0 + 8, 1, "Approval evidence in supplied packet")
    wr.cell(r0 + 8, 2, f"={a_approval_plain}")
    wr.cell(r0 + 9, 1, "Modeling basis (Kiln ownership date)")
    wr.cell(r0 + 9, 2, f"={a_ownership_model_plain}")
    wr.cell(r0 + 10, 1, "Ownership Date Evidence Status")
    wr.cell(r0 + 10, 2, f"={a_ownership_evidence_plain}")
    wr.cell(r0 + 11, 1, "Post-Close Billing Support")
    wr.cell(r0 + 11, 2, f"={a_billing_support_plain}")

    wr.cell(
        r0 + 13,
        1,
        "Case comparison (Case 1 = selected board method; Case 2 = conditional only)",
    ).font = HEADER_FONT
    wr.cell(r0 + 14, 1, "Metric").font = HEADER_FONT
    wr.cell(r0 + 14, 2, "Case 1: 50/92 proration (selected final)").font = HEADER_FONT
    wr.cell(r0 + 14, 3, "Case 2: Conditional cutoff (not evidenced)").font = HEADER_FONT

    wr.cell(r0 + 15, 1, "Hearth Kitchen Co USD")
    wr.cell(r0 + 15, 2, f"=E{r0 + 1}")
    wr.cell(r0 + 15, 3, f"=E{r0 + 1}")

    wr.cell(r0 + 16, 1, "Brine & Barrel EU USD")
    wr.cell(r0 + 16, 2, f"=E{r0 + 2}")
    wr.cell(r0 + 16, 3, f"=E{r0 + 2}")

    wr.cell(r0 + 17, 1, "Kiln Spice Works USD")
    wr.cell(
        r0 + 17,
        2,
        (
            f"=ROUND(SUMIF(Line_Detail!$B$2:$B${last_row},\"Kiln Spice Works\","
            f"Line_Detail!$J$2:$J${last_row})*{a_proration},2)"
        ),
    )
    wr.cell(
        r0 + 17,
        3,
        (
            f"=ROUND(SUMIF(Line_Detail!$B$2:$B${last_row},\"Kiln Spice Works\","
            f"Line_Detail!$K$2:$K${last_row}),2)"
        ),
    )

    wr.cell(r0 + 18, 1, "Total consolidated Q3 2025 USD").font = HEADER_FONT
    wr.cell(r0 + 18, 2, f"=ROUND(B{r0 + 15}+B{r0 + 16}+B{r0 + 17},2)")
    wr.cell(r0 + 18, 3, f"=ROUND(C{r0 + 15}+C{r0 + 16}+C{r0 + 17},2)")

    wr.cell(r0 + 19, 1, "Case 2 reporting label")
    wr.cell(r0 + 19, 3, f"={a_cutoff_label_plain}")

    wr.cell(r0 + 21, 1, "Method note")
    wr.cell(
        r0 + 21,
        2,
        (
            "Selected final board figure uses 50/92 proration governed by "
            "kiln_spice_acquisition_memo.txt. Case 2 recalculates the transaction-date "
            "cutoff from Line_Detail kiln_post_close_eligible and is labeled a "
            "Conditional transaction-date cutoff scenario because approval evidence "
            "is not present in the supplied packet. Ownership-date support (Schedule 2.1 "
            "/ bank wire) and post-close billing shipment/service support remain open "
            "follow-ups. Do not mix methods. See Kiln_Audit for every Kiln row."
        ),
    )
    wr.cell(r0 + 22, 1, "Brine FX note")
    wr.cell(
        r0 + 22,
        2,
        (
            "Governing Brine USD uses sum-then-round (native EUR sum × 1.0874, rounded to cents). "
            f"Line-by-line convert-then-round control equals ${brine_line_round_control:,.2f} "
            "and is not used for the consolidated Brine total."
        ),
    )

    br_hearth = f"Brand_Rollup!E{r0 + 1}"
    br_brine = f"Brand_Rollup!E{r0 + 2}"
    br_kiln = f"Brand_Rollup!E{r0 + 3}"
    br_total = f"Brand_Rollup!E{r0 + 5}"
    br_case1_kiln = f"Brand_Rollup!B{r0 + 17}"
    br_case1_total = f"Brand_Rollup!B{r0 + 18}"
    br_case2_kiln = f"Brand_Rollup!C{r0 + 17}"
    br_case2_total = f"Brand_Rollup!C{r0 + 18}"

    fmt_usd(
        wr,
        f"B{r0 + 1}",
        f"E{r0 + 1}",
        f"E{r0 + 2}",
        f"B{r0 + 3}",
        f"E{r0 + 3}",
        f"E{r0 + 5}",
        f"B{r0 + 15}",
        f"C{r0 + 15}",
        f"B{r0 + 16}",
        f"C{r0 + 16}",
        f"B{r0 + 17}",
        f"C{r0 + 17}",
        f"B{r0 + 18}",
        f"C{r0 + 18}",
    )
    fmt_eur(wr, f"B{r0 + 2}")
    wr[f"D{r0 + 2}"].number_format = RATE_FORMAT

    formula_cache["Brand_Rollup"] = {
        f"B{r0 + 1}": hearth_usd,
        f"E{r0 + 1}": hearth_usd,
        f"B{r0 + 2}": brine_eur_gross,
        f"D{r0 + 2}": EUR_USD,
        f"E{r0 + 2}": brine_usd,
        f"B{r0 + 3}": kiln_native_base,
        f"D{r0 + 3}": factor_cache if SELECTED_METHOD == METHOD_PRORATION else 1.0,
        f"E{r0 + 3}": kiln_base,
        f"E{r0 + 5}": total,
        f"B{r0 + 7}": SELECTED_METHOD,
        f"B{r0 + 8}": APPROVAL_EVIDENCE,
        f"B{r0 + 9}": OWNERSHIP_MODELING_BASIS,
        f"B{r0 + 10}": OWNERSHIP_EVIDENCE_STATUS,
        f"B{r0 + 11}": POST_CLOSE_BILLING_SUPPORT,
        f"B{r0 + 15}": hearth_usd,
        f"C{r0 + 15}": hearth_usd,
        f"B{r0 + 16}": brine_usd,
        f"C{r0 + 16}": brine_usd,
        f"B{r0 + 17}": kiln_prorated,
        f"C{r0 + 17}": kiln_alt,
        f"B{r0 + 18}": total_proration,
        f"C{r0 + 18}": total_alt,
        f"C{r0 + 19}": CUTOFF_LABEL,
    }

    for idx, width in enumerate([42, 36, 14, 18, 36], start=1):
        wr.column_dimensions[get_column_letter(idx)].width = width
    wr.freeze_panes = "A6"

    # --- Kiln_Audit ---
    waud = wb.create_sheet("Kiln_Audit")
    aud_headers = [
        "transaction_id",
        "transaction_date",
        "amount_native",
        "status",
        "status_period_include",
        "pre_or_post_close",
        "cutoff_include",
        "cutoff_exclude_reason",
        "amount_to_cutoff_revenue",
        "amount_to_full_quarter_eligible",
        "Line_Detail_row",
    ]
    for col, h in enumerate(aud_headers, 1):
        waud.cell(1, col, h).font = HEADER_FONT

    aud_cache: dict[str, float | str] = {}
    for aud_i, line_i in enumerate(kiln_line_rows, start=2):
        waud.cell(aud_i, 1, f"=Line_Detail!A{line_i}")
        waud.cell(aud_i, 2, f"=Line_Detail!C{line_i}")
        waud.cell(aud_i, 3, f"=Line_Detail!E{line_i}")
        waud.cell(aud_i, 4, f"=Line_Detail!G{line_i}")
        waud.cell(aud_i, 5, f"=Line_Detail!H{line_i}")
        waud.cell(aud_i, 6, f'=IF(B{aud_i}<{a_kiln_close},"pre-close","post-close")')
        waud.cell(
            aud_i,
            7,
            f'=IF(AND(E{aud_i}="yes",F{aud_i}="post-close"),"included","excluded")',
        )
        waud.cell(
            aud_i,
            8,
            (
                f'=IF(G{aud_i}="included","",'
                f'IF(E{aud_i}<>"yes",IF(Line_Detail!I{line_i}<>"",Line_Detail!I{line_i},"status_or_period"),'
                f'"pre-close_ownership"))'
            ),
        )
        waud.cell(aud_i, 9, f"=Line_Detail!K{line_i}")
        waud.cell(aud_i, 10, f"=Line_Detail!J{line_i}")
        waud.cell(aud_i, 11, line_i)
        waud.cell(aud_i, 3).number_format = USD_FORMAT
        waud.cell(aud_i, 9).number_format = USD_FORMAT
        waud.cell(aud_i, 10).number_format = USD_FORMAT

        src = rows[line_i - 2]
        include = classify_status_period(src)[1]
        amt = float(src["amount"])
        date = src["transaction_date"]
        post = date >= KILN_CLOSE
        aud_cache[f"A{aud_i}"] = src["transaction_id"]
        aud_cache[f"B{aud_i}"] = date
        aud_cache[f"C{aud_i}"] = amt
        aud_cache[f"D{aud_i}"] = src["status"]
        aud_cache[f"E{aud_i}"] = "yes" if include else "no"
        aud_cache[f"F{aud_i}"] = "post-close" if post else "pre-close"
        aud_cache[f"G{aud_i}"] = "included" if (include and post) else "excluded"
        aud_cache[f"I{aud_i}"] = amt if (include and post) else 0.0
        aud_cache[f"J{aud_i}"] = amt if include else 0.0

    aud_last = 1 + len(kiln_line_rows)
    sum_row = aud_last + 1
    waud.cell(sum_row, 1, "Kiln method totals (formula sums)").font = HEADER_FONT
    waud.cell(sum_row, 9, f"=SUM(I2:I{aud_last})")
    waud.cell(sum_row, 10, f"=SUM(J2:J{aud_last})")
    waud.cell(sum_row, 9).number_format = USD_FORMAT
    waud.cell(sum_row, 10).number_format = USD_FORMAT
    aud_cache[f"I{sum_row}"] = kiln_alt
    aud_cache[f"J{sum_row}"] = kiln_full_gross

    waud.cell(sum_row + 2, 1, "Cutoff Kiln USD (Case 2 — conditional)")
    waud.cell(sum_row + 2, 2, f"=I{sum_row}")
    waud.cell(sum_row + 3, 1, "Full-quarter Kiln eligible USD (pre-proration)")
    waud.cell(sum_row + 3, 2, f"=J{sum_row}")
    waud.cell(sum_row + 4, 1, "50/92 Kiln USD (Case 1 — selected final)")
    waud.cell(sum_row + 4, 2, f"=ROUND(J{sum_row}*{a_proration},2)")
    waud.cell(sum_row + 5, 1, "Governing close date (modeling)")
    waud.cell(sum_row + 5, 2, f"={a_kiln_close}")
    waud.cell(sum_row + 6, 1, "Rule")
    waud.cell(
        sum_row + 6,
        2,
        (
            "50/92 ownership/proration method — governed by kiln_spice_acquisition_memo.txt "
            "(selected final). Transaction-date cutoff may be calculated as a Conditional "
            "transaction-date cutoff scenario only; approval evidence is not present in the "
            "supplied packet. Ownership-date and post-close billing support remain open "
            "follow-ups. Do not mix methods."
        ),
    )
    waud.cell(sum_row + 7, 1, "Approval evidence in supplied packet")
    waud.cell(sum_row + 7, 2, f"={a_approval_plain}")
    waud.cell(sum_row + 8, 1, "Modeling basis (Kiln ownership date)")
    waud.cell(sum_row + 8, 2, f"={a_ownership_model_plain}")
    waud.cell(sum_row + 9, 1, "Ownership Date Evidence Status")
    waud.cell(sum_row + 9, 2, f"={a_ownership_evidence_plain}")
    waud.cell(sum_row + 10, 1, "Supporting evidence status (ownership date)")
    waud.cell(sum_row + 10, 2, f"={a_ownership_support_plain}")
    waud.cell(sum_row + 11, 1, "Post-Close Billing Support")
    waud.cell(sum_row + 11, 2, f"={a_billing_support_plain}")
    for coord in (f"B{sum_row + 2}", f"B{sum_row + 3}", f"B{sum_row + 4}"):
        waud[coord].number_format = USD_FORMAT
    aud_cache[f"B{sum_row + 2}"] = kiln_alt
    aud_cache[f"B{sum_row + 3}"] = kiln_full_gross
    aud_cache[f"B{sum_row + 4}"] = kiln_prorated
    aud_cache[f"B{sum_row + 5}"] = KILN_CLOSE
    aud_cache[f"B{sum_row + 7}"] = APPROVAL_EVIDENCE
    aud_cache[f"B{sum_row + 8}"] = OWNERSHIP_MODELING_BASIS
    aud_cache[f"B{sum_row + 9}"] = OWNERSHIP_EVIDENCE_STATUS
    aud_cache[f"B{sum_row + 10}"] = OWNERSHIP_SUPPORTING_EVIDENCE
    aud_cache[f"B{sum_row + 11}"] = POST_CLOSE_BILLING_SUPPORT
    formula_cache["Kiln_Audit"] = aud_cache
    style_table(
        waud,
        1,
        aud_last,
        11,
        [14, 14, 14, 12, 16, 14, 12, 20, 18, 22, 12],
    )

    # --- Reconciliation: formula-driven raw→adjusted bridge + exclusion log ---
    wx = wb.create_sheet("Reconciliation")
    recon_cache: dict[str, float | str] = {}

    ld_b = f"Line_Detail!$B$2:$B${last_row}"
    ld_e = f"Line_Detail!$E$2:$E${last_row}"
    ld_i = f"Line_Detail!$I$2:$I${last_row}"
    ld_j = f"Line_Detail!$J$2:$J${last_row}"

    def sumif_brand(amount_range: str, brand: str) -> str:
        return f'=SUMIF({ld_b},"{brand}",{amount_range})'

    def void_cancelled(brand: str) -> str:
        return (
            f"=SUMIFS({ld_e},{ld_b},\"{brand}\",{ld_i},\"status_void\")"
            f"+SUMIFS({ld_e},{ld_b},\"{brand}\",{ld_i},\"status_cancelled\")"
        )

    def period_excl(brand: str) -> str:
        return (
            f"=SUMIFS({ld_e},{ld_b},\"{brand}\",{ld_i},\"prior_period\")"
            f"+SUMIFS({ld_e},{ld_b},\"{brand}\",{ld_i},\"post_period\")"
            f"+SUMIFS({ld_e},{ld_b},\"{brand}\",{ld_i},\"status_out_of_period\")"
        )

    def duplicate_excl(brand: str) -> str:
        return f'=SUMIFS({ld_e},{ld_b},"{brand}",{ld_i},"status_duplicate")'

    hearth = "Hearth Kitchen Co"
    brine = "Brine & Barrel EU"
    kiln = "Kiln Spice Works"

    # Pre-compute category totals for formula cache
    def _excl_amt(brand: str, *reasons: str) -> float:
        return round2(
            sum(
                float(r["amount"])
                for r in rows
                if r["brand"] == brand and classify_status_period(r)[0] in reasons
            )
        )

    hearth_raw = round2(
        sum(float(r["amount"]) for r in rows if r["brand"] == hearth)
    )
    brine_raw = round2(
        sum(float(r["amount"]) for r in rows if r["brand"] == brine)
    )
    kiln_raw = round2(
        sum(float(r["amount"]) for r in rows if r["brand"] == kiln)
    )
    hearth_void = _excl_amt(hearth, "status_void", "status_cancelled")
    brine_void = _excl_amt(brine, "status_void", "status_cancelled")
    kiln_void = _excl_amt(kiln, "status_void", "status_cancelled")
    hearth_dup = _excl_amt(hearth, "status_duplicate")
    brine_dup = _excl_amt(brine, "status_duplicate")
    kiln_dup = _excl_amt(kiln, "status_duplicate")
    hearth_period = _excl_amt(hearth, "prior_period", "post_period", "status_out_of_period")
    brine_period = _excl_amt(brine, "prior_period", "post_period", "status_out_of_period")
    kiln_period = _excl_amt(kiln, "prior_period", "post_period", "status_out_of_period")
    kiln_own_adj = round2(kiln_prorated - kiln_full_gross)

    wx["A1"] = "Reconciliation bridge — raw extract gross to Brand_Rollup"
    wx["A1"].font = TITLE_FONT
    wx.merge_cells("A1:H1")
    wx["A2"] = (
        "Formula-driven bridge. Editing Line_Detail updates raw gross, exclusions, "
        "D12/F12 (Brine), D13 (Kiln), Brand_Rollup, and control difference."
    )
    wx["A2"].font = Font(italic=True, size=9, color="595959")
    wx.merge_cells("A2:H2")

    headers = [
        "Category",
        "Hearth Kitchen Co (USD)",
        "Brine & Barrel EU (EUR)",
        "Native eligible link",
        "Kiln Spice Works (USD)",
        "Brine / board USD",
        "Consolidated USD",
        "Source / note",
    ]
    for col, h in enumerate(headers, 1):
        wx.cell(3, col, h).font = HEADER_FONT
        wx.cell(3, col).fill = SECTION_FILL

    # Row 4: Raw extract gross
    wx["A4"] = "Raw extract gross"
    wx["B4"] = sumif_brand(ld_e, hearth)
    wx["C4"] = sumif_brand(ld_e, brine)
    wx["E4"] = sumif_brand(ld_e, kiln)
    wx["F4"] = f"=ROUND(C4*{a_fx},2)"
    wx["G4"] = "=ROUND(B4+F4+E4,2)"
    wx["H4"] = "vesper_culinary_brands_brand_ledger.csv (all extract rows)"

    # Row 5: void/cancelled
    wx["A5"] = "Less: void / cancelled status exclusions"
    wx["B5"] = void_cancelled(hearth)
    wx["C5"] = void_cancelled(brine)
    wx["E5"] = void_cancelled(kiln)
    wx["F5"] = f"=ROUND(C5*{a_fx},2)"
    wx["G5"] = "=ROUND(B5+F5+E5,2)"
    wx["H5"] = "Ledger status void/cancelled; brine_barrel_fx_and_period_reference.txt"

    # Row 6: duplicate
    wx["A6"] = "Less: duplicate exclusions"
    wx["B6"] = duplicate_excl(hearth)
    wx["C6"] = duplicate_excl(brine)
    wx["E6"] = duplicate_excl(kiln)
    wx["F6"] = f"=ROUND(C6*{a_fx},2)"
    wx["G6"] = "=ROUND(B6+F6+E6,2)"
    wx["H6"] = "Ledger status=duplicate"

    # Row 7: period
    wx["A7"] = "Less: period exclusions (prior / post / out_of_period)"
    wx["B7"] = period_excl(hearth)
    wx["C7"] = period_excl(brine)
    wx["E7"] = period_excl(kiln)
    wx["F7"] = f"=ROUND(C7*{a_fx},2)"
    wx["G7"] = "=ROUND(B7+F7+E7,2)"
    wx["H7"] = "Q3 window + out_of_period status per governing period rules"

    # Row 8: eligible subtotal (links toward D12/D13)
    wx["A8"] = "Subtotal: eligible after status / period"
    wx["B8"] = sumif_brand(ld_j, hearth)
    wx["C8"] = "=D12"  # live Brine eligible EUR
    wx["E8"] = "=D13"  # live Kiln eligible USD
    wx["F8"] = "=F12"
    wx["G8"] = "=ROUND(B8+F8+E8,2)"
    wx["H8"] = "Line_Detail eligible_native; D12/D13 live links"

    # Row 9: ownership adjustment (Kiln 50/92 haircut only)
    wx["A9"] = "Ownership adjustment (Kiln 50/92; selected board method)"
    wx["B9"] = 0
    wx["C9"] = 0
    wx["E9"] = f"=ROUND(D13*{a_proration},2)-D13"
    wx["F9"] = 0
    wx["G9"] = "=E9"
    wx["H9"] = "50/92 — kiln_spice_acquisition_memo.txt"

    # Row 10: adjusted board revenue
    wx["A10"] = "Adjusted board revenue"
    wx["B10"] = "=B8"
    wx["C10"] = "=D12"
    wx["E10"] = f"=ROUND(D13*{a_proration},2)"
    wx["F10"] = "=F12"
    wx["G10"] = "=ROUND(B10+F10+E10,2)"
    wx["H10"] = "Board figure under selected 50/92 method"

    # Row 11: Brand_Rollup control
    wx["A11"] = "Brand_Rollup control total"
    wx["B11"] = f"={br_hearth}"
    wx["C11"] = "=D12"
    wx["E11"] = f"={br_kiln}"
    wx["F11"] = f"={br_brine}"
    wx["G11"] = f"={br_total}"
    wx["H11"] = "Must equal adjusted board revenue"

    # Rows 12–13: reviewer-required live links (D12 / D13 / F12)
    wx["A12"] = "Brine eligible EUR (live from Line_Detail)"
    wx["C12"] = brine
    wx["D12"] = sumif_brand(ld_j, brine)
    wx["E12"] = "EUR"
    wx["F12"] = f"=ROUND(D12*{a_fx},2)"
    wx["G12"] = "→ Brand_Rollup Brine USD"
    wx["H12"] = "brine_barrel_fx_and_period_reference.txt (EUR/USD 1.0874; sum-then-round)"

    wx["A13"] = "Kiln full-quarter eligible USD (live from Kiln_Audit)"
    wx["C13"] = kiln
    wx["D13"] = "=Kiln_Audit!J15"
    wx["E13"] = "USD"
    wx["F13"] = f"=ROUND(D13*{a_proration},2)"
    wx["G13"] = "→ Brand_Rollup Kiln USD (50/92)"
    wx["H13"] = "kiln_spice_acquisition_memo.txt (50/92 selected final)"

    # Row 14–18: explicit controls
    wx["A14"] = "Control checks"
    wx["A14"].font = HEADER_FONT
    wx["A15"] = "Control 1 — raw Line_Detail by brand = raw bridge total"
    wx["B15"] = f"=ROUND(B4-SUMIF({ld_b},\"{hearth}\",{ld_e}),2)"
    wx["C15"] = f"=ROUND(C4-SUMIF({ld_b},\"{brine}\",{ld_e}),2)"
    wx["E15"] = f"=ROUND(E4-SUMIF({ld_b},\"{kiln}\",{ld_e}),2)"
    wx["G15"] = '=IF(AND(ABS(B15)<0.005,ABS(C15)<0.005,ABS(E15)<0.005),"PASS","FAIL")'
    wx["H15"] = "Difference must be zero"

    wx["A16"] = "Control 2 — raw less exclusion categories = eligible / adjusted path"
    wx["B16"] = "=ROUND(B4-B5-B6-B7-B8,2)"
    wx["C16"] = "=ROUND(C4-C5-C6-C7-C8,2)"
    wx["E16"] = "=ROUND(E4-E5-E6-E7-E8,2)"
    wx["G16"] = '=IF(AND(ABS(B16)<0.005,ABS(C16)<0.005,ABS(E16)<0.005),"PASS","FAIL")'
    wx["H16"] = "Bridge arithmetic identity"

    wx["A17"] = "Control 3 — adjusted brand total = Brand_Rollup"
    wx["B17"] = "=ROUND(B10-B11,2)"
    wx["F17"] = "=ROUND(F10-F11,2)"
    wx["E17"] = "=ROUND(E10-E11,2)"
    wx["G17"] = "=ROUND(G10-G11,2)"
    wx["H17"] = '=IF(AND(ABS(B17)<0.005,ABS(F17)<0.005,ABS(E17)<0.005,ABS(G17)<0.005),"PASS","FAIL")'

    wx["A18"] = "Control 4 — consolidated difference"
    wx["G18"] = "=G17"
    wx["H18"] = '=IF(ABS(G18)<0.005,"PASS","FAIL")'

    wx["A19"] = "Bridge control status (all controls)"
    wx["G19"] = '=IF(AND(G15="PASS",G16="PASS",H17="PASS",H18="PASS"),"PASS","FAIL")'
    wx["G19"].font = HEADER_FONT

    for coord in (
        "B4", "E4", "F4", "G4",
        "B5", "E5", "F5", "G5",
        "B6", "E6", "F6", "G6",
        "B7", "E7", "F7", "G7",
        "B8", "E8", "F8", "G8",
        "B9", "E9", "G9",
        "B10", "E10", "F10", "G10",
        "B11", "E11", "F11", "G11",
        "D12", "F12", "D13", "F13",
        "B15", "C15", "E15",
        "B16", "C16", "E16",
        "B17", "F17", "E17", "G17", "G18",
    ):
        wx[coord].number_format = USD_FORMAT
    for coord in ("C4", "C5", "C6", "C7", "C8", "C10", "C11", "D12"):
        wx[coord].number_format = EUR_FORMAT
    wx["C9"].number_format = EUR_FORMAT

    recon_cache.update(
        {
            "B4": hearth_raw,
            "C4": brine_raw,
            "E4": kiln_raw,
            "F4": round2(brine_raw * EUR_USD),
            "G4": round2(hearth_raw + round2(brine_raw * EUR_USD) + kiln_raw),
            "B5": hearth_void,
            "C5": brine_void,
            "E5": kiln_void,
            "F5": round2(brine_void * EUR_USD),
            "G5": round2(hearth_void + round2(brine_void * EUR_USD) + kiln_void),
            "B6": hearth_dup,
            "C6": brine_dup,
            "E6": kiln_dup,
            "F6": round2(brine_dup * EUR_USD),
            "G6": round2(hearth_dup + round2(brine_dup * EUR_USD) + kiln_dup),
            "B7": hearth_period,
            "C7": brine_period,
            "E7": kiln_period,
            "F7": round2(brine_period * EUR_USD),
            "G7": round2(hearth_period + round2(brine_period * EUR_USD) + kiln_period),
            "B8": hearth_usd,
            "C8": brine_eur_gross,
            "E8": kiln_full_gross,
            "F8": brine_usd,
            "G8": round2(hearth_usd + brine_usd + kiln_full_gross),
            "B9": 0.0,
            "C9": 0.0,
            "E9": kiln_own_adj,
            "F9": 0.0,
            "G9": kiln_own_adj,
            "B10": hearth_usd,
            "C10": brine_eur_gross,
            "E10": kiln_prorated,
            "F10": brine_usd,
            "G10": total,
            "B11": hearth_usd,
            "C11": brine_eur_gross,
            "E11": kiln_base,
            "F11": brine_usd,
            "G11": total,
            "D12": brine_eur_gross,
            "F12": brine_usd,
            "D13": kiln_full_gross,
            "F13": kiln_prorated,
            "B15": 0.0,
            "C15": 0.0,
            "E15": 0.0,
            "G15": "PASS",
            "B16": 0.0,
            "C16": 0.0,
            "E16": 0.0,
            "G16": "PASS",
            "B17": 0.0,
            "F17": 0.0,
            "E17": 0.0,
            "G17": 0.0,
            "H17": "PASS",
            "G18": 0.0,
            "H18": "PASS",
            "G19": "PASS",
        }
    )

    # Exclusion log (detail) — still required for FX impact on excluded Brine rows
    log_header = 21
    wx.cell(log_header, 1, "Exclusion log (transaction detail)").font = HEADER_FONT
    wx.merge_cells(start_row=log_header, start_column=1, end_row=log_header, end_column=8)
    log_cols = [
        "Issue",
        "transaction_id",
        "brand",
        "native_amount",
        "currency",
        "USD_Impact",
        "Action",
        "Source",
    ]
    for col, h in enumerate(log_cols, 1):
        wx.cell(log_header + 1, col, h).font = HEADER_FONT

    row_i = log_header + 2
    for tid, brand, reason, amt, cur in exclusions:
        wx.cell(row_i, 1, reason)
        wx.cell(row_i, 2, tid)
        wx.cell(row_i, 3, brand)
        wx.cell(row_i, 4, amt)
        wx.cell(row_i, 5, cur)
        wx.cell(row_i, 6, f'=IF(E{row_i}="EUR",ROUND(D{row_i}*{a_fx},2),D{row_i})')
        wx.cell(row_i, 7, "Excluded from consolidated Q3 revenue")
        wx.cell(row_i, 8, "vesper_culinary_brands_brand_ledger.csv")
        if cur == "EUR":
            wx.cell(row_i, 4).number_format = EUR_FORMAT
            recon_cache[f"F{row_i}"] = round2(amt * EUR_USD)
        else:
            wx.cell(row_i, 4).number_format = USD_FORMAT
            recon_cache[f"F{row_i}"] = round2(amt)
        wx.cell(row_i, 6).number_format = USD_FORMAT
        row_i += 1

    wx.cell(row_i, 1, "Source authority summary").font = HEADER_FONT
    row_i += 1
    wx.cell(row_i, 1, "50/92 ownership method")
    wx.cell(row_i, 8, "kiln_spice_acquisition_memo.txt")
    row_i += 1
    wx.cell(row_i, 1, "EUR/USD board close rate")
    wx.cell(row_i, 8, "brine_barrel_fx_and_period_reference.txt")
    row_i += 1
    wx.cell(row_i, 1, "Status / period treatment")
    wx.cell(
        row_i,
        8,
        "vesper_culinary_brands_brand_ledger.csv; brine_barrel_fx_and_period_reference.txt",
    )
    row_i += 1
    wx.cell(row_i, 1, "Approval evidence in supplied packet")
    wx.cell(row_i, 8, APPROVAL_EVIDENCE)
    row_i += 1
    wx.cell(row_i, 1, "Ownership Date Evidence Status")
    wx.cell(row_i, 8, OWNERSHIP_EVIDENCE_STATUS)
    row_i += 1
    wx.cell(row_i, 1, "Post-Close Billing Support")
    wx.cell(row_i, 8, POST_CLOSE_BILLING_SUPPORT)

    formula_cache["Reconciliation"] = recon_cache
    for idx, width in enumerate([48, 18, 18, 18, 18, 16, 16, 52], start=1):
        wx.column_dimensions[get_column_letter(idx)].width = width
    wx.freeze_panes = "A4"

    # --- Notes ---
    wn = wb.create_sheet("Notes")
    rn = write_title_block(wn, "Notes", cols=3)
    wn.cell(rn, 1, "Open Follow-Up Items").font = HEADER_FONT
    wn.cell(rn + 1, 1, "Item").font = HEADER_FONT
    wn.cell(rn + 1, 2, "Status").font = HEADER_FONT
    wn.cell(rn + 1, 3, "Evidence needed").font = HEADER_FONT
    wn.cell(rn + 2, 1, "Kiln ownership date")
    wn.cell(rn + 2, 2, "Open follow-up")
    wn.cell(rn + 2, 3, "Schedule 2.1 and bank wire confirmation")
    wn.cell(rn + 3, 1, "Post-close billings")
    wn.cell(rn + 3, 2, "Open follow-up")
    wn.cell(rn + 3, 3, "Shipment/service support")

    wn.cell(rn + 5, 1, "Ownership-date support").font = HEADER_FONT
    wn.cell(rn + 6, 1, OWNERSHIP_FOLLOWUP_NOTE)
    wn.merge_cells(start_row=rn + 6, start_column=1, end_row=rn + 6, end_column=3)
    wn.cell(rn + 6, 1).alignment = Alignment(wrap_text=True, vertical="top")
    wn.row_dimensions[rn + 6].height = 48

    wn.cell(rn + 7, 1, "Post-close billing support").font = HEADER_FONT
    wn.cell(rn + 8, 1, BILLING_FOLLOWUP_NOTE)
    wn.merge_cells(start_row=rn + 8, start_column=1, end_row=rn + 8, end_column=3)
    wn.cell(rn + 8, 1).alignment = Alignment(wrap_text=True, vertical="top")
    wn.row_dimensions[rn + 8].height = 48

    wn.cell(rn + 10, 1, "Kiln ownership methods and judgment").font = HEADER_FONT
    wn.cell(
        rn + 12,
        1,
        (
            "kiln_spice_acquisition_memo.txt establishes the permitted Q3 ownership treatment, "
            "including the 50/92 proration method used as the selected final board method. "
            "A transaction-date ownership cutoff (eligible Kiln rows on or after "
            f"{KILN_CLOSE} after status/period exclusions) is memo-supported only as a "
            "conditional alternative. Approval evidence is not present in the supplied packet. "
            "The August 12 ownership date is the modeling basis from the memo; independent "
            "confirmation against Schedule 2.1 and the bank wire confirmation remains an open "
            "follow-up. Do not mix methods or apply proration on top of an already "
            "ownership-cutoff sum."
        ),
    )
    wn.merge_cells(start_row=rn + 12, start_column=1, end_row=rn + 12, end_column=3)
    wn.cell(rn + 12, 1).alignment = Alignment(wrap_text=True, vertical="top")
    wn.row_dimensions[rn + 12].height = 78

    wn.cell(rn + 14, 1, "Selected final method (board)")
    wn.cell(rn + 14, 2, f"={a_method_plain}")
    wn.cell(rn + 15, 1, "Kiln consolidated USD — selected final")
    wn.cell(rn + 15, 2, f"={br_kiln}")
    wn.cell(rn + 16, 1, "Consolidated total USD — selected final")
    wn.cell(rn + 16, 2, f"={br_total}")
    wn.cell(rn + 17, 1, "Approval evidence in supplied packet")
    wn.cell(rn + 17, 2, f"={a_approval_plain}")
    wn.cell(rn + 18, 1, "Modeling basis (Kiln ownership date)")
    wn.cell(rn + 18, 2, f"={a_ownership_model_plain}")
    wn.cell(rn + 19, 1, "Ownership Date Evidence Status")
    wn.cell(rn + 19, 2, f"={a_ownership_evidence_plain}")
    wn.cell(rn + 20, 1, "Post-Close Billing Support")
    wn.cell(rn + 20, 2, f"={a_billing_support_plain}")

    wn.cell(rn + 22, 1, "50/92 method (Case 1 — selected final)").font = HEADER_FONT
    wn.cell(
        rn + 23,
        1,
        (
            "Sum status/period-eligible Kiln rows for the full Q3 window, then multiply by the "
            "50/92 proration factor. Governing source: kiln_spice_acquisition_memo.txt. "
            "This workbook uses 50/92 as the selected final board method. Ownership-date "
            "supporting evidence remains an open diligence item and does not change the "
            "selected method."
        ),
    )
    wn.merge_cells(start_row=rn + 23, start_column=1, end_row=rn + 23, end_column=3)
    wn.cell(rn + 23, 1).alignment = Alignment(wrap_text=True, vertical="top")
    wn.row_dimensions[rn + 23].height = 54
    wn.cell(rn + 24, 1, "Case 1 Kiln USD")
    wn.cell(rn + 24, 2, f"={br_case1_kiln}")
    wn.cell(rn + 25, 1, "Case 1 consolidated total USD")
    wn.cell(rn + 25, 2, f"={br_case1_total}")

    wn.cell(rn + 27, 1, "Transaction-date method (Case 2 — conditional)").font = HEADER_FONT
    wn.cell(
        rn + 28,
        1,
        (
            "After status/period exclusions, include only Kiln rows dated on or after the close "
            f"date {KILN_CLOSE}. Do not apply 50/92 on top of that ownership-cutoff sum. "
            f"{CUTOFF_LABEL}. {APPROVAL_EVIDENCE}. The Case 2 amounts below are analytical "
            "only and are not the unqualified final board figure."
        ),
    )
    wn.merge_cells(start_row=rn + 28, start_column=1, end_row=rn + 28, end_column=3)
    wn.cell(rn + 28, 1).alignment = Alignment(wrap_text=True, vertical="top")
    wn.row_dimensions[rn + 28].height = 54
    wn.cell(rn + 29, 1, "Case 2 Kiln USD (conditional)")
    wn.cell(rn + 29, 2, f"={br_case2_kiln}")
    wn.cell(rn + 30, 1, "Case 2 consolidated total USD (conditional)")
    wn.cell(rn + 30, 2, f"={br_case2_total}")
    wn.cell(rn + 31, 1, "Case 2 label")
    wn.cell(rn + 31, 2, f"={a_cutoff_label_plain}")

    wn.cell(rn + 33, 1, "Board reporting note").font = HEADER_FONT
    wn.cell(
        rn + 34,
        1,
        (
            '="Board reporting note: As of the Q3 2025 close in early October 2025, '
            f'Q3 2025 consolidated net revenue is reported at "&TEXT({br_total},"{TEXT_USD}")&'
            '". Brine & Barrel EU revenue is translated at the governing EUR/USD rate of "&'
            f'TEXT({a_fx_plain},"0.0000")&" using sum-then-round treatment. Kiln Spice Works '
            f'is included using the "&{a_method_plain}&" ownership method under '
            f'kiln_spice_acquisition_memo.txt. "&{a_approval_plain}&" Ownership-date and '
            'post-close billing support remain open follow-ups."'
        ),
    )
    wn.merge_cells(start_row=rn + 34, start_column=1, end_row=rn + 34, end_column=3)
    wn.cell(rn + 34, 1).alignment = Alignment(wrap_text=True, vertical="top")
    wn.row_dimensions[rn + 34].height = 78

    wn.cell(rn + 36, 1, "Cited sources").font = HEADER_FONT
    wn.cell(rn + 37, 1, "vesper_culinary_brands_brand_ledger.csv")
    wn.cell(rn + 38, 1, "brine_barrel_fx_and_period_reference.txt")
    wn.cell(rn + 39, 1, "kiln_spice_acquisition_memo.txt")

    fmt_usd(
        wn,
        f"B{rn + 15}",
        f"B{rn + 16}",
        f"B{rn + 24}",
        f"B{rn + 25}",
        f"B{rn + 29}",
        f"B{rn + 30}",
    )
    wn.column_dimensions["A"].width = 48
    wn.column_dimensions["B"].width = 52
    wn.column_dimensions["C"].width = 42
    wn.freeze_panes = "A6"

    board_note_text = (
        f"Board reporting note: As of the Q3 2025 close in early October 2025, "
        f"Q3 2025 consolidated net revenue is reported at ${total:,.2f}. "
        f"Brine & Barrel EU revenue is translated at the governing EUR/USD rate of "
        f"{EUR_USD} using sum-then-round treatment. Kiln Spice Works is included using the "
        f"{SELECTED_METHOD} ownership method under kiln_spice_acquisition_memo.txt. "
        f"{APPROVAL_EVIDENCE} Ownership-date and post-close billing support remain open "
        f"follow-ups."
    )
    formula_cache["Notes"] = {
        f"B{rn + 14}": SELECTED_METHOD,
        f"B{rn + 15}": kiln_base,
        f"B{rn + 16}": total,
        f"B{rn + 17}": APPROVAL_EVIDENCE,
        f"B{rn + 18}": OWNERSHIP_MODELING_BASIS,
        f"B{rn + 19}": OWNERSHIP_EVIDENCE_STATUS,
        f"B{rn + 20}": POST_CLOSE_BILLING_SUPPORT,
        f"B{rn + 24}": kiln_prorated,
        f"B{rn + 25}": total_proration,
        f"B{rn + 29}": kiln_alt,
        f"B{rn + 30}": total_alt,
        f"B{rn + 31}": CUTOFF_LABEL,
        f"A{rn + 34}": board_note_text,
    }

    # --- Recommendation ---
    wrec = wb.create_sheet("Recommendation")
    rr = write_title_block(wrec, "Recommendation", cols=2)
    wrec.cell(rr, 1, "Board recommendation").font = HEADER_FONT
    wrec.cell(rr + 2, 1, "Consolidated Q3 2025 net revenue — selected final (USD)")
    wrec.cell(rr + 2, 2, f"={br_total}")
    fmt_usd(wrec, f"B{rr + 2}")

    wrec.cell(rr + 3, 1, "Selected final Kiln ownership method")
    wrec.cell(rr + 3, 2, f"={a_method_plain}")
    wrec.cell(rr + 4, 1, "Approval evidence in supplied packet")
    wrec.cell(rr + 4, 2, f"={a_approval_plain}")
    wrec.cell(rr + 5, 1, "Conditional cutoff total (not final board figure)")
    wrec.cell(rr + 5, 2, f"={br_case2_total}")
    fmt_usd(wrec, f"B{rr + 5}")
    wrec.cell(rr + 6, 1, "Conditional cutoff label")
    wrec.cell(rr + 6, 2, f"={a_cutoff_label_plain}")
    wrec.cell(rr + 7, 1, "Ownership Date Evidence Status")
    wrec.cell(rr + 7, 2, f"={a_ownership_evidence_plain}")
    wrec.cell(rr + 8, 1, "Post-Close Billing Support")
    wrec.cell(rr + 8, 2, f"={a_billing_support_plain}")

    wrec.cell(rr + 10, 1, "Board reporting note").font = HEADER_FONT
    wrec.cell(
        rr + 11,
        1,
        (
            '="Board reporting note: As of the Q3 2025 close in early October 2025, '
            f'Q3 2025 consolidated net revenue is reported at "&TEXT({br_total},"{TEXT_USD}")&'
            '". Brine & Barrel EU revenue is translated at the governing EUR/USD rate of "&'
            f'TEXT({a_fx_plain},"0.0000")&" using sum-then-round treatment. Kiln Spice Works '
            f'is included using the "&{a_method_plain}&" ownership method under '
            f'kiln_spice_acquisition_memo.txt. "&{a_approval_plain}&" Ownership-date and '
            'post-close billing support remain open follow-ups."'
        ),
    )
    wrec.merge_cells(start_row=rr + 11, start_column=1, end_row=rr + 11, end_column=2)
    wrec.cell(rr + 11, 1).alignment = Alignment(wrap_text=True, vertical="top")
    wrec.row_dimensions[rr + 11].height = 78

    wrec.cell(rr + 13, 1, "Recommendation text").font = HEADER_FONT
    wrec.cell(
        rr + 14,
        1,
        (
            "As of the Q3 2025 close in early October 2025, use the brand-level rollup "
            "selected-final consolidated total for the board packet. Apply status, period, and "
            "ownership filters from the governing sources before rolling revenue. Brine & "
            "Barrel EU is converted at 1.0874 USD per EUR per "
            "brine_barrel_fx_and_period_reference.txt (not the 1.1025 spot aside), using "
            "sum-then-round treatment. Kiln Spice Works follows kiln_spice_acquisition_memo.txt: "
            "50/92 proration of post-exclusion full-quarter Kiln gross is the selected final "
            "board method. A transaction-date cutoff from "
            f"{KILN_CLOSE} is shown only as a Conditional transaction-date cutoff "
            f"scenario — {APPROVAL_EVIDENCE} — and is not the unqualified final consolidated "
            "revenue. Ownership-date confirmation against Schedule 2.1 and bank wire "
            "confirmation, and post-close billing shipment/service tie-out, remain open "
            "follow-ups because those records are not in the packet. Do not adopt the "
            "sales-ops full-quarter Kiln note, do not mix ownership methods, and do not treat "
            "void or cancelled rows as Q3 revenue."
        ),
    )
    wrec.merge_cells(start_row=rr + 14, start_column=1, end_row=rr + 14, end_column=2)
    wrec.cell(rr + 14, 1).alignment = Alignment(wrap_text=True, vertical="top")
    wrec.row_dimensions[rr + 14].height = 130
    wrec.column_dimensions["A"].width = 52
    wrec.column_dimensions["B"].width = 72
    wrec.freeze_panes = "A6"

    formula_cache["Recommendation"] = {
        f"B{rr + 2}": total,
        f"B{rr + 3}": SELECTED_METHOD,
        f"B{rr + 4}": APPROVAL_EVIDENCE,
        f"B{rr + 5}": total_alt,
        f"B{rr + 6}": CUTOFF_LABEL,
        f"B{rr + 7}": OWNERSHIP_EVIDENCE_STATUS,
        f"B{rr + 8}": POST_CLOSE_BILLING_SUPPORT,
        f"A{rr + 11}": board_note_text,
    }

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    cache_xlsx_formula_values(OUT, formula_cache)
    sanitize_xlsx(OUT)

    print(f"Wrote {OUT}")
    print(f"Hearth USD: {hearth_usd}")
    print(f"Brine EUR gross: {brine_eur_gross} -> USD {brine_usd} (sum-then-round)")
    print(f"Brine line-round control (not governing): {brine_line_round_control}")
    print(f"Kiln full gross: {kiln_full_gross}")
    print(f"Kiln cutoff gross: {kiln_cutoff_gross}")
    print(f"Selected method: {SELECTED_METHOD}")
    print(f"Kiln base-case USD: {kiln_base}")
    print(f"TOTAL base case: {total}")
    print(f"Case1 50/92 Kiln / Total: {kiln_prorated} / {total_proration}")
    print(f"Case2 cutoff Kiln / Total: {kiln_alt} / {total_alt}")
    print(f"Excluded rows: {len(exclusions)}")
    print(f"Kiln audit rows: {len(kiln_line_rows)}")


if __name__ == "__main__":
    main()
