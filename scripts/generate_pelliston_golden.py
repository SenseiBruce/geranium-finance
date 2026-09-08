#!/usr/bin/env python3
"""Generate golden self-insured WC reserve workbook for Pelliston (formula-driven)."""

from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from pelliston_constants import (
    BEGINNING_RESERVE,
    CLAIM_LARGE,
    CLAIM_SUBRO,
    DELIVERABLE,
    IBNR_FACTORS,
    RETENTION,
    SLUG,
    STALE_2025_FACTOR,
)
from sanitize_office import cache_xlsx_formula_values

ROOT = Path(__file__).resolve().parent.parent
TASK = ROOT / "tasks" / SLUG
INPUTS = TASK / "inputs"
GOLDEN = TASK / "golden"

header_font = Font(bold=True)
section_font = Font(bold=True, size=12)
CLOSED_CLAIM = "CLM-2024-0901"


def money(x: float) -> float:
    return round(float(x), 2)


def load_inventory() -> list[dict]:
    wb = load_workbook(INPUTS / "pelliston_open_claims_triangle.xlsx", data_only=True)
    ws = wb["Open Inventory"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for vals in ws.iter_rows(min_row=2, values_only=True):
        if not vals[0]:
            continue
        rows.append(dict(zip(headers, vals)))
    return rows


def load_triangle_paid() -> dict[int, float]:
    wb = load_workbook(INPUTS / "pelliston_open_claims_triangle.xlsx", data_only=True)
    out: dict[int, float] = {}
    for vals in wb["AY Triangle"].iter_rows(min_row=2, values_only=True):
        if not vals[0]:
            continue
        out[int(vals[0])] = money(vals[5])
    return out


def open_inventory_paid_by_ay(inventory: list[dict], *, exclude_closed: bool) -> dict[int, float]:
    totals: dict[int, float] = defaultdict(float)
    for r in inventory:
        claim = str(r["claim_number"])
        if exclude_closed and claim == CLOSED_CLAIM:
            continue
        totals[int(r["accident_year"])] += float(r["paid_to_date"])
    return {ay: money(v) for ay, v in totals.items()}


def calendar_2025_paid() -> float:
    total = 0.0
    with (INPUTS / "pelliston_claims_payment_ledger.csv").open(newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            if str(row["payment_date"]).startswith("2025"):
                total += float(row["amount"])
    return money(total)


def build_case_rows(inventory: list[dict]) -> list[dict]:
    out = []
    for r in inventory:
        claim = str(r["claim_number"])
        case = float(r["case_reserve"])
        paid = float(r["paid_to_date"])
        ay = int(r["accident_year"])
        note = ""
        closed = claim == CLOSED_CLAIM
        if closed:
            note = "Removed — ledger final payment 2025-12-18; closed"
            si_case = 0.0
            excess = 0.0
        elif claim == CLAIM_LARGE:
            si_case = money(min(case, RETENTION))
            excess = money(max(0.0, case - si_case))
            note = f"TPA case capped at SI retention via Assumptions"
        else:
            si_case = money(case)
            excess = 0.0
            if claim == CLAIM_SUBRO:
                note = "Expected subrogation not credited (cash not received)"
        out.append(
            {
                "claim": claim,
                "ay": ay,
                "claimant": r["claimant"],
                "tpa_case": money(case),
                "si_case": si_case,
                "excess": excess,
                "paid": money(paid),
                "closed": closed,
                "note": note,
            }
        )
    return out


def main() -> None:
    inventory = load_inventory()
    cases = build_case_rows(inventory)
    triangle_paid = load_triangle_paid()
    oi_paid_excl = open_inventory_paid_by_ay(inventory, exclude_closed=True)

    ay_si: dict[int, float] = defaultdict(float)
    for c in cases:
        ay_si[c["ay"]] += c["si_case"]

    ibnr_by_ay: dict[int, float] = {}
    base_by_ay: dict[int, float] = {}
    for ay in sorted(triangle_paid):
        base = money(triangle_paid[ay] + ay_si[ay])
        base_by_ay[ay] = base
        ibnr_by_ay[ay] = money(base * IBNR_FACTORS[ay])

    total_si_case = money(sum(c["si_case"] for c in cases))
    total_ibnr = money(sum(ibnr_by_ay.values()))
    ending = money(total_si_case + total_ibnr)
    paid_2025 = calendar_2025_paid()
    incurred = money(ending - BEGINNING_RESERVE + paid_2025)
    large = next(c for c in cases if c["claim"] == CLAIM_LARGE)
    large_excess = large["excess"]

    # AY2024 conflict figures (documented, not reconciled)
    oi_2024 = oi_paid_excl[2024]
    tri_2024 = triangle_paid[2024]
    conflict_delta = money(oi_2024 - tri_2024)

    wb = Workbook()

    # ---- Assumptions (central) ----
    wa = wb.active
    wa.title = "Assumptions"
    wa["A1"] = "Governing assumptions — SI WC reserve 2025-12-31"
    wa["A1"].font = section_font
    wa["A2"] = "Source: actuarial_funding_memo.txt (retention, factors, beginning reserve)"
    wa.append([])
    wa.append(["parameter", "value", "notes"])
    for c in wa[4]:
        c.font = header_font
    params = [
        ("per_occurrence_retention", RETENTION, "SI retention cap"),
        ("beginning_si_reserve_2024_12_31", BEGINNING_RESERVE, "Audited opening balance"),
        ("ibnr_factor_ay2022", IBNR_FACTORS[2022], "Memo factor"),
        ("ibnr_factor_ay2023", IBNR_FACTORS[2023], "Memo factor"),
        ("ibnr_factor_ay2024", IBNR_FACTORS[2024], "Memo factor"),
        ("ibnr_factor_ay2025", IBNR_FACTORS[2025], "Memo factor — not superseded 0.40"),
        ("superseded_ay2025_factor", STALE_2025_FACTOR, "Do not apply"),
        ("closed_claim_exclude", CLOSED_CLAIM, "Case forced to zero per memo"),
        ("large_loss_claim", CLAIM_LARGE, "Retention cap applies"),
        ("subrogation_claim", CLAIM_SUBRO, "No credit until cash received"),
        ("valuation_date", "2025-12-31", "Year-end"),
    ]
    for i, (name, val, note) in enumerate(params, start=5):
        wa.cell(i, 1, name)
        cell = wa.cell(i, 2, val)
        if isinstance(val, float):
            if "factor" in name:
                cell.number_format = "0.00"
            else:
                cell.number_format = "#,##0.00"
        wa.cell(i, 3, note)

    # Named row map for formula refs
    ret_row = 5
    beg_row = 6
    factor_rows = {2022: 7, 2023: 8, 2024: 9, 2025: 10}

    # ---- Case Inventory ----
    ws = wb.create_sheet("Case Inventory")
    ws["A1"] = "Self-insured case reserves at 2025-12-31 (retention-capped)"
    ws["A1"].font = section_font
    ws["A2"] = (
        "Source values: pelliston_open_claims_triangle.xlsx Open Inventory. "
        "SI case = MIN(TPA case, Assumptions retention); closed claim forced to zero."
    )
    ws.append([])
    headers = [
        "claim_number",
        "accident_year",
        "claimant",
        "tpa_case_reserve",
        "si_case_reserve",
        "excess_above_retention",
        "oi_paid_to_date",
        "notes",
    ]
    ws.append(headers)
    for c in ws[4]:
        c.font = header_font
    start = 5
    for i, c in enumerate(cases):
        r = start + i
        ws.cell(r, 1, c["claim"])
        ws.cell(r, 2, c["ay"])
        ws.cell(r, 3, c["claimant"])
        ws.cell(r, 4, c["tpa_case"]).number_format = "#,##0.00"
        # Formula: closed claim → 0; else MIN(TPA, retention) / excess above retention
        ws.cell(
            r,
            5,
            f'=IF(A{r}=Assumptions!$B$12,0,MIN(D{r},Assumptions!$B${ret_row}))',
        )
        ws.cell(r, 5).number_format = "#,##0.00"
        ws.cell(
            r,
            6,
            f"=IF(A{r}=Assumptions!$B$12,0,MAX(0,D{r}-E{r}))",
        )
        ws.cell(r, 6).number_format = "#,##0.00"
        ws.cell(r, 7, c["paid"]).number_format = "#,##0.00"
        ws.cell(r, 8, c["note"])
    end = start + len(cases) - 1
    tot = end + 1
    ws.cell(tot, 3, "Totals").font = Font(bold=True)
    ws.cell(tot, 5, f"=SUM(E{start}:E{end})")
    ws.cell(tot, 5).number_format = "#,##0.00"
    ws.cell(tot, 6, f"=SUM(F{start}:F{end})")
    ws.cell(tot, 6).number_format = "#,##0.00"

    # ---- IBNR by AY ----
    wi = wb.create_sheet("IBNR by AY")
    wi["A1"] = "Accident-year IBNR — Triangle cumulative paid + SI case × memo factor"
    wi["A1"].font = section_font
    wi["A2"] = (
        "Cumulative paid from AY Triangle (cumulative_paid_as_of_2025_12). "
        "SI case from Case Inventory via SUMIF. Factors from Assumptions / actuarial_funding_memo.txt."
    )
    wi.append([])
    wi.append(
        [
            "accident_year",
            "cumulative_paid_triangle",
            "si_case_reserve",
            "ibnr_base",
            "ibnr_factor",
            "ibnr",
            "paid_source",
        ]
    )
    for c in wi[4]:
        c.font = header_font
    ay_rows: dict[int, int] = {}
    r = 5
    for ay in sorted(triangle_paid):
        wi.cell(r, 1, ay)
        wi.cell(r, 2, triangle_paid[ay]).number_format = "#,##0.00"
        wi.cell(r, 3, f"=SUMIF('Case Inventory'!B$5:B${end},A{r},'Case Inventory'!E$5:E${end})")
        wi.cell(r, 3).number_format = "#,##0.00"
        wi.cell(r, 4, f"=B{r}+C{r}")
        wi.cell(r, 4).number_format = "#,##0.00"
        wi.cell(r, 5, f"=Assumptions!B{factor_rows[ay]}")
        wi.cell(r, 5).number_format = "0.00"
        wi.cell(r, 6, f"=ROUND(D{r}*E{r},2)")
        wi.cell(r, 6).number_format = "#,##0.00"
        wi.cell(r, 7, "AY Triangle cumulative_paid_as_of_2025_12")
        ay_rows[ay] = r
        r += 1
    ibnr_end = r - 1
    ibnr_total_row = r
    wi.cell(r, 1, "Total IBNR")
    wi.cell(r, 1).font = Font(bold=True)
    wi.cell(r, 6, f"=SUM(F5:F{ibnr_end})")
    wi.cell(r, 6).number_format = "#,##0.00"
    wi.cell(r + 2, 1, "Superseded factor not used")
    wi.cell(r + 2, 2, f"=Assumptions!B11")
    wi.cell(r + 2, 2).number_format = "0.00"
    wi.cell(r + 2, 3, "Mid-year 2025 exhibit — do not apply")

    # ---- Excess Bridge ----
    we = wb.create_sheet("Excess Bridge")
    we["A1"] = "Excess / ceded recoverable (not SI liability)"
    we["A1"].font = section_font
    we["A2"] = "Gross TPA case vs Assumptions retention; excess = gross − retention amount"
    we.append([])
    we.append(
        [
            "claim_number",
            "gross_tpa_case",
            "retention_amount",
            "si_case_capped",
            "excess_ceded_recoverable",
            "rule",
        ]
    )
    for c in we[4]:
        c.font = header_font
    # Locate large-loss row on Case Inventory
    large_row = next(start + i for i, c in enumerate(cases) if c["claim"] == CLAIM_LARGE)
    we["A5"] = CLAIM_LARGE
    we["B5"] = f"='Case Inventory'!D{large_row}"
    we["B5"].number_format = "#,##0.00"
    we["C5"] = f"=Assumptions!$B${ret_row}"
    we["C5"].number_format = "#,##0.00"
    we["D5"] = f"=MIN(B5,C5)"
    we["D5"].number_format = "#,##0.00"
    we["E5"] = f"=MAX(0,B5-C5)"
    we["E5"].number_format = "#,##0.00"
    we["F5"] = "actuarial_funding_memo.txt — SI retention; not SI liability"

    # ---- Reserve Rollforward ----
    wr = wb.create_sheet("Reserve Rollforward")
    wr["A1"] = "Self-insured WC reserve rollforward to 2025-12-31"
    wr["A1"].font = section_font
    wr.append([])
    wr.append(["line", "amount", "source"])
    for c in wr[3]:
        c.font = header_font
    wr["A4"] = "Beginning SI reserve (2024-12-31)"
    wr["B4"] = f"=Assumptions!B{beg_row}"
    wr["B4"].number_format = "#,##0.00"
    wr["C4"] = "Assumptions / actuarial_funding_memo.txt"
    wr["A5"] = "Net paid (calendar 2025, voids netted)"
    wr["B5"] = paid_2025
    wr["B5"].number_format = "#,##0.00"
    wr["C5"] = "pelliston_claims_payment_ledger.csv"
    wr["A6"] = "Net incurred (plug)"
    wr["B6"] = "=ROUND(B8+B5-B4,2)"
    wr["B6"].number_format = "#,##0.00"
    wr["C6"] = "Identity: ending − beginning + paid"
    wr["A7"] = "SI case reserves (capped)"
    wr["B7"] = f"='Case Inventory'!E{tot}"
    wr["B7"].number_format = "#,##0.00"
    wr["C7"] = "Case Inventory total"
    wr["A8"] = "Ending SI reserve (case + IBNR)"
    wr["B8"] = f"=ROUND(B7+'IBNR by AY'!F{ibnr_total_row},2)"
    wr["B8"].number_format = "#,##0.00"
    wr["C8"] = "Case Inventory + IBNR by AY"

    # ---- Notes / methodology ----
    wn = wb.create_sheet("Notes")
    wn["A1"] = "Methodology and source conflicts"
    wn["A1"].font = section_font
    notes = [
        (
            "IBNR paid source",
            "actuarial_funding_memo.txt requires IBNR on cumulative paid_to_date + SI case. "
            "Cumulative paid is taken from pelliston_open_claims_triangle.xlsx AY Triangle "
            "(cumulative_paid_as_of_2025_12), which covers the full accident-year population "
            "including closed claims. SI case comes from Open Inventory after retention caps.",
        ),
        (
            "Mature years AY2022 / AY2023",
            f"Open Inventory open-claim paid sums "
            f"(${oi_paid_excl[2022]:,.2f} / ${oi_paid_excl[2023]:,.2f}) are below Triangle "
            f"cumulative paid (${triangle_paid[2022]:,.2f} / ${triangle_paid[2023]:,.2f}). "
            "Using only Open Inventory paid would understate the IBNR base on mature years. "
            "Triangle is the selected AY-level paid source.",
        ),
        (
            "AY2024 paid conflict",
            f"Open Inventory open-claim paid excluding {CLOSED_CLAIM} = ${oi_2024:,.2f}. "
            f"AY Triangle cumulative_paid_as_of_2025_12 for AY2024 = ${tri_2024:,.2f}. "
            f"Open-subset paid exceeds AY-level cumulative by ${conflict_delta:,.2f}. "
            "If both were literal cumulative figures for the same population, an open subset "
            "cannot exceed the AY total — the sources cannot both be correct as stated. "
            "No bridge is provided in the source workbook. Selected source for IBNR: AY Triangle "
            f"(${tri_2024:,.2f}). Effect: IBNR base uses Triangle paid + retention-capped SI case; "
            "Open Inventory paid is not used in the IBNR base for AY2024.",
        ),
        (
            "Sources cited",
            "actuarial_funding_memo.txt; pelliston_open_claims_triangle.xlsx (AY Triangle and "
            "Open Inventory); pelliston_claims_payment_ledger.csv (calendar net paid).",
        ),
        (
            "Memo treatments applied",
            f"{CLAIM_LARGE} capped at retention; {CLOSED_CLAIM} case zeroed; "
            f"{CLAIM_SUBRO} subrogation not credited; void pair excluded from net paid; "
            f"AY2025 factor {IBNR_FACTORS[2025]:.2f} (not {STALE_2025_FACTOR:.2f}).",
        ),
    ]
    wn["A3"] = "topic"
    wn["B3"] = "documentation"
    wn["A3"].font = header_font
    wn["B3"].font = header_font
    for i, (topic, body) in enumerate(notes, start=4):
        wn.cell(i, 1, topic)
        cell = wn.cell(i, 2, body)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        wn.row_dimensions[i].height = 60
    wn.column_dimensions["A"].width = 28
    wn.column_dimensions["B"].width = 96

    # ---- Controller Recommendation ----
    wc = wb.create_sheet("Controller Recommendation")
    wc["A1"] = "Controller recommendation — SI WC reserve for audit / bank pack"
    wc["A1"].font = section_font
    wc["A3"] = "Ending SI liability to book"
    wc["B3"] = "='Reserve Rollforward'!B8"
    wc["B3"].number_format = "#,##0.00"
    wc["A4"] = "Excess / ceded recoverable (disclosure)"
    wc["B4"] = "='Excess Bridge'!E5"
    wc["B4"].number_format = "#,##0.00"
    wc["A5"] = "Narrative"
    wc["A6"] = (
        f"Book self-insured workers compensation liability equal to the ending reserve on "
        f"Reserve Rollforward (SI case from Case Inventory plus IBNR from IBNR by AY). Cap "
        f"{CLAIM_LARGE} at the Assumptions retention per actuarial_funding_memo.txt and "
        f"disclose Excess Bridge recoverable as excess-carrier recoverable, not SI liability. "
        f"Do not credit expected subrogation on {CLAIM_SUBRO} until cash is received. Apply "
        f"AY2025 IBNR factor from Assumptions ({IBNR_FACTORS[2025]:.2f}), not the superseded "
        f"{STALE_2025_FACTOR:.2f}. Remove {CLOSED_CLAIM} from case inventory per December "
        f"ledger closure. AY IBNR uses Triangle cumulative paid; see Notes for the AY2024 "
        f"Open Inventory vs Triangle paid conflict."
    )
    wc["A6"].alignment = Alignment(wrap_text=True, vertical="top")
    wc.merge_cells("A6:D10")
    wc["A12"] = "Prohibited treatments not used"
    wc["A13"] = (
        "Did not credit unpaid subrogation; did not carry uncapped TPA case on the large loss "
        "as SI liability; did not apply the mid-year 0.40 factor to AY2025; did not hide the "
        "AY2024 paid conflict between Open Inventory and AY Triangle."
    )
    wc["A13"].alignment = Alignment(wrap_text=True)
    wc.merge_cells("A13:D15")

    for sheet in wb.worksheets:
        for col in sheet.columns:
            letter = get_column_letter(col[0].column)
            if sheet.title == "Notes" and letter == "B":
                continue
            width = min(48, max(12, max(len(str(c.value or "")) for c in col[:30]) + 2))
            sheet.column_dimensions[letter].width = width

    GOLDEN.mkdir(parents=True, exist_ok=True)
    for old in GOLDEN.glob("*.xlsx"):
        old.unlink()
    out = GOLDEN / DELIVERABLE
    wb.save(out)

    formula_cache = {
        "Case Inventory": {
            **{
                f"E{start + i}": c["si_case"]
                for i, c in enumerate(cases)
            },
            **{
                f"F{start + i}": c["excess"]
                for i, c in enumerate(cases)
            },
            f"E{tot}": total_si_case,
            f"F{tot}": money(sum(c["excess"] for c in cases)),
        },
        "IBNR by AY": {
            **{f"C{ay_rows[ay]}": money(ay_si[ay]) for ay in ay_rows},
            **{f"D{ay_rows[ay]}": base_by_ay[ay] for ay in ay_rows},
            **{f"E{ay_rows[ay]}": IBNR_FACTORS[ay] for ay in ay_rows},
            **{f"F{ay_rows[ay]}": ibnr_by_ay[ay] for ay in ay_rows},
            f"F{ibnr_total_row}": total_ibnr,
            f"B{ibnr_total_row + 2}": STALE_2025_FACTOR,
        },
        "Excess Bridge": {
            "B5": large["tpa_case"],
            "C5": RETENTION,
            "D5": large["si_case"],
            "E5": large_excess,
        },
        "Reserve Rollforward": {
            "B4": BEGINNING_RESERVE,
            "B6": incurred,
            "B7": total_si_case,
            "B8": ending,
        },
        "Controller Recommendation": {
            "B3": ending,
            "B4": large_excess,
        },
    }
    cache_xlsx_formula_values(out, formula_cache)

    print(f"Wrote {out}")
    print(f"SI case {total_si_case:.2f}  IBNR {total_ibnr:.2f}  Ending {ending:.2f}")
    print(f"Paid 2025 {paid_2025:.2f}  Incurred {incurred:.2f}")
    print(f"IBNR by AY: {ibnr_by_ay}")
    print(f"Large excess {large_excess:.2f}")
    print(f"AY2024 conflict OI {oi_2024:.2f} vs Triangle {tri_2024:.2f} delta {conflict_delta:.2f}")


if __name__ == "__main__":
    main()
