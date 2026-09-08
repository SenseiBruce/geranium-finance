#!/usr/bin/env python3
"""Generate golden covenant headroom workbook for Northline Industrial Holdings."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

from northline_constants import (
    ADD_ACQ_COSTS,
    ADD_INVENTORY_STEPUP,
    ADD_LITIGATION,
    ADD_RESTRUCT_CAP,
    ADD_RESTRUCT_CLAIMED,
    ADD_RESTRUCT_EXCESS,
    ADD_SBC,
    ADD_SPONSOR_FEE_CAP,
    ADD_SPONSOR_FEE_CLAIMED,
    ADD_SPONSOR_FEE_EXCESS,
    ADD_SYNERGY_RUNRATE,
    ADJUSTED_EBITDA,
    AS_OF,
    CA_TXT,
    CAPEX_BASKET_LIMIT,
    CAPEX_BASKET_USAGE,
    CAPEX_EXCLUDED_ACQ,
    CAPEX_GROSS_YTD,
    CAPEX_REMAINING,
    CASH_INTEREST_TTM,
    DEBT_HEADROOM,
    DELIVERABLE,
    ENTITY,
    FINANCE_LEASES,
    FUNDED_DEBT,
    ICR,
    INTERCOMPANY_GUARANTOR,
    LC_UNDRAWN,
    LEVERAGE,
    MAX_DEBT_AT_LIMIT,
    MAX_LEVERAGE,
    MIN_ICR,
    REPORTED_EBITDA,
    REVOLVER_DRAWN,
    SELLER_NOTE,
    SLUG,
    TERM_LOAN_A,
    TTM_LABEL,
)
from sanitize_office import cache_xlsx_formula_values, sanitize_xlsx

TASK = Path(__file__).resolve().parent.parent / "tasks" / SLUG
OUT = TASK / "golden" / DELIVERABLE

header_fill = PatternFill("solid", fgColor="37474F")
header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
title_font = Font(bold=True, size=14, name="Calibri")
section_font = Font(bold=True, size=11, name="Calibri")
thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
pass_fill = PatternFill("solid", fgColor="C8E6C9")
fail_fill = PatternFill("solid", fgColor="FFCDD2")
money_fmt = '#,##0.00'
ratio_fmt = '0.0000'
pct_fmt = '0.00%'


def style_header(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin


def money(x: float) -> float:
    return round(float(x), 2)


def main() -> None:
    wb = Workbook()

    # ===== Funded Debt =====
    ws = wb.active
    ws.title = "Funded_Debt"
    ws["A1"] = f"{ENTITY} — Consolidated Funded Debt"
    ws["A1"].font = title_font
    ws["A2"] = f"Test Date: {AS_OF}"
    ws["A3"] = f"Governing source: {CA_TXT} (Consolidated Funded Debt definition)"

    headers = ["Component", "Amount", "Include?", "Basis"]
    for i, h in enumerate(headers, 1):
        ws.cell(5, i, h)
    style_header(ws, 5, 4)

    debt_rows = [
        ("Term Loan A outstanding", TERM_LOAN_A, "Yes", "CA §1.01 — Term Loans"),
        ("Revolving Loans drawn", REVOLVER_DRAWN, "Yes", "CA §1.01 — drawn Revolving Loans only"),
        ("Finance Lease Obligations", FINANCE_LEASES, "Yes", "CA §1.01 Finance Lease Obligations"),
        ("Bolt & Die seller note", SELLER_NOTE, "Yes", "CA §1.01 expressly includes seller note"),
        (
            "Intercompany note — Northline Holdco LLC",
            INTERCOMPANY_GUARANTOR,
            "No",
            "Excluded: Indebtedness owing to Guarantor / Loan Party",
        ),
        (
            "Standby LCs undrawn face",
            LC_UNDRAWN,
            "No",
            "Excluded: undrawn letters of credit",
        ),
    ]
    for r, (comp, amt, inc, basis) in enumerate(debt_rows, 6):
        ws.cell(r, 1, comp)
        ws.cell(r, 2, money(amt)).number_format = money_fmt
        ws.cell(r, 3, inc)
        ws.cell(r, 4, basis)
        for c in range(1, 5):
            ws.cell(r, c).border = thin

    ws["A13"] = "Consolidated Funded Debt"
    ws["A13"].font = section_font
    # Live formula: sum of included rows only (B6:B9)
    ws["B13"] = "=B6+B7+B8+B9"
    ws["B13"].number_format = money_fmt
    ws["B13"].font = section_font
    ws["C13"] = "Yes"
    ws["D13"] = "Sum of included components (excludes rows marked No)"

    ws["A15"] = "Stale Draft_Q1 balances were not used; Facilities tab as-of Test Date governs."
    ws["A16"] = f"Control total (rounded): {FUNDED_DEBT:,.2f}"

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 62

    # ===== Adjusted EBITDA =====
    ws2 = wb.create_sheet("Adjusted_EBITDA")
    ws2["A1"] = f"{ENTITY} — Consolidated Adjusted EBITDA"
    ws2["A1"].font = title_font
    ws2["A2"] = TTM_LABEL
    ws2["A3"] = f"Governing source: {CA_TXT} (Consolidated Adjusted EBITDA definition)"

    for i, h in enumerate(
        ["Line", "Borrower amount", "Lender-allowed", "Disposition", "Agreement basis"], 1
    ):
        ws2.cell(5, i, h)
    style_header(ws2, 5, 5)

    ebitda_rows = [
        (
            "Reported Consolidated EBITDA (borrower pack)",
            REPORTED_EBITDA,
            REPORTED_EBITDA,
            "Base",
            "Starting Consolidated EBITDA",
        ),
        (
            "Non-cash stock-based compensation",
            ADD_SBC,
            ADD_SBC,
            "Allow",
            "CA clause (a) non-cash SBC",
        ),
        (
            "Dayton restructuring (Permitted Restructuring Charges)",
            ADD_RESTRUCT_CLAIMED,
            ADD_RESTRUCT_CAP,
            "Cap",
            f"CA clause (b) cap $2,000,000 TTM; excess {ADD_RESTRUCT_EXCESS:,.2f} disallowed",
        ),
        (
            "Bolt & Die transaction costs",
            ADD_ACQ_COSTS,
            ADD_ACQ_COSTS,
            "Allow",
            "CA clause (c) Permitted Acquisition fees within 12 months of close",
        ),
        (
            "Bolt & Die run-rate synergies (projected)",
            ADD_SYNERGY_RUNRATE,
            0.0,
            "Reject",
            "CA negative list (i) — projected/run-rate synergies not realized",
        ),
        (
            "Sponsor management fee — Ridgepath Capital",
            ADD_SPONSOR_FEE_CLAIMED,
            ADD_SPONSOR_FEE_CAP,
            "Cap",
            f"CA clause (e) $500,000 TTM basket; excess {ADD_SPONSOR_FEE_EXCESS:,.2f} disallowed",
        ),
        (
            "Inventory step-up amortization",
            ADD_INVENTORY_STEPUP,
            ADD_INVENTORY_STEPUP,
            "Allow",
            "CA clause (d) non-cash purchase accounting",
        ),
        (
            "Customer warranty dispute settlement (Atlas Forge)",
            ADD_LITIGATION,
            0.0,
            "Reject",
            "Ordinary Course Warranty Matter — CA negative list (ii)",
        ),
    ]
    for r, (line, borr, lend, disp, basis) in enumerate(ebitda_rows, 6):
        ws2.cell(r, 1, line)
        ws2.cell(r, 2, money(borr)).number_format = money_fmt
        ws2.cell(r, 3, money(lend)).number_format = money_fmt
        ws2.cell(r, 4, disp)
        ws2.cell(r, 5, basis)
        for c in range(1, 6):
            ws2.cell(r, c).border = thin

    ws2["A15"] = "Consolidated Adjusted EBITDA"
    ws2["A15"].font = section_font
    ws2["C15"] = "=C6+C7+C8+C9+C10+C11+C12+C13"
    ws2["C15"].number_format = money_fmt
    ws2["C15"].font = section_font
    ws2["D15"] = "Lender-adjusted"
    ws2["E15"] = f"Control: {ADJUSTED_EBITDA:,.2f}"

    ws2["A17"] = "Rejected / capped amounts (not in Adjusted EBITDA)"
    ws2["A17"].font = section_font
    ws2["A18"] = "Restructuring excess above $2,000,000 cap"
    ws2["C18"] = ADD_RESTRUCT_EXCESS
    ws2["C18"].number_format = money_fmt
    ws2["A19"] = "Run-rate synergies rejected"
    ws2["C19"] = ADD_SYNERGY_RUNRATE
    ws2["C19"].number_format = money_fmt
    ws2["A20"] = "Sponsor fee excess above $500,000 basket"
    ws2["C20"] = ADD_SPONSOR_FEE_EXCESS
    ws2["C20"].number_format = money_fmt
    ws2["A21"] = "Ordinary-course warranty settlement rejected"
    ws2["C21"] = ADD_LITIGATION
    ws2["C21"].number_format = money_fmt

    for col, w in enumerate([52, 16, 16, 12, 70], 1):
        ws2.column_dimensions[get_column_letter(col)].width = w

    # ===== Covenant Tests =====
    ws3 = wb.create_sheet("Covenant_Tests")
    ws3["A1"] = f"{ENTITY} — Financial covenant tests"
    ws3["A1"].font = title_font
    ws3["A2"] = f"Test Date {AS_OF} | {TTM_LABEL}"

    for i, h in enumerate(["Metric", "Value", "Limit", "Headroom / cushion", "Result"], 1):
        ws3.cell(4, i, h)
    style_header(ws3, 4, 5)

    # Row 5 Total Leverage
    ws3["A5"] = "Total Leverage Ratio"
    ws3["B5"] = f"=Funded_Debt!B13/Adjusted_EBITDA!C15"
    ws3["B5"].number_format = ratio_fmt
    ws3["C5"] = MAX_LEVERAGE
    ws3["C5"].number_format = '0.00'
    ws3["D5"] = f"=C5-B5"
    ws3["D5"].number_format = ratio_fmt
    ws3["E5"] = "PASS"
    ws3["E5"].fill = pass_fill

    # Row 6 Interest Coverage
    ws3["A6"] = "Interest Coverage Ratio"
    ws3["B6"] = f"=Adjusted_EBITDA!C15/{CASH_INTEREST_TTM}"
    ws3["B6"].number_format = ratio_fmt
    ws3["C6"] = MIN_ICR
    ws3["C6"].number_format = '0.00'
    ws3["D6"] = f"=B6-C6"
    ws3["D6"].number_format = ratio_fmt
    ws3["E6"] = "PASS"
    ws3["E6"].fill = pass_fill

    # CapEx section
    ws3["A8"] = "CapEx basket (calendar 2026 YTD)"
    ws3["A8"].font = section_font
    for i, h in enumerate(["Item", "Amount", "Notes"], 1):
        ws3.cell(9, i, h)
    style_header(ws3, 9, 3)

    ws3["A10"] = "Gross CapEx YTD"
    ws3["B10"] = CAPEX_GROSS_YTD
    ws3["B10"].number_format = money_fmt
    ws3["C10"] = "From CapEx_YTD tab in debt schedule"

    ws3["A11"] = "Less: Permitted Acquisition CapEx (Bolt & Die integration)"
    ws3["B11"] = CAPEX_EXCLUDED_ACQ
    ws3["B11"].number_format = money_fmt
    ws3["C11"] = "CA §1.01 Permitted Acquisition CapEx / §6.12 carve-out"

    ws3["A12"] = "CapEx Basket Usage"
    ws3["A12"].font = section_font
    ws3["B12"] = "=B10-B11"
    ws3["B12"].number_format = money_fmt
    ws3["C12"] = "Gross minus carve-out"

    ws3["A13"] = "CapEx Basket Limit"
    ws3["B13"] = CAPEX_BASKET_LIMIT
    ws3["B13"].number_format = money_fmt
    ws3["C13"] = "CA §6.12 — $8,500,000 calendar year (not prorated)"

    ws3["A14"] = "Remaining basket"
    ws3["B14"] = "=B13-B12"
    ws3["B14"].number_format = money_fmt
    ws3["E14"] = "PASS"
    ws3["E14"].fill = pass_fill

    ws3["A16"] = "Supporting inputs"
    ws3["A16"].font = section_font
    ws3["A17"] = "Cash Interest Expense (TTM)"
    ws3["B17"] = CASH_INTEREST_TTM
    ws3["B17"].number_format = money_fmt
    ws3["C17"] = "Excludes commitment fees and LC fees per CA §1.01"

    ws3["A18"] = "Max Funded Debt at 4.50x"
    ws3["B18"] = f"=Adjusted_EBITDA!C15*{MAX_LEVERAGE}"
    ws3["B18"].number_format = money_fmt
    ws3["A19"] = "Debt dollars of headroom"
    ws3["B19"] = f"=B18-Funded_Debt!B13"
    ws3["B19"].number_format = money_fmt

    # Control labels
    ws3["A21"] = f"Control leverage {LEVERAGE} | ICR {ICR} | CapEx usage {CAPEX_BASKET_USAGE:,.2f} | remaining {CAPEX_REMAINING:,.2f}"

    for col, w in enumerate([55, 16, 55, 18, 10], 1):
        ws3.column_dimensions[get_column_letter(col)].width = w

    # ===== Exceptions =====
    ws4 = wb.create_sheet("Exceptions")
    ws4["A1"] = "Exception and disputed-item log"
    ws4["A1"].font = title_font
    ws4["A2"] = "Items where borrower treatment differs from lender-adjusted certificate"

    for i, h in enumerate(
        ["Item", "Borrower treatment", "Lender disposition", "$ impact vs borrower", "Citation"], 1
    ):
        ws4.cell(4, i, h)
    style_header(ws4, 4, 5)

    exceptions = [
        (
            "Dayton restructuring add-back",
            f"Include full {ADD_RESTRUCT_CLAIMED:,.2f}",
            f"Allow only {ADD_RESTRUCT_CAP:,.2f} under Permitted Restructuring cap",
            -ADD_RESTRUCT_EXCESS,
            f"{CA_TXT} Consolidated Adjusted EBITDA (b)",
        ),
        (
            "Bolt & Die run-rate synergies",
            f"Include {ADD_SYNERGY_RUNRATE:,.2f} projected savings",
            "Reject — not realized in consolidated results",
            -ADD_SYNERGY_RUNRATE,
            f"{CA_TXT} Consolidated Adjusted EBITDA negative list (i)",
        ),
        (
            "Ridgepath sponsor management fee",
            f"Include full {ADD_SPONSOR_FEE_CLAIMED:,.2f}",
            f"Allow only {ADD_SPONSOR_FEE_CAP:,.2f} affiliate fee basket",
            -ADD_SPONSOR_FEE_EXCESS,
            f"{CA_TXT} Consolidated Adjusted EBITDA (e) / §6.05",
        ),
        (
            "Atlas Forge warranty settlement",
            f"Include {ADD_LITIGATION:,.2f} as extraordinary litigation",
            "Reject — Ordinary Course Warranty Matter",
            -ADD_LITIGATION,
            f"{CA_TXT} Ordinary Course Warranty Matters / negative list (ii)",
        ),
        (
            "Intercompany note to Northline Holdco LLC",
            "Borrower flagged as possible Funded Debt",
            "Exclude — owing to Guarantor Loan Party",
            0.0,
            f"{CA_TXT} Consolidated Funded Debt exclusion (2)",
        ),
        (
            "Undrawn standby LC face",
            "Listed on debt schedule at face",
            "Exclude undrawn face from Funded Debt",
            0.0,
            f"{CA_TXT} Consolidated Funded Debt exclusion (1)",
        ),
        (
            "Bolt & Die integration CapEx",
            "Included in gross CapEx vs basket",
            "Carve out as Permitted Acquisition CapEx",
            -CAPEX_EXCLUDED_ACQ,
            f"{CA_TXT} §6.12 / Permitted Acquisition CapEx",
        ),
    ]
    for r, row in enumerate(exceptions, 5):
        for c, val in enumerate(row, 1):
            cell = ws4.cell(r, c, val)
            cell.border = thin
            if c == 4 and isinstance(val, (int, float)):
                cell.number_format = money_format if False else money_fmt

    for col, w in enumerate([40, 42, 48, 18, 55], 1):
        ws4.column_dimensions[get_column_letter(col)].width = w

    # ===== Recommendation =====
    ws5 = wb.create_sheet("Recommendation")
    ws5["A1"] = f"{ENTITY} — Portfolio committee recommendation"
    ws5["A1"].font = title_font
    ws5["A2"] = f"Q2 covenant certificate as of {AS_OF}"

    ws5["A4"] = "Recommendation"
    ws5["A4"].font = section_font
    ws5["A5"] = (
        "PASS — certify compliance on Total Leverage, Interest Coverage, and CapEx basket "
        "using lender-adjusted Consolidated Funded Debt and Consolidated Adjusted EBITDA. "
        "Do not accept the borrower-claimed Adjusted EBITDA figure for the certificate."
    )
    ws5["A5"].alignment = Alignment(wrap_text=True)
    ws5.merge_cells("A5:D5")
    ws5.row_dimensions[5].height = 60

    ws5["A7"] = "Covenant results (lender-adjusted)"
    ws5["A7"].font = section_font
    ws5["A8"] = "Total Leverage"
    ws5["B8"] = f"=Covenant_Tests!B5"
    ws5["B8"].number_format = ratio_fmt
    ws5["C8"] = f"vs max {MAX_LEVERAGE:.2f}x — PASS"
    ws5["A9"] = "Interest Coverage"
    ws5["B9"] = f"=Covenant_Tests!B6"
    ws5["B9"].number_format = ratio_fmt
    ws5["C9"] = f"vs min {MIN_ICR:.2f}x — PASS"
    ws5["A10"] = "CapEx Basket Usage"
    ws5["B10"] = f"=Covenant_Tests!B12"
    ws5["B10"].number_format = money_fmt
    ws5["C10"] = f"vs limit {CAPEX_BASKET_LIMIT:,.2f} — PASS (see Covenant_Tests remaining basket)"

    ws5["A12"] = "Key adjustments vs borrower pack"
    ws5["A12"].font = section_font
    ws5["A13"] = (
        f"Adjusted EBITDA reduced for: restructuring excess ${ADD_RESTRUCT_EXCESS:,.2f}; "
        f"run-rate synergies ${ADD_SYNERGY_RUNRATE:,.2f}; sponsor fee excess ${ADD_SPONSOR_FEE_EXCESS:,.2f}; "
        f"ordinary-course warranty ${ADD_LITIGATION:,.2f}. "
        f"Funded Debt includes seller note ${SELLER_NOTE:,.2f}; excludes intercompany "
        f"${INTERCOMPANY_GUARANTOR:,.2f} and undrawn LCs ${LC_UNDRAWN:,.2f}."
    )
    ws5["A13"].alignment = Alignment(wrap_text=True)
    ws5.merge_cells("A13:D13")
    ws5.row_dimensions[13].height = 75

    ws5["A15"] = "Residual monitoring"
    ws5["A15"].font = section_font
    ws5["A16"] = (
        "Watch CapEx remaining basket (~$1.22mm) if H2 growth projects accelerate; "
        "confirm no further restructuring charges that would exhaust the $2.0mm TTM add-back cap; "
        "reject any future run-rate synergy add-backs until reflected in consolidated results. "
        "Citations: credit_agreement_covenant_excerpts.txt; debt schedule Facilities/CapEx_YTD; "
        "ebitda bridge add-back codes RESTRUCT, SYNERGY, MGMT_FEE, LIT."
    )
    ws5["A16"].alignment = Alignment(wrap_text=True)
    ws5.merge_cells("A16:D16")
    ws5.row_dimensions[16].height = 75

    ws5["A18"] = "Prepared for Mid-Market Credit portfolio committee — certificate-ready figures only."
    ws5.column_dimensions["A"].width = 28
    ws5.column_dimensions["B"].width = 16
    ws5.column_dimensions["C"].width = 55
    ws5.column_dimensions["D"].width = 20

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    sanitize_xlsx(OUT)

    formula_cache = {
        "Funded_Debt": {
            "B13": FUNDED_DEBT,
        },
        "Adjusted_EBITDA": {
            "C15": ADJUSTED_EBITDA,
        },
        "Covenant_Tests": {
            "B5": LEVERAGE,
            "D5": round(MAX_LEVERAGE - LEVERAGE, 4),
            "B6": ICR,
            "D6": round(ICR - MIN_ICR, 4),
            "B12": CAPEX_BASKET_USAGE,
            "B14": CAPEX_REMAINING,
            "B18": MAX_DEBT_AT_LIMIT,
            "B19": DEBT_HEADROOM,
        },
        "Recommendation": {
            "B8": LEVERAGE,
            "B9": ICR,
            "B10": CAPEX_BASKET_USAGE,
        },
    }
    cache_xlsx_formula_values(OUT, formula_cache)
    print(f"Wrote {OUT}")
    print(f"  Funded Debt {FUNDED_DEBT:,.2f}")
    print(f"  Adj EBITDA {ADJUSTED_EBITDA:,.2f}")
    print(f"  Leverage {LEVERAGE} ICR {ICR}")
    print(f"  CapEx usage {CAPEX_BASKET_USAGE:,.2f} remain {CAPEX_REMAINING:,.2f}")


if __name__ == "__main__":
    main()
