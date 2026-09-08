#!/usr/bin/env python3
"""Generate golden workbook for sawtooth-july-renewal-desk."""

from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook

TASK = Path(__file__).resolve().parent.parent / "tasks" / "sawtooth-july-renewal-desk"
INPUTS = TASK / "inputs"
OUT = TASK / "golden" / "renewal_underwriting_review.xlsx"


def load_csv(name: str) -> list[dict[str, str]]:
    with (INPUTS / name).open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def attritional_incurred(losses: list[dict[str, str]]) -> float:
    return sum(
        float(x["paid_amount"]) + float(x["reserved_amount"])
        for x in losses
        if not x["cat_code"].strip()
    )


def loc_incurred(losses: list[dict[str, str]]) -> float:
    return attritional_incurred(losses)


def route_account(
    acct: dict[str, str],
    locs: list[dict[str, str]],
    losses: list[dict[str, str]],
) -> tuple[str, str, float, float, list[str]]:
    premium = float(acct["expiring_premium"])
    incurred = attritional_incurred(losses)
    lr = incurred / premium * 100 if premium else 0.0
    list_lr = float(acct["account_loss_ratio_36mo"])
    tiv_acct = float(acct["account_tiv"])
    tiv_loc = sum(float(x["building_tiv"]) for x in locs)
    tiv_var_pct = abs(tiv_acct - tiv_loc) / tiv_acct * 100 if tiv_acct else 0.0

    triggers: list[str] = []
    rationale_parts: list[str] = []

    if lr >= 70:
        triggers.append("decline_lr")
        rationale_parts.append(f"Section 3 decline: recomputed LR {lr:.1f}% from loss_run_36mo.csv")
    elif lr >= 45:
        triggers.append("refer_lr")
        rationale_parts.append(f"Section 3 refer: recomputed LR {lr:.1f}% from loss_run_36mo.csv")

    for loc in locs:
        loc_losses = [x for x in losses if x["location_id"] == loc["location_id"]]
        loc_total = loc_incurred(loc_losses)
        if loc_total >= 250_000:
            triggers.append("refer_loc")
            rationale_parts.append(
                f"Section 3 refer: {loc['location_id']} incurred ${loc_total:,.0f} per loss_run_36mo.csv"
            )
        if loc["mitigation_status"].startswith("pending"):
            triggers.append("refer_mit")
            rationale_parts.append(
                f"Section 5 refer: {loc['location_id']} mitigation {loc['mitigation_status']} per policy_location_schedule.csv"
            )
        if loc["coastal_wind_zone"] in ("CW-1", "CW-2"):
            triggers.append("coastal")
            rationale_parts.append(
                f"Section 5 coastal endorsement required: {loc['location_id']} zone {loc['coastal_wind_zone']}"
            )
        if loc["vacancy_flag"] == "Y" and loc["broker_confirmation"] == "no":
            triggers.append("ask_broker")
            rationale_parts.append(
                f"Section 6 Ask Broker: {loc['location_id']} vacancy without broker confirmation"
            )

    if tiv_var_pct > 3:
        triggers.append("refer_tiv")
        rationale_parts.append(
            f"Section 4 refer: TIV variance {tiv_var_pct:.1f}% between renewal_account_list.csv and policy_location_schedule.csv"
        )

    if "decline_lr" in triggers:
        disposition = "Decline"
    elif "ask_broker" in triggers:
        disposition = "Ask Broker"
    elif any(t in triggers for t in ("refer_lr", "refer_loc", "refer_mit", "refer_tiv", "coastal")):
        disposition = "Refer"
    else:
        disposition = "Quote"
        rationale_parts = ["Clean renewal path under underwriting_rules_notes.txt"]

    rationale = "; ".join(dict.fromkeys(rationale_parts))
    return disposition, rationale, lr, list_lr, triggers


def main() -> None:
    accounts = load_csv("renewal_account_list.csv")
    locations = load_csv("policy_location_schedule.csv")
    losses = load_csv("loss_run_36mo.csv")

    loc_by_acct: dict[str, list[dict[str, str]]] = defaultdict(list)
    for loc in locations:
        loc_by_acct[loc["account_id"]].append(loc)

    loss_by_acct: dict[str, list[dict[str, str]]] = defaultdict(list)
    for loss in losses:
        loss_by_acct[loss["account_id"]].append(loss)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Account_Routing"
    ws.append([
        "account_id", "named_insured", "expiring_premium", "list_loss_ratio",
        "recomputed_loss_ratio", "disposition", "rationale",
    ])

    routing_rows: list[tuple[str, str]] = []
    loss_recon_rows: list[list] = []
    location_exceptions: list[list] = []

    for acct in accounts:
        acct_id = acct["account_id"]
        locs = loc_by_acct[acct_id]
        acct_losses = loss_by_acct[acct_id]
        disposition, rationale, lr, list_lr, triggers = route_account(acct, locs, acct_losses)
        routing_rows.append((acct_id, disposition))
        ws.append([
            acct_id,
            acct["named_insured"],
            float(acct["expiring_premium"]),
            list_lr,
            round(lr, 1),
            disposition,
            rationale,
        ])
        if abs(lr - list_lr) > 5:
            loss_recon_rows.append([
                acct_id,
                list_lr,
                round(lr, 1),
                round(lr - list_lr, 1),
                "loss_run_36mo.csv vs renewal_account_list.csv",
            ])
        for loc in locs:
            if loc["coastal_wind_zone"] in ("CW-1", "CW-2"):
                location_exceptions.append([
                    loc["location_id"],
                    acct_id,
                    "Coastal wind endorsement required",
                    loc["coastal_wind_zone"],
                    "policy_location_schedule.csv",
                ])
            if loc["mitigation_status"].startswith("pending"):
                location_exceptions.append([
                    loc["location_id"],
                    acct_id,
                    "Pending mitigation",
                    loc["mitigation_status"],
                    "policy_location_schedule.csv",
                ])
            if loc["vacancy_flag"] == "Y" and loc["broker_confirmation"] == "no":
                location_exceptions.append([
                    loc["location_id"],
                    acct_id,
                    "Vacancy pending broker update",
                    "vacancy_flag=Y",
                    "policy_location_schedule.csv",
                ])
        if "refer_tiv" in triggers:
            tiv_acct = float(acct["account_tiv"])
            tiv_loc = sum(float(x["building_tiv"]) for x in locs)
            location_exceptions.append([
                acct_id,
                acct_id,
                "Account TIV does not tie to locations",
                f"variance ${tiv_acct - tiv_loc:,.0f}",
                "renewal_account_list.csv; policy_location_schedule.csv",
            ])

    ws2 = wb.create_sheet("Location_Exceptions")
    ws2.append(["location_or_account", "account_id", "exception_type", "detail", "source"])
    for row in location_exceptions:
        ws2.append(row)

    ws3 = wb.create_sheet("Loss_Reconciliation")
    ws3.append(["account_id", "list_loss_ratio", "recomputed_loss_ratio", "delta", "source"])
    for row in loss_recon_rows:
        ws3.append(row)

    ws4 = wb.create_sheet("Rules_Citations")
    ws4.append(["rule_topic", "section", "source_file"])
    citations = [
        ("Loss ratio decline/refer thresholds", "Section 3", "underwriting_rules_notes.txt"),
        ("Single-location incurred referral", "Section 3", "underwriting_rules_notes.txt"),
        ("TIV reconciliation tolerance", "Section 4", "underwriting_rules_notes.txt"),
        ("Pending mitigation referral", "Section 5", "underwriting_rules_notes.txt"),
        ("Coastal wind endorsement requirement", "Section 5", "underwriting_rules_notes.txt"),
        ("Vacancy broker confirmation", "Section 6", "underwriting_rules_notes.txt"),
        ("Cat loss exclusion from attritional LR", "Section 9", "underwriting_rules_notes.txt"),
    ]
    for row in citations:
        ws4.append(row)

    ws5 = wb.create_sheet("Summary")
    ws5["A1"] = "July renewal batch disposition summary"
    ws5["A2"] = "Accounts reviewed"
    ws5["B2"] = "=COUNTA(Account_Routing!A2:A100)-COUNTIF(Account_Routing!A2:A100,\"\")"
    ws5["A3"] = "Quote"
    ws5["B3"] = '=COUNTIF(Account_Routing!F2:F100,"Quote")'
    ws5["A4"] = "Refer"
    ws5["B4"] = '=COUNTIF(Account_Routing!F2:F100,"Refer")'
    ws5["A5"] = "Decline"
    ws5["B5"] = '=COUNTIF(Account_Routing!F2:F100,"Decline")'
    ws5["A6"] = "Ask Broker"
    ws5["B6"] = '=COUNTIF(Account_Routing!F2:F100,"Ask Broker")'
    ws5["A8"] = "Coastal locations flagged"
    ws5["B8"] = '=COUNTIF(Location_Exceptions!D2:D200,"CW-1")+COUNTIF(Location_Exceptions!D2:D200,"CW-2")'

    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Accounts: {len(accounts)}")
    from collections import Counter
    print("Dispositions:", Counter(d for _, d in routing_rows))


if __name__ == "__main__":
    main()
