#!/usr/bin/env python3
"""Generate inputs for Crowhaven Logistics insured claims triangle reserve opinion."""

from __future__ import annotations

import csv
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from crowhaven_constants import (
    AY_AGE_MONTHS,
    BEGINNING_BOOKED_RESERVE,
    CLAIM_CLOSED_STALE,
    CLAIM_CLOSED_STALE_CASE,
    CLAIM_DUP,
    CLAIM_DUP_CASE_NEW,
    CLAIM_DUP_CASE_OLD,
    CLAIM_LARGE,
    CLAIM_LARGE_CASE,
    CLAIM_LARGE_PAID,
    CLAIM_VOID,
    CUM_PAID,
    DECLINE_LR,
    EARNED_PREMIUM,
    ENTITY,
    EVAL_DATE,
    LDF,
    POLICY,
    RENEWAL_DATE,
    REFER_LR,
    SLUG,
    STALE_LDF,
    TARGET_LR,
)

TASK_INPUTS = Path(__file__).resolve().parent.parent / "tasks" / SLUG / "inputs"
rng = random.Random(20260904)
header_font = Font(bold=True)


def money(x: float) -> float:
    return round(float(x), 2)


def build_prior_diagonal(latest: float, ages: list[int], ay: int) -> dict[int, float]:
    """Back into earlier development ages with irregular emergence."""
    # ages ascending e.g. [12,24,36,48]; latest is at max age
    factors = {12: 0.42, 24: 0.68, 36: 0.86, 48: 1.0}
    if ay == 2025:
        return {12: latest}
    if ay == 2024:
        return {12: money(latest * 0.61 + rng.uniform(-800, 800)), 24: latest}
    if ay == 2023:
        p12 = money(latest * 0.48 + rng.uniform(-1200, 900))
        p24 = money(latest * 0.79 + rng.uniform(-900, 700))
        return {12: p12, 24: p24, 36: latest}
    # 2022
    p12 = money(latest * 0.41 + rng.uniform(-1500, 1100))
    p24 = money(latest * 0.67 + rng.uniform(-1000, 800))
    p36 = money(latest * 0.88 + rng.uniform(-700, 600))
    return {12: p12, 24: p24, 36: p36, 48: latest}


def write_triangle(path: Path) -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "Paid Triangle"
    headers = [
        "accident_year",
        "line_of_business",
        "dev_age_months",
        "cumulative_paid",
        "evaluation_date",
        "source_system",
    ]
    ws.append(headers)
    for c in ws[1]:
        c.font = header_font

    rows = []
    for ay, latest in CUM_PAID.items():
        ages = sorted(a for a in (12, 24, 36, 48) if a <= AY_AGE_MONTHS[ay])
        diag = build_prior_diagonal(latest, ages, ay)
        lob = "Commercial Auto" if ay != 2025 else "Mixed Auto/GL"
        for age in ages:
            rows.append(
                [
                    ay,
                    lob,
                    age,
                    money(diag[age]),
                    f"2025-{min(12, age // 12 + (ay - 2022) * 0 + 12):02d}-31"
                    if age == AY_AGE_MONTHS[ay]
                    else f"{ay + age // 12 - 1}-12-31",
                    "TPA_CUM_PAID_V3",
                ]
            )
    # Fix evaluation dates to be coherent: latest diagonal = EVAL_DATE
    for r in rows:
        if r[2] == AY_AGE_MONTHS[int(r[0])]:
            r[4] = EVAL_DATE
        else:
            # prior year-end snapshots
            matured = int(r[0]) + r[2] // 12 - 1
            r[4] = f"{matured}-12-31"
    for r in rows:
        ws.append(r)

    # Noise rows — interim quarterly snapshots that should not replace year-end diagonal
    ws.append([2024, "Commercial Auto", 18, money(CUM_PAID[2024] * 0.72), "2025-06-30", "TPA_CUM_PAID_V3"])
    ws.append([2025, "Commercial Auto", 6, money(CUM_PAID[2025] * 0.51), "2025-06-30", "TPA_CUM_PAID_V3"])

    ws2 = wb.create_sheet("AY Summary")
    ws2.append(
        [
            "accident_year",
            "maturity_months",
            "latest_cum_paid",
            "large_loss_paid_in_ay",
            "notes",
        ]
    )
    for c in ws2[1]:
        c.font = header_font
    for ay, paid in CUM_PAID.items():
        ll = CLAIM_LARGE_PAID if ay == 2023 else 0.0
        note = ""
        if ay == 2023:
            note = f"Includes {CLAIM_LARGE} paid {CLAIM_LARGE_PAID:.2f}; see actuarial memo carve-out"
        ws2.append([ay, AY_AGE_MONTHS[ay], money(paid), money(ll), note])

    ws3 = wb.create_sheet("Definitions")
    ws3["A1"] = "Crowhaven Logistics — TPA triangle field definitions"
    ws3["A1"].font = header_font
    defs = [
        ("cumulative_paid", "Indemnity + medical + ALAE paid to evaluation date, net of recoveries already received."),
        ("dev_age_months", "Months from start of accident year to evaluation date for that diagonal cell."),
        ("line_of_business", "Primary commercial auto liability or GL as coded by TPA; Mixed for current year roll-up."),
        ("source_system", "Extract job id from Sedgwick-style warehouse; V3 is production."),
        ("stale_ldf_warning", "Mid-year LDF table below was published 2025-06-30 for pricing draft only."),
        ("large_loss", f"{CLAIM_LARGE} is flagged in AY Summary; actuarial memo governs IBNR treatment."),
        ("insured_program", "Primary insured package — not a self-insured retention program."),
    ]
    ws3.append(["field", "definition"])
    for c in ws3[2]:
        c.font = header_font
    for row in defs:
        ws3.append(list(row))

    ws3.append([])
    ws3.append(["Mid-year paid LDF to ultimate (DRAFT — 2025-06-30)", "", ""])
    ws3["A12"].font = header_font
    ws3.append(["dev_age_months", "paid_ldf_to_ultimate", "status"])
    for age, f in STALE_LDF.items():
        ws3.append([age, f, "superseded_pending_year_end"])

    ws3.append([])
    ws3.append(
        [
            "Note",
            "Do not use mid-year LDFs for 2025-12-31 reserve opinion; consulting actuary year-end memo controls.",
        ]
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def open_claim_rows() -> list[dict]:
    """Build messy open inventory with known trap rows."""
    base: list[dict] = []

    # Attritional open claims — irregular cases
    auto_claims = [
        ("CLM-AUTO-2022-0188", 2022, "Commercial Auto", "R. Holcomb", 14_220.18, 4_810.55),
        ("CLM-AUTO-2022-0441", 2022, "Commercial Auto", "S. Mercado", 9_880.40, 2_110.00),
        ("CLM-AUTO-2022-0612", 2022, "Commercial Auto", "P. Nguyen", 22_640.55, 5_200.18),
        ("CLM-AUTO-2022-0790", 2022, "Commercial Auto", "L. Ortiz", 6_110.22, 1_840.55),
        ("CLM-AUTO-2023-0091", 2023, "Commercial Auto", "K. Brennan", 18_420.40, 6_110.18),
        ("CLM-AUTO-2023-0275", 2023, "Commercial Auto", "D. Singh", 31_220.18, 8_880.40),
        ("CLM-AUTO-2023-0510", 2023, "Commercial Auto", "M. Alvarez", 12_640.55, 4_220.18),
        ("CLM-AUTO-2023-0666", 2023, "Commercial Auto", "J. Whitaker", 27_110.00, 7_640.55),
        ("CLM-AUTO-2023-0822", 2023, "Commercial Auto", "T. Brooks", 8_880.40, 3_220.18),
        ("CLM-AUTO-2024-0044", 2024, "Commercial Auto", "A. Patel", 24_110.22, 12_880.40),
        ("CLM-AUTO-2024-0318", 2024, "Commercial Auto", "C. Ramirez", 16_420.18, 9_640.55),
        ("CLM-AUTO-2024-0555", 2024, "Commercial Auto", "E. Cho", 33_220.40, 18_110.18),
        ("CLM-AUTO-2024-0712", 2024, "Commercial Auto", "N. Foster", 11_880.55, 7_220.40),
        ("CLM-AUTO-2024-0880", 2024, "Commercial Auto", "B. Keller", 19_640.18, 11_410.00),
        ("CLM-AUTO-2025-0019", 2025, "Commercial Auto", "H. Quinn", 8_220.40, 14_640.55),
        ("CLM-AUTO-2025-0140", 2025, "Commercial Auto", "I. Delgado", 14_880.18, 19_220.40),
        ("CLM-AUTO-2025-0277", 2025, "Commercial Auto", "F. Okonkwo", 6_640.55, 11_110.22),
        ("CLM-AUTO-2025-0391", 2025, "Commercial Auto", "G. Marsh", 21_110.00, 16_880.40),
        ("CLM-AUTO-2025-0502", 2025, "Commercial Auto", "Y. Tran", 9_420.18, 9_640.55),
    ]
    gl_claims = [
        ("CLM-GL-2022-0033", 2022, "General Liability", "Warehouse slip", 7_220.18, 2_840.55),
        ("CLM-GL-2023-0118", 2023, "General Liability", "Dock damage", 15_640.40, 5_110.18),
        ("CLM-GL-2023-0290", 2023, "General Liability", "Customer injury", 28_880.55, 8_220.40),
        ("CLM-GL-2024-0077", 2024, "General Liability", "Forklift strike", 12_110.18, 8_640.55),
        ("CLM-GL-2024-0215", 2024, "General Liability", "Cargo spill", 18_420.40, 10_880.18),
        ("CLM-GL-2024-0388", 2024, "General Liability", "Premise fall", 9_880.55, 6_220.40),
        ("CLM-GL-2025-0021", 2025, "General Liability", "Visitor injury", 11_220.18, 12_640.55),
        ("CLM-GL-2025-0088", 2025, "General Liability", "Property damage", 16_640.40, 14_110.18),
        ("CLM-GL-2025-0205", 2025, "General Liability", "Product allegation", 22_880.55, 18_220.40),
    ]

    as_of = "2025-12-28"
    for claim, ay, lob, claimant, paid, case in auto_claims + gl_claims:
        base.append(
            {
                "claim_number": claim,
                "accident_year": ay,
                "line_of_business": lob,
                "claimant_or_desc": claimant,
                "status": "Open",
                "paid_to_date": f"{money(paid):.2f}",
                "case_reserve": f"{money(case):.2f}",
                "as_of_date": as_of,
                "large_loss_flag": "N",
                "tpa_note": "",
            }
        )

    # Large loss
    base.append(
        {
            "claim_number": CLAIM_LARGE,
            "accident_year": 2023,
            "line_of_business": "Commercial Auto",
            "claimant_or_desc": "Multi-vehicle I-65",
            "status": "Open",
            "paid_to_date": f"{money(CLAIM_LARGE_PAID):.2f}",
            "case_reserve": f"{money(CLAIM_LARGE_CASE):.2f}",
            "as_of_date": as_of,
            "large_loss_flag": "Y",
            "tpa_note": "Severity committee; see actuarial memo",
        }
    )

    # Closed but still Open
    base.append(
        {
            "claim_number": CLAIM_CLOSED_STALE,
            "accident_year": 2024,
            "line_of_business": "Commercial Auto",
            "claimant_or_desc": "C. Delgado",
            "status": "Open",
            "paid_to_date": "41_880.40".replace("_", ""),
            "case_reserve": f"{money(CLAIM_CLOSED_STALE_CASE):.2f}",
            "as_of_date": "2025-12-05",
            "large_loss_flag": "N",
            "tpa_note": "Final indemnity issued 2025-12-12; file closed — inventory lag",
        }
    )

    # Duplicate GL claim — older higher case, newer lower case
    base.append(
        {
            "claim_number": CLAIM_DUP,
            "accident_year": 2025,
            "line_of_business": "General Liability",
            "claimant_or_desc": "Terminal visitor",
            "status": "Open",
            "paid_to_date": "4_220.18".replace("_", ""),
            "case_reserve": f"{money(CLAIM_DUP_CASE_OLD):.2f}",
            "as_of_date": "2025-11-30",
            "large_loss_flag": "N",
            "tpa_note": "Preliminary case",
        }
    )
    base.append(
        {
            "claim_number": CLAIM_DUP,
            "accident_year": 2025,
            "line_of_business": "General Liability",
            "claimant_or_desc": "Terminal visitor",
            "status": "Open",
            "paid_to_date": "6_880.40".replace("_", ""),
            "case_reserve": f"{money(CLAIM_DUP_CASE_NEW):.2f}",
            "as_of_date": "2025-12-28",
            "large_loss_flag": "N",
            "tpa_note": "Adjusted after defense counsel update",
        }
    )

    # VOID
    base.append(
        {
            "claim_number": CLAIM_VOID,
            "accident_year": 2024,
            "line_of_business": "Commercial Auto",
            "claimant_or_desc": "Cancelled FNOL",
            "status": "VOID",
            "paid_to_date": "0.00",
            "case_reserve": "99999.00",
            "as_of_date": "2025-01-02",
            "large_loss_flag": "N",
            "tpa_note": "Voided extract artifact — ignore",
        }
    )

    # Extra thin claims for depth (realistic IDs/names; economics unchanged)
    filler = [
        (2022, "Commercial Auto", "CLM-AUTO-2022-1033", "W. Carmichael"),
        (2023, "General Liability", "CLM-GL-2023-0442", "Loading bay leak"),
        (2024, "Commercial Auto", "CLM-AUTO-2024-1190", "V. Serrano"),
        (2025, "General Liability", "CLM-GL-2025-0317", "Sidewalk trip"),
        (2022, "Commercial Auto", "CLM-AUTO-2022-1564", "Q. Bartlett"),
        (2023, "General Liability", "CLM-GL-2023-0701", "Gate hinge injury"),
        (2024, "Commercial Auto", "CLM-AUTO-2024-1628", "Z. Ingram"),
        (2025, "General Liability", "CLM-GL-2025-0559", "Scale house claim"),
    ]
    for i, (ay, lob, claim_no, claimant) in enumerate(filler):
        base.append(
            {
                "claim_number": claim_no,
                "accident_year": ay,
                "line_of_business": lob,
                "claimant_or_desc": claimant,
                "status": "Open",
                "paid_to_date": f"{money(3_110.18 + i * 880.4):.2f}",
                "case_reserve": f"{money(2_220.40 + i * 640.18):.2f}",
                "as_of_date": as_of,
                "large_loss_flag": "N",
                "tpa_note": "",
            }
        )

    rng.shuffle(base)
    return base


def write_open_claims(path: Path) -> None:
    rows = open_claim_rows()
    fields = [
        "claim_number",
        "accident_year",
        "line_of_business",
        "claimant_or_desc",
        "status",
        "paid_to_date",
        "case_reserve",
        "as_of_date",
        "large_loss_flag",
        "tpa_note",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def write_memo(path: Path) -> None:
    text = f"""
{ENTITY}
Primary Package Reserve Factor Memo — Year-End {EVAL_DATE}
Policy {POLICY} | Renewal target {RENEWAL_DATE}
Prepared for Midwest Regional Underwriting Desk
Consulting Actuary: Northbridge Reserve Advisors LLC

Purpose
This memo sets the governing factors and claim-handling rules for the {EVAL_DATE} loss
reserve opinion on Crowhaven Logistics. Crowhaven buys a fully insured primary commercial
auto liability and general liability package. There is no self-insured retention, no
deductible funding corridor, and no captive. No per-claim retention cap is applied to
case reserves because no such cap is specified in the governing materials.

Evaluation basis
Use the TPA cumulative paid triangle (latest year-end diagonal only) and the open-claim
case listing extracted as of {EVAL_DATE}. Ignore interim 6-month and 18-month snapshot
rows that appear in the triangle extract; those are warehouse artifacts from mid-year
pricing drafts.

Governing paid loss development factors (to ultimate)
The year-end paid LDFs below replace every mid-year draft factor printed on the triangle
workbook Definitions sheet. The Definitions table labeled DRAFT 2025-06-30 is superseded
as of 2025-12-15 and must not be used for this opinion.

Development age (months) | Paid LDF to ultimate
12 | {LDF[12]:.2f}
24 | {LDF[24]:.2f}
36 | {LDF[36]:.2f}
48 | {LDF[48]:.2f}

Apply each accident year's LDF to attritional cumulative paid at the matching maturity:
AY 2022 at 48 months, AY 2023 at 36 months, AY 2024 at 24 months, AY 2025 at 12 months.

Large-loss carve-out — {CLAIM_LARGE}
{CLAIM_LARGE} is a multi-vehicle severity claim in accident year 2023. Cumulative paid on
that claim is ${CLAIM_LARGE_PAID:,.2f} and the TPA case reserve is ${CLAIM_LARGE_CASE:,.2f}.
For triangle development, remove that claim's paid dollars from AY 2023 cumulative paid
before multiplying by the 36-month LDF. Do not generate pure IBNR on the carved-out paid.
Hold the large-loss case reserve in the case inventory. Indicated ultimate for the large
loss equals paid plus case (no additional pure IBNR). Attritional pure IBNR for AY 2023
equals attritional ultimate minus attritional paid minus attritional cleaned case.

Case inventory cleaning rules
1. Drop any claim whose TPA note states the file closed after a final payment even if
   status still reads Open. Specifically, remove {CLAIM_CLOSED_STALE} (final payment
   2025-12-12; inventory lag).
2. If the same claim_number appears more than once, keep the row with the latest
   as_of_date and discard earlier duplicates. For {CLAIM_DUP}, keep the 2025-12-28 row
   (case ${CLAIM_DUP_CASE_NEW:,.2f}) and discard the 2025-11-30 preliminary case of
   ${CLAIM_DUP_CASE_OLD:,.2f}.
3. Ignore rows with status VOID, including {CLAIM_VOID}.
4. Expected salvage or subrogation that has not been received in cash must not reduce
   case or pure IBNR.

Indicated ultimate and pure IBNR
For each accident year after large-loss carve-out:
  Attritional ultimate = attritional cumulative paid × governing LDF
  Pure IBNR = max(0, attritional ultimate − attritional cumulative paid − attritional cleaned case)
Total reserve opinion = cleaned case (all retained open claims, including large-loss case)
  + pure IBNR across accident years.
Do not double-count case inside pure IBNR.

Premium and referral thresholds
Subject earned premium for the experience used in this opinion is ${EARNED_PREMIUM:,.2f}.
Target ultimate loss ratio for quote authority is {TARGET_LR:.1%} of that premium.
If indicated ultimate loss ratio exceeds {REFER_LR:.1%}, recommend Refer to regional
referral underwriting. If indicated ultimate loss ratio exceeds {DECLINE_LR:.1%},
recommend Decline. Otherwise recommend Quote at standard authority.

Booked context
The prior booked total reserve (case + IBNR combined) at 2024-12-31 was
${BEGINNING_BOOKED_RESERVE:,.2f}. Use that figure only for narrative context in the
opinion; this assignment is not a self-insured liability rollforward.

What this memo does not provide
This memo supplies governing factors and claim-handling rules only. Final reserve
totals, the cleaned open inventory, and the Quote / Refer / Decline disposition are
determined in the underwriting reserve opinion workbook after the triangle,
inventory, and rules above are reconciled.

Distribution
For underwriting desk use on Crowhaven Logistics {RENEWAL_DATE} renewal pricing only.
Not for reclaim to the insured.
""".strip()
    # Expand slightly for word count with authentic paragraphs
    extra = """

Additional working notes for desk reviewers
Commercial auto dominates Crowhaven's historical paid emergence; GL is thinner but more
volatile on premise and dock claims at the Indianapolis and Louisville terminals. When
an accident year shows Mixed Auto/GL on the triangle roll-up, still apply a single paid
LDF by maturity — do not blend stale mid-year Auto-only and GL-only draft factors from
the June pricing pack. Those line-level drafts were never adopted for year-end.

Case reserves are not reduced by an informal retention-style ceiling; Crowhaven's
primary policy responds within limits on a fully insured basis. Excess above policy
limits is outside this opinion's scope and is not a self-insured layer on Crowhaven's
balance sheet.

Document control: Memo YE-2025-Crowhaven-LDF-03, supersedes YE-2025-Crowhaven-LDF-02
(mid-year). Questions to Northbridge Reserve Advisors, Midwest casualty practice.
"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text + extra, encoding="utf-8")


def main() -> None:
    write_triangle(TASK_INPUTS / "crowhaven-logistics_loss_triangle.xlsx")
    write_open_claims(TASK_INPUTS / "crowhaven-logistics_open_claims.csv")
    write_memo(TASK_INPUTS / "actuarial_factor_memo.txt")
    from sanitize_office import sanitize_xlsx

    sanitize_xlsx(TASK_INPUTS / "crowhaven-logistics_loss_triangle.xlsx")
    print(f"Wrote inputs to {TASK_INPUTS}")


if __name__ == "__main__":
    main()
